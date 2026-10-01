#!/usr/bin/env python3
"""Build a .pptx deck from a deck spec + template map.

Usage: build.py --deck deck.json --map template-map.json [--template T]
                -o out.pptx
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

import jsonschema
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.opc.package import PackURI, Part, XmlPart
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
SCHEMA_DIR = Path(__file__).resolve().parents[1] / "schema"
MIN_NODE_W = Inches(1.1)

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
}
# part prefixes that must be deep-copied per slide instead of shared
CLONE_PREFIXES = ("/ppt/charts/", "/ppt/diagrams/", "/ppt/embeddings/")

SHAPE_TAGS = {
    qn("p:sp"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:pic"),
    qn("p:cxnSp"), qn("p:contentPart"),
}


class BuildError(Exception):
    pass


def fail(msg):
    raise BuildError(msg)


def load_json(path):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        fail(f"cannot read {path}: {e}")


def check_schema(instance, schema_name, label):
    schema = json.loads((SCHEMA_DIR / schema_name).read_text(encoding="utf-8"))
    try:
        jsonschema.validate(instance, schema)
    except jsonschema.ValidationError as e:
        where = "/".join(str(p) for p in e.absolute_path) or "(root)"
        fail(f"{label} violates {schema_name}: {e.message} (at {where})")


# ---------- validation ----------

def _check_orgchart(node, where):
    if not isinstance(node, dict) or not isinstance(node.get("name"), str):
        fail(f"{where}: orgchart node must be an object with a 'name' string")
    if "role" in node and not isinstance(node["role"], str):
        fail(f"{where}: orgchart 'role' must be a string")
    children = node.get("children", [])
    if not isinstance(children, list):
        fail(f"{where}: orgchart 'children' must be an array")
    for child in children:
        _check_orgchart(child, where)


def check_value(value, slot_def, where):
    t = slot_def["type"]
    if t == "text":
        if not isinstance(value, str):
            fail(f"{where}: slot expects a string")
    elif t == "list":
        if not (isinstance(value, list)
                and all(isinstance(v, str) for v in value)):
            fail(f"{where}: slot expects an array of strings")
    elif t == "table":
        if not isinstance(value, dict) or "header" not in value \
                or "rows" not in value:
            fail(f"{where}: table slot expects {{'header': [...], "
                 f"'rows': [[...], ...]}}")
        if not isinstance(value["header"], list) \
                or not isinstance(value["rows"], list):
            fail(f"{where}: table 'header'/'rows' must be arrays")
        for row in value["rows"]:
            if not isinstance(row, list):
                fail(f"{where}: table rows must be arrays")
        if "col_widths" in value and not isinstance(value["col_widths"], list):
            fail(f"{where}: table 'col_widths' must be an array of numbers")
    elif t == "orgchart":
        _check_orgchart(value, where)
    elif t == "timeline":
        if not isinstance(value, list):
            fail(f"{where}: timeline slot expects an array")
        for step in value:
            if not isinstance(step, dict) or "label" not in step \
                    or "period" not in step:
                fail(f"{where}: timeline step needs 'label' and 'period'")
    elif t == "chart":
        _check_chart(value, where)


def _check_chart(value, where):
    if not isinstance(value, dict):
        fail(f"{where}: chart slot expects an object")
    if "type" in value and value["type"] not in CHART_TYPES:
        fail(f"{where}: chart 'type' must be one of {sorted(CHART_TYPES)}")
    cats = value.get("categories")
    series = value.get("series")
    if not isinstance(cats, list) or not cats \
            or not all(isinstance(c, (str, int, float)) for c in cats):
        fail(f"{where}: chart 'categories' must be a non-empty array")
    if not isinstance(series, list) or not series:
        fail(f"{where}: chart 'series' must be a non-empty array")
    for s in series:
        if not isinstance(s, dict) or not isinstance(s.get("name"), str):
            fail(f"{where}: each chart series needs a 'name' string")
        vals = s.get("values")
        if not isinstance(vals, list) \
                or not all(isinstance(v, (int, float))
                           and not isinstance(v, bool) for v in vals):
            fail(f"{where}: series '{s.get('name')}' values must be "
                 f"an array of numbers")
        if len(vals) != len(cats):
            fail(f"{where}: series '{s['name']}' has {len(vals)} values "
                 f"but there are {len(cats)} categories")
    if "number_format" in value \
            and not isinstance(value["number_format"], str):
        fail(f"{where}: chart 'number_format' must be a string")


def effective_slots(spec, comp, idx):
    slots = dict(spec.get("slots") or {})
    if "title" in comp["slots"] and "title" not in slots:
        message = spec.get("message")
        if message:
            slots["title"] = message
    where = f"slide {idx} (component '{spec['component']}')"
    for name in slots:
        if name not in comp["slots"]:
            fail(f"{where}: unknown slot '{name}' "
                 f"(defined: {sorted(comp['slots'])})")
    for name, slot_def in comp["slots"].items():
        if slot_def.get("required") and name not in slots:
            fail(f"{where}: required slot '{name}' is missing")
    for name, value in slots.items():
        check_value(value, comp["slots"][name], f"{where}, slot '{name}'")
    return slots


def validate_deck(deck, tmap):
    check_schema(deck, "deck.schema.json", "deck.json")
    check_schema(tmap, "template-map.schema.json", "template-map.json")
    for i, spec in enumerate(deck["slides"], start=1):
        comp_name = spec["component"]
        if comp_name not in tmap["components"]:
            fail(f"slide {i}: component '{comp_name}' not in template map "
                 f"(have: {sorted(tmap['components'])})")
        comp = tmap["components"][comp_name]
        if comp.get("kind") == "content":
            if not spec.get("message"):
                fail(f"slide {i}: content slide requires a non-empty "
                     f"'message' (it is the slide's action title)")
            if not spec.get("notes"):
                fail(f"slide {i}: content slide requires speaker 'notes'")
        effective_slots(spec, comp, i)


# ---------- shape tree / geometry helpers ----------

def _el_name(el):
    """Name stored in the element's cNvPr, or None."""
    for child in el.iter():
        if child.tag.endswith("}cNvPr") and "name" in child.attrib:
            return child.get("name")
    return None


def _shape_children(container_el):
    """Direct shape children of an spTree or grpSp element."""
    return [c for c in container_el if c.tag in SHAPE_TAGS]


def _xfrm_el(el):
    """xfrm element of a shape (a:xfrm under *Pr, p:xfrm on graphicFrame)."""
    if el.tag == qn("p:graphicFrame"):
        return el.find(qn("p:xfrm"))
    if el.tag == qn("p:grpSp"):
        pr = el.find(qn("p:grpSpPr"))
    elif el.tag == qn("p:cxnSp"):
        pr = el.find(qn("p:spPr"))
    else:
        pr = el.find(qn("p:spPr"))
    if pr is None:
        return None
    return pr.find(qn("a:xfrm"))


def _raw_bbox(el):
    """(x, y, w, h) in the element's own coordinate frame, or None."""
    xfrm = _xfrm_el(el)
    if xfrm is None:
        return None
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    if off is None or ext is None:
        return None
    return (int(off.get("x", 0)), int(off.get("y", 0)),
            int(ext.get("cx", 0)), int(ext.get("cy", 0)))


def absolute_bbox(el):
    """(x, y, w, h) in slide coordinates, applying ancestor group transforms."""
    bbox = _raw_bbox(el)
    if bbox is None:
        return None
    x, y, w, h = bbox
    parent = el.getparent()
    while parent is not None and parent.tag == qn("p:grpSp"):
        xfrm = _xfrm_el(parent)
        if xfrm is not None:
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            ch_off, ch_ext = (xfrm.find(qn("a:chOff")),
                              xfrm.find(qn("a:chExt")))
            if off is not None and ext is not None \
                    and ch_off is not None and ch_ext is not None:
                cx, cy = int(ch_ext.get("cx", 0)), int(ch_ext.get("cy", 0))
                ex, ey = int(ext.get("cx", 0)), int(ext.get("cy", 0))
                sx = ex / cx if cx else 1.0
                sy = ey / cy if cy else 1.0
                x = int(off.get("x", 0)) + int((x - int(ch_off.get("x", 0))) * sx)
                y = int(off.get("y", 0)) + int((y - int(ch_off.get("y", 0))) * sy)
                w = int(w * sx)
                h = int(h * sy)
        parent = parent.getparent()
    return (x, y, w, h)


def _find_all(container_el, name):
    """All shape elements named `name` under a container, group-aware."""
    hits = []
    for c in _shape_children(container_el):
        if _el_name(c) == name:
            hits.append(c)
        if c.tag == qn("p:grpSp"):
            hits.extend(_find_all(c, name))
    return hits


def find_shape_el(slide, name, where):
    """Find a shape element by name, recursing into groups.

    Raises if zero or more than one match.
    """
    hits = _find_all(slide.shapes._spTree, name)
    if not hits:
        fail(f"{where}: shape '{name}' not found on source slide")
    if len(hits) > 1:
        fail(f"{where}: shape name '{name}' matches {len(hits)} shapes "
             f"on the slide; template-map targets must be unique")
    return hits[0]


def remove_shape_el(el):
    """Remove a shape element; remove ancestor groups that become empty."""
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)
    while parent is not None and parent.tag == qn("p:grpSp") \
            and not _shape_children(parent):
        gp = parent.getparent()
        if gp is not None:
            gp.remove(parent)
        parent = gp


def _graphic_uri(el):
    """graphicData uri of a p:graphicFrame, or None."""
    if el.tag != qn("p:graphicFrame"):
        return None
    gd = el.find(f"{qn('a:graphic')}/{qn('a:graphicData')}")
    return gd.get("uri") if gd is not None else None


_KIND_JA = {
    "shape": "図形", "group": "グループ", "picture": "画像",
    "connector": "コネクタ", "smartart": "SmartArt", "chart": "グラフ",
    "table": "表", "ole": "OLEオブジェクト", "graphic": "グラフィック枠",
    "other": "図形",
}


def _target_kind(el):
    """Classify a shape element for region/prototype handling."""
    if el.tag == qn("p:sp"):
        return "shape"
    if el.tag == qn("p:grpSp"):
        return "group"
    if el.tag == qn("p:pic"):
        return "picture"
    if el.tag == qn("p:cxnSp"):
        return "connector"
    if el.tag == qn("p:graphicFrame"):
        uri = _graphic_uri(el) or ""
        if uri == DIAGRAM_URI:
            return "smartart"
        if uri == CHART_URI:
            return "chart"
        if uri.endswith("/table"):
            return "table"
        if "ole" in uri.lower():
            return "ole"
        return "graphic"
    return "other"


def _ph_bbox(slide, ph, where):
    """Resolved (x, y, w, h) of a slide placeholder via layout/master
    inheritance."""
    box = (ph.left, ph.top, ph.width, ph.height)
    if all(v is not None for v in box):
        return tuple(int(v) for v in box)
    idx = ph.placeholder_format.idx
    for container in (slide.slide_layout, slide.slide_layout.slide_master):
        for ref in container.placeholders:
            if ref.placeholder_format.idx == idx:
                box = (ref.left, ref.top, ref.width, ref.height)
                if all(v is not None for v in box):
                    return tuple(int(v) for v in box)
    fail(f"{where}: placeholder idx {idx} has no geometry")


# ---------- text fill (keeps prototype formatting) ----------

def _txbody(sp_el):
    for tag in (qn("p:txBody"), qn("a:txBody")):
        el = sp_el.find(tag)
        if el is not None:
            return el
    fail("shape has no text body")


def _clone_run(proto_p):
    proto_r = proto_p.find(qn("a:r"))
    if proto_r is not None:
        new_r = copy.deepcopy(proto_r)
        for t in new_r.findall(qn("a:t")):
            new_r.remove(t)
        return new_r
    new_r = etree.Element(qn("a:r"))
    epr = proto_p.find(qn("a:endParaRPr"))
    if epr is not None:
        rpr = copy.deepcopy(epr)
        rpr.tag = qn("a:rPr")
        if "lang" not in rpr.attrib:
            rpr.set("lang", "ja-JP")
        new_r.append(rpr)
    return new_r


def _fill_txbody(txBody, lines, per_paragraph=False):
    protos = [copy.deepcopy(p) for p in txBody.findall(qn("a:p"))]
    if not protos:
        fail("text body has no paragraph to use as format prototype")
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for i, line in enumerate(lines):
        proto_p = protos[min(i, len(protos) - 1)] if per_paragraph \
            else protos[0]
        new_p = copy.deepcopy(proto_p)
        for child in list(new_p):
            if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                new_p.remove(child)
        new_r = _clone_run(proto_p)
        t = etree.SubElement(new_r, qn("a:t"))
        t.text = line
        pPr = new_p.find(qn("a:pPr"))
        if pPr is not None:
            pPr.addnext(new_r)
        else:
            new_p.insert(0, new_r)
        txBody.append(new_p)


def fill_text(shape_el, value):
    lines = str(value).split("\n")
    _fill_txbody(_txbody(shape_el), lines)


def fill_list(shape_el, value):
    _fill_txbody(_txbody(shape_el), [str(v) for v in value])


def _text_bearing_sps(el):
    """Descendant p:sp elements containing at least one run, in doc order."""
    out = []
    for sp in el.iter(qn("p:sp")):
        tx = sp.find(qn("p:txBody"))
        if tx is None:
            continue
        if any(ch.tag in (qn("a:r"), qn("a:fld"), qn("a:br"))
               for p in tx.findall(qn("a:p")) for ch in p):
            out.append(sp)
    return out


def fill_node_text(proto_el, lines):
    """Distribute node lines into a node prototype.

    Single shape: paragraph i gets line i. Group: the i-th text-bearing
    child gets line i, extra lines appended to the last one; a single
    text-bearing child receives every line as paragraphs.
    """
    if proto_el.tag == qn("p:grpSp"):
        targets = _text_bearing_sps(proto_el)
        if not targets:
            return
        if len(targets) == 1:
            _fill_txbody(_txbody(targets[0]), lines, per_paragraph=True)
            return
        n = len(targets)
        chunks = list(lines[:n - 1])
        chunks.append("\n".join(lines[n - 1:]) if lines else "")
        for i, sp in enumerate(targets):
            chunk = chunks[i] if i < len(chunks) else ""
            _fill_txbody(_txbody(sp), chunk.split("\n") if chunk else [""])
        return
    _fill_txbody(_txbody(proto_el), lines, per_paragraph=True)


# ---------- slide assembly ----------

def find_layout(prs, name):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    fail(f"layout '{name}' not found in template "
         f"(have: {[l.name for l in prs.slide_layouts]})")


class BuildCtx:
    """Per-build state: fallback report + donor lookups."""

    def __init__(self, prs, n_original):
        self.prs = prs
        self.n_original = n_original
        self.fallbacks = []
        self.native = 0
        self.slide_i = 0
        self.component = ""
        self.slot = ""
        self._donor_cache = None

    def fb(self, target, reason, substitute, style_source="theme"):
        self.fallbacks.append({
            "slide": self.slide_i,
            "component": self.component,
            "slot": self.slot,
            "target": str(target),
            "reason": reason,
            "substitute": substitute,
            "style_source": style_source,
        })


def _rel_ids(el):
    ids = []
    for e in el.iter():
        for attr, val in e.attrib.items():
            if attr.startswith("{" + R + "}"):
                ids.append((e, attr, val))
    return ids


def _should_clone_part(part):
    return any(str(part.partname).startswith(p) for p in CLONE_PREFIXES)


def _used_partnames(pkg):
    return {str(p.partname) for p in pkg.iter_parts()}


def _fresh_partname(used, partname):
    """Unused partname derived from `partname` by bumping its number."""
    s = str(partname)
    m = re.match(r"^(.*?)(\d+)(\.[^.]*)$", s)
    if not m:
        base = s + "-copy"
        name = base
        i = 1
        while name in used:
            i += 1
            name = f"{base}{i}"
        used.add(name)
        return PackURI(name)
    pre, num, ext = m.groups()
    i = int(num) + 1
    while f"{pre}{i}{ext}" in used:
        i += 1
    name = f"{pre}{i}{ext}"
    used.add(name)
    return PackURI(name)


def _clone_part(pkg, src_part, used):
    """Deep-copy a part (XML or binary) and recursively clone its
    related parts; remap r:* attributes inside copied XML."""
    newname = _fresh_partname(used, src_part.partname)
    if isinstance(src_part, XmlPart):
        new_el = copy.deepcopy(src_part._element)
        # keep the concrete class (e.g. ChartPart) so .chart etc. work
        new_part = type(src_part)(newname, src_part.content_type, pkg,
                                  new_el)
    else:
        new_el = None
        new_part = type(src_part)(newname, src_part.content_type, pkg,
                                  src_part.blob)
    rid_map = {}
    for rid, rel in src_part.rels.items():
        if rel.is_external:
            nrid = new_part.relate_to(rel.target_ref, rel.reltype,
                                      is_external=True)
        else:
            tp = rel.target_part
            ntp = _clone_part(pkg, tp, used) if _should_clone_part(tp) else tp
            nrid = new_part.relate_to(ntp, rel.reltype)
        rid_map[rid] = nrid
    if new_el is not None:
        for e, attr, rid in _rel_ids(new_el):
            if rid in rid_map:
                e.set(attr, rid_map[rid])
    return new_part


def clone_slide(prs, src_slide):
    dest = prs.slides.add_slide(src_slide.slide_layout)
    for ph in list(dest.placeholders):
        ph._element.getparent().remove(ph._element)
    copied = []
    for shape in src_slide.shapes:
        el = copy.deepcopy(shape._element)
        dest.shapes._spTree.append(el)
        copied.append(el)
    # remap relationship ids (images, media, hyperlinks, charts, ...)
    src_rels = src_slide.part.rels
    pkg = prs.part.package
    used = _used_partnames(pkg)
    for el in copied:
        for e, attr, rid in _rel_ids(el):
            if rid not in src_rels:
                continue
            rel = src_rels[rid]
            if rel.is_external:
                new_rid = dest.part.relate_to(rel.target_ref, rel.reltype,
                                              is_external=True)
            else:
                tp = rel.target_part
                if _should_clone_part(tp):
                    tp = _clone_part(pkg, tp, used)
                new_rid = dest.part.relate_to(tp, rel.reltype)
            e.set(attr, new_rid)
    return dest


def renumber_shape_ids(slide):
    """Give every shape a unique positive cNvPr id.

    Cloned prototype copies share the prototype's id; connectors may also
    collide. The spTree's own cNvPr (id 1) is left alone. stCxn/endCxn
    endpoint references are remapped when the old id is unambiguous.
    """
    tree = slide.shapes._spTree
    root_nv = tree.find(qn("p:nvGrpSpPr"))
    root = root_nv.find(qn("p:cNvPr")) if root_nv is not None else None
    next_id = 2
    remap = {}
    for e in tree.iter(qn("p:cNvPr")):
        if e is root:
            continue
        old = e.get("id")
        new = str(next_id)
        next_id += 1
        if old is not None:
            remap.setdefault(old, []).append(new)
        e.set("id", new)
    for e in tree.iter():
        if e.tag in (qn("a:stCxn"), qn("a:endCxn")):
            old = e.get("id")
            targets = remap.get(old)
            if targets and len(targets) == 1:
                e.set("id", targets[0])


def remove_unfilled_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph._element.tag == qn("p:graphicFrame"):
            continue  # populated table/chart frames are content, not prompts
        text = "".join(
            t.text or "" for t in ph._element.iter(qn("a:t")))
        if not text.strip():
            ph._element.getparent().remove(ph._element)


def inherit_placeholder_geometry(slide):
    """Stamp resolved layout/master geometry onto slide placeholders so
    renderers that skip layout-level xfrm inheritance (LibreOffice) place
    them correctly."""
    layout = slide.slide_layout
    master = layout.slide_master
    for ph in slide.placeholders:
        spPr = ph._element.find(qn("p:spPr"))
        if spPr is not None and spPr.find(qn("a:xfrm")) is not None:
            continue
        idx = ph.placeholder_format.idx
        for container in (layout, master):
            for ref in container.placeholders:
                if ref.placeholder_format.idx == idx:
                    src = ref._element.find(qn("p:spPr"))
                    xfrm = src.find(qn("a:xfrm")) if src is not None else None
                    if xfrm is not None:
                        if spPr is None:
                            spPr = ph._element.find(qn("p:spPr"))
                        if spPr is None:
                            break
                        spPr.append(copy.deepcopy(xfrm))
                        break
            else:
                continue
            break


# ---------- regions ----------

def _drop_frame_rels(slide, frame_el):
    """Drop slide relationships referenced only by `frame_el`, plus any
    unreferenced rels into /ppt/diagrams/ (SmartArt drawing parts)."""
    rids = {v for _, _, v in _rel_ids(frame_el)}
    keep = set()

    def rec(e):
        if e is frame_el:
            return
        for attr, val in e.attrib.items():
            if attr.startswith("{" + R + "}"):
                keep.add(val)
        for c in e:
            rec(c)

    rec(slide.shapes._spTree)
    rels = slide.part.rels
    drop = set()
    for rid in rids - keep:
        if rid in rels:
            drop.add(rid)
    for rid, rel in list(rels.items()):
        if rel.is_external:
            continue
        if str(rel.target_partname).startswith("/ppt/diagrams/") \
                and rid not in keep:
            drop.add(rid)
    for rid in drop:
        slide.part.drop_rel(rid)


def _region_bbox(el, where):
    """Absolute (x, y, w, h) of a sample-mode region element."""
    bbox = absolute_bbox(el)
    if bbox is None:
        fail(f"{where}: region shape has no geometry")
    return bbox


def _remove_region(slide, el):
    """Remove a region element, dropping rels it alone referenced."""
    if el.tag == qn("p:graphicFrame"):
        _drop_frame_rels(slide, el)
    remove_shape_el(el)


def _prototype_element(ctx, slide, slot_def, where, kind,
                       layout_mode=False):
    """Deep-copied node prototype element (removed from the slide), or a
    synthesized theme-style prototype when node_prototype is absent.

    In layout mode the prototype may live on the slide layout's shape
    tree; it is copied but not removed there. A prototype that is not a
    shape/group (picture, OLE, graphicFrame, ...) is unusable: fall back
    to a synthesized node and record it.
    """
    name = slot_def.get("node_prototype")
    if not name:
        return synth_node_proto(kind)
    hits = _find_all(slide.shapes._spTree, name)
    remove = True
    if not hits and layout_mode:
        hits = _find_all(slide.slide_layout.shapes._spTree, name)
        remove = False
    if not hits:
        fail(f"{where}: node prototype '{name}' not found")
    if len(hits) > 1:
        fail(f"{where}: node prototype name '{name}' matches "
             f"{len(hits)} shapes; names must be unique")
    el = hits[0]
    if _target_kind(el) not in ("shape", "group"):
        ctx.fb(name,
               f"ノード雛形が{_KIND_JA.get(_target_kind(el), '図形')}"
               f"のため複製できない",
               "テーマスタイルのノードを生成", "theme")
        if remove:
            if el.tag == qn("p:graphicFrame"):
                _drop_frame_rels(slide, el)
            remove_shape_el(el)
        return synth_node_proto(kind)
    proto = copy.deepcopy(el)
    if remove:
        remove_shape_el(el)
    return proto


def synth_node_proto(kind):
    """Synthesized theme-styled node prototype (p:sp XML string).

    orgchart: rounded rect, accent1 fill, white centered text, name bold /
    role smaller. timeline: 3 paragraphs label/period/detail.
    """
    if kind == "org":
        w, h = 2011680, 822960  # 2.2in x 0.9in
        paras = [(1500, True), (1150, False)]
        radius = "12000"
    else:
        w, h = 2926080, 1828800  # 3.2in x 2.0in
        paras = [(1600, True), (1300, False), (1150, False)]
        radius = "8000"
    p_xml = ""
    for sz, bold in paras:
        b = ' b="1"' if bold else ""
        p_xml += (
            f'<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="ja-JP" sz="{sz}"{b}>'
            f'<a:solidFill><a:schemeClr val="lt1"/></a:solidFill>'
            f'<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/>'
            f'</a:rPr><a:t>node</a:t></a:r></a:p>')
    xml = (
        f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/'
        f'2006/main" xmlns:a="{A}">'
        f'<p:nvSpPr><p:cNvPr id="0" name="SynthNode"/>'
        f'<p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/>'
        f'</a:xfrm>'
        f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" '
        f'fmla="val {radius}"/></a:avLst></a:prstGeom>'
        f'<a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
        f'<a:ln><a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
        f'</a:ln></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square" lIns="91440" rIns="91440" '
        f'tIns="45720" bIns="45720" anchor="mid"/>'
        f'<a:lstStyle/>{p_xml}</p:txBody></p:sp>')
    return etree.fromstring(xml)


# shapes drawn by the builder carry a descr marker so qa can check
# whether a redrawn fallback or a generated node collides with
# content that was kept from the sample slide
FALLBACK_MARK = "pptx-deck-fallback"
GENERATED_MARK = "pptx-deck-generated"


def _mark_el(el, mark):
    """Set descr on the element's own cNvPr."""
    for child in el:
        for c in child.iter(qn("p:cNvPr")):
            c.set("descr", mark)
            return


# ---------- table ----------

def _find_parent_frame(el):
    """Ancestor p:graphicFrame of an element, or None."""
    p = el.getparent()
    while p is not None:
        if p.tag == qn("p:graphicFrame"):
            return p
        p = p.getparent()
    return None


def _tc_fill(tc):
    tcPr = tc.find(qn("a:tcPr"))
    f = tcPr.find(qn("a:solidFill")) if tcPr is not None else None
    return copy.deepcopy(f) if f is not None else None


def _rpr_info(tc):
    """First run property info of a cell: {sz, b, fill(solidFill el)}."""
    if tc is None:
        return {}
    for rpr in tc.iter(qn("a:rPr"), qn("a:endParaRPr")):
        info = {}
        if rpr.get("sz"):
            info["sz"] = rpr.get("sz")
        if rpr.get("b") in ("1", "true"):
            info["b"] = True
        f = rpr.find(qn("a:solidFill"))
        if f is not None:
            info["fill"] = copy.deepcopy(f)
        return info
    return {}


def _extract_donor_fills(tbl):
    """Harvest explicit cell styling from a template table.

    Returns a dict: header_fill, band1, band2, border (solidFill
    elements), header_rpr/body_rpr (sz/b/fill), row_h. None when the
    table has no explicit fills.
    """
    trs = tbl.findall(qn("a:tr"))
    if len(trs) < 2:
        return None
    donor = {}
    header_tc = trs[0].find(qn("a:tc"))
    donor["header_fill"] = _tc_fill(header_tc)
    donor["header_rpr"] = _rpr_info(header_tc)
    donor["body_rpr"] = _rpr_info(trs[1].find(qn("a:tc")))
    donor["row_h"] = trs[1].get("h") or trs[0].get("h")
    row1 = trs[1].find(qn("a:tc"))
    donor["band1"] = _tc_fill(row1)
    donor["band2"] = _tc_fill(trs[2].find(qn("a:tc"))) \
        if len(trs) > 2 else None
    for tcPr in tbl.iter(qn("a:tcPr")):
        for edge in ("a:lnB", "a:lnT", "a:lnL", "a:lnR"):
            ln = tcPr.find(qn(edge))
            f = ln.find(qn("a:solidFill")) if ln is not None else None
            if f is not None:
                donor["border"] = copy.deepcopy(f)
                break
        if donor.get("border") is not None:
            break
    if donor.get("header_fill") is None and donor.get("band1") is None:
        return None
    return donor


def _theme_donor():
    """Fallback styling built from theme colors."""
    def sf(inner):
        e = etree.Element(qn("a:solidFill"))
        e.append(etree.fromstring(inner))
        return e
    lt1 = sf(f'<a:schemeClr xmlns:a="{A}" val="lt1"/>')
    tx1 = sf(f'<a:schemeClr xmlns:a="{A}" val="tx1"/>')
    return {
        "header_fill": sf(f'<a:schemeClr xmlns:a="{A}" val="accent1"/>'),
        "header_rpr": {"sz": "1200", "b": True, "fill": lt1},
        "body_rpr": {"sz": "1200", "fill": tx1},
        "row_h": "365760",  # 0.4in
        "band1": sf(f'<a:schemeClr xmlns:a="{A}" val="lt1"/>'),
        "band2": sf(f'<a:schemeClr xmlns:a="{A}" val="accent1">'
                    f'<a:lumMod val="20000"/><a:lumOff val="80000"/>'
                    f'</a:schemeClr>'),
        "border": sf(f'<a:schemeClr xmlns:a="{A}" val="accent1">'
                     f'<a:lumMod val="40000"/><a:lumOff val="60000"/>'
                     f'</a:schemeClr>'),
    }


def _table_style_id_of(tbl):
    tblPr = tbl.find(qn("a:tblPr"))
    el = tblPr.find(qn("a:tableStyleId")) if tblPr is not None else None
    return el.text if el is not None and el.text else None


def _donor_tables(ctx):
    """(slide_idx, frame_name, tbl) for each table in original template
    slides, in slide order."""
    if ctx._donor_cache is not None:
        return ctx._donor_cache
    found = []
    for i, s in enumerate(list(ctx.prs.slides)[:ctx.n_original],
                          start=1):
        for tbl in s.shapes._spTree.iter(qn("a:tbl")):
            frame = _find_parent_frame(tbl)
            name = _el_name(frame) if frame is not None else "?"
            found.append((i, name, tbl))
    ctx._donor_cache = found
    return found


def resolve_table_donor(ctx):
    """Donor chain for redrawn tables.

    Returns (mode, payload, style_source): mode 'fills' applies explicit
    cell fills, 'styleId' applies a table style GUID.
    """
    donors = _donor_tables(ctx)
    for i, name, tbl in donors:
        d = _extract_donor_fills(tbl)
        if d:
            return "fills", d, f"slide {i} '{name}' (explicit fills)"
    for i, name, tbl in donors:
        sid = _table_style_id_of(tbl)
        if sid:
            return "styleId", sid, f"tableStyleId {sid} from slide {i}"
    default = _default_table_style_id(ctx.prs)
    if default:
        return "styleId", default, "tableStyles.xml def"
    return "fills", _theme_donor(), "theme"


def _apply_donor_fills(tbl, donor):
    trs = tbl.findall(qn("a:tr"))
    if not trs:
        return

    def set_fill(tc, fill):
        if fill is None:
            return
        tcPr = tc.find(qn("a:tcPr"))
        if tcPr is None:
            tcPr = etree.SubElement(tc, qn("a:tcPr"))
        for old in tcPr.findall(qn("a:solidFill")):
            tcPr.remove(old)
        tcPr.insert(0, copy.deepcopy(fill))

    def set_border(tc, fill):
        if fill is None:
            return
        tcPr = tc.find(qn("a:tcPr"))
        if tcPr is None:
            tcPr = etree.SubElement(tc, qn("a:tcPr"))
        for old in tcPr.findall(qn("a:lnB")):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn("a:lnB"))
        ln.set("w", "9525")
        ln.set("cap", "flat")
        ln.append(copy.deepcopy(fill))

    def set_runs(tc, info):
        if not info:
            return
        tx = tc.find(qn("a:txBody"))
        if tx is None:
            return
        for r in tx.iter(qn("a:r")):
            rpr = r.find(qn("a:rPr"))
            if rpr is None:
                rpr = etree.Element(qn("a:rPr"))
                r.insert(0, rpr)
            if info.get("sz"):
                rpr.set("sz", str(info["sz"]))
            if info.get("b"):
                rpr.set("b", "1")
            for old in rpr.findall(qn("a:solidFill")):
                rpr.remove(old)
            if info.get("fill") is not None:
                rpr.insert(0, copy.deepcopy(info["fill"]))

    row_h = donor.get("row_h")
    for ri, tr in enumerate(trs):
        if row_h:
            tr.set("h", str(row_h))
        for tc in tr.findall(qn("a:tc")):
            if ri == 0:
                set_fill(tc, donor.get("header_fill"))
            else:
                band = donor.get("band1") if ri % 2 == 1 \
                    else donor.get("band2")
                if band is None:
                    band = donor.get("band1")
                set_fill(tc, band)
            set_border(tc, donor.get("border"))
            set_runs(tc, donor.get("header_rpr") if ri == 0
                     else donor.get("body_rpr"))


def _apply_table_style(tbl, mode, payload):
    if mode == "styleId":
        _set_table_style(tbl, payload)
    else:
        _apply_donor_fills(tbl, payload)


def _default_table_style_id(prs):
    """`def` style id of ppt/tableStyles.xml, or None."""
    for part in prs.part.package.iter_parts():
        if str(part.partname) == "/ppt/tableStyles.xml":
            try:
                root = etree.fromstring(part.blob)
            except Exception:
                return None
            return root.get("def")
    return None


def _set_table_style(tbl, style_id):
    tblPr = tbl.find(qn("a:tblPr"))
    if tblPr is None:
        tblPr = etree.Element(qn("a:tblPr"))
        tbl.insert(0, tblPr)
        tblPr.set("firstRow", "1")
        tblPr.set("bandRow", "1")
    el = tblPr.find(qn("a:tableStyleId"))
    if el is None:
        el = etree.SubElement(tblPr, qn("a:tableStyleId"))
    el.text = style_id


def _adjust_cols(tbl, n_cols, where):
    grid = tbl.find(qn("a:tblGrid"))
    cols = grid.findall(qn("a:gridCol"))
    trs = tbl.findall(qn("a:tr"))
    while len(cols) < n_cols:
        new_col = copy.deepcopy(cols[-1])
        grid.append(new_col)
        cols.append(new_col)
        for tr in trs:
            tr.append(copy.deepcopy(tr.findall(qn("a:tc"))[-1]))
    while len(cols) > n_cols:
        grid.remove(cols.pop())
        for tr in trs:
            tr.remove(tr.findall(qn("a:tc"))[-1])
    return cols


def _distribute_widths(cols, value, total_w, n_cols, where):
    weights = value.get("col_widths") or [1.0] * n_cols
    if len(weights) != n_cols:
        fail(f"{where}: col_widths has {len(weights)} entries, "
             f"table has {n_cols} columns")
    s = float(sum(weights))
    widths = [int(round(total_w * w / s)) for w in weights]
    widths[-1] += int(total_w) - sum(widths)
    for col, w in zip(cols, widths):
        col.set("w", str(w))


def _fill_table_cells(tbl, header, rows, where):
    trs = tbl.findall(qn("a:tr"))
    n_cols = len(header)
    for row in rows:
        if len(row) != n_cols:
            fail(f"{where}: table row {row} has {len(row)} cells, "
                 f"header has {n_cols}")
    # cycle template body rows so banded fills alternate on extra rows
    body_protos = list(trs[1:]) or [trs[0]]
    while len(trs) < 1 + len(rows):
        new_tr = copy.deepcopy(body_protos[(len(trs) - 1) % len(body_protos)])
        tbl.append(new_tr)
        trs.append(new_tr)
    while len(trs) > 1 + len(rows):
        tbl.remove(trs.pop())
    for tc, text in zip(trs[0].findall(qn("a:tc")), header):
        _fill_txbody(tc.find(qn("a:txBody")), text.split("\n"))
    for tr, row in zip(trs[1:], rows):
        for tc, text in zip(tr.findall(qn("a:tc")), row):
            _fill_txbody(tc.find(qn("a:txBody")), text.split("\n"))


def fill_table(ctx, slide, slot_def, value, where):
    el = find_shape_el(slide, slot_def["target"], where)
    kind = _target_kind(el)
    gd = el.find(f"{qn('a:graphic')}/{qn('a:graphicData')}")
    tbl = gd.find(qn("a:tbl")) if kind == "table" and gd is not None \
        else None
    header = [str(c) for c in value["header"]]
    rows = [[str(c) for c in row] for row in value["rows"]]
    n_cols = len(header)
    if tbl is not None and len(tbl.findall(qn("a:tr"))) >= 2:
        # native path: reuse the template table's styling
        cols = _adjust_cols(tbl, n_cols, where)
        bbox = absolute_bbox(el)
        total = bbox[2] if bbox else int(el.find(qn("p:xfrm")).find(
            qn("a:ext")).get("cx"))
        _distribute_widths(cols, value, total, n_cols, where)
        _fill_table_cells(tbl, header, rows, where)
        return
    # fallback: the target exists but is not a usable table
    bbox = absolute_bbox(el)
    if bbox is None:
        fail(f"{where}: table slot target '{slot_def['target']}' "
             f"has no geometry")
    if kind == "table":
        reason = "対象の表にデータ行が無い（ヘッダ＋1行以上が必要）"
    else:
        reason = (f"対象が{_KIND_JA.get(kind, '図形')}のため、"
                  f"表として使えない")
    _remove_region(slide, el)
    frame = slide.shapes.add_table(1 + len(rows), n_cols, *bbox)
    try:
        frame.name = slot_def["target"]
    except Exception:
        pass
    _mark_el(frame._element, FALLBACK_MARK)
    new_tbl = frame._element.graphic.graphicData.tbl
    cols = _adjust_cols(new_tbl, n_cols, where)
    _distribute_widths(cols, value, bbox[2], n_cols, where)
    _fill_table_cells(new_tbl, header, rows, where)
    # donor styling goes last: it restyles the runs just written
    mode, payload, src = resolve_table_donor(ctx)
    _apply_table_style(new_tbl, mode, payload)
    ctx.fb(slot_def["target"], reason,
           "テンプレ内の表のスタイルを参考に新しい表を描画", src)


def fill_table_layout(ctx, slide, slot_def, value, ph, bbox, where):
    header = [str(c) for c in value["header"]]
    rows = [[str(c) for c in row] for row in value["rows"]]
    n_cols, n_rows = len(header), 1 + len(rows)
    if hasattr(ph, "insert_table"):
        frame = ph.insert_table(n_rows, n_cols)
    else:
        frame = slide.shapes.add_table(n_rows, n_cols, *bbox)
        ph._element.getparent().remove(ph._element)
    tbl = frame._element.graphic.graphicData.tbl
    cols = _adjust_cols(tbl, n_cols, where)
    _distribute_widths(cols, value, bbox[2], n_cols, where)
    _fill_table_cells(tbl, header, rows, where)
    # styling goes last: donor fills restyle the runs just written
    style_id = slot_def.get("table_style_id") \
        or _default_table_style_id(ctx.prs)
    if style_id:
        _set_table_style(tbl, style_id)
    else:
        mode, payload, src = resolve_table_donor(ctx)
        _apply_table_style(tbl, mode, payload)
        _mark_el(frame._element, FALLBACK_MARK)
        ctx.fb(ph.placeholder_format.idx,
               "適用できる表スタイル（tableStyleId）がテンプレに無い",
               "ドナー表の配色をコピーして描画", src)


# ---------- orgchart / timeline ----------

def _tree_depth(node):
    kids = node.get("children") or []
    return 1 + max((_tree_depth(k) for k in kids), default=0)


def _count_leaves(node):
    kids = node.get("children") or []
    if not kids:
        return 1
    return sum(_count_leaves(k) for k in kids)


def _set_xfrm(sp_el, x, y, w, h):
    """Set off/ext of an sp or grpSp element (grpSp keeps chOff/chExt)."""
    if sp_el.tag == qn("p:grpSp"):
        pr = sp_el.find(qn("p:grpSpPr"))
    else:
        pr = sp_el.find(qn("p:spPr"))
    if pr is None:
        fail("node prototype has no spPr/grpSpPr")
    xfrm = pr.find(qn("a:xfrm"))
    if xfrm is None:
        xfrm = etree.Element(qn("a:xfrm"))
        pr.insert(0, xfrm)
    for tag in ("a:off", "a:ext"):
        if xfrm.find(qn(tag)) is None:
            etree.SubElement(xfrm, qn(tag))
    xfrm.find(qn("a:off")).set("x", str(int(x)))
    xfrm.find(qn("a:off")).set("y", str(int(y)))
    xfrm.find(qn("a:ext")).set("cx", str(int(w)))
    xfrm.find(qn("a:ext")).set("cy", str(int(h)))


def _proto_size(proto_el):
    bbox = _raw_bbox(proto_el)
    if bbox is None:
        fail("node prototype has no geometry")
    return bbox[2], bbox[3]


def _node_text(node):
    lines = [node["name"]]
    if node.get("role"):
        lines.append(node["role"])
    return lines


def _style_connector(conn, scheme_color):
    ln = conn.line._get_or_add_ln()
    ln.set("w", "19050")
    for child in list(ln):
        if child.tag == qn("a:solidFill"):
            ln.remove(child)
    fill = etree.Element(qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:schemeClr"))
    clr.set("val", scheme_color)
    prst = ln.find(qn("a:prstDash"))
    ln.insert(list(ln).index(prst) if prst is not None else 0, fill)


def draw_orgchart(slide, tree, bbox, proto_el, connector_color, where):
    if _tree_depth(tree) > 3:
        fail(f"{where}: orgchart supports at most 3 levels")
    rx, ry, rw, rh = bbox
    proto_w, ph = _proto_size(proto_el)
    leaves = _count_leaves(tree)
    levels = _tree_depth(tree)
    node_w = min(proto_w, Emu(int(rw / leaves - Inches(0.15))))
    if node_w < MIN_NODE_W:
        fail(f"{where}: orgchart does not fit ({leaves} leaf nodes in "
             f"{rw / 914400:.2f}in; each box would be < 1.1in wide)")
    band_h = rh / levels
    leaf_i = [0]
    pos = {}

    def walk(node, level):
        kids = node.get("children") or []
        if kids:
            xs = [walk(k, level + 1) for k in kids]
            cx = (xs[0] + xs[-1]) / 2
        else:
            cx = rx + (leaf_i[0] + 0.5) * (rw / leaves)
            leaf_i[0] += 1
        y = ry + level * band_h + (band_h - ph) / 2
        pos[id(node)] = (cx, y)
        return cx

    walk(tree, 0)

    def emit(node):
        kids = node.get("children") or []
        cx, y = pos[id(node)]
        el = copy.deepcopy(proto_el)
        _set_xfrm(el, cx - node_w / 2, y, node_w, ph)
        _mark_el(el, GENERATED_MARK)
        fill_node_text(el, _node_text(node))
        slide.shapes._spTree.append(el)
        for k in kids:
            kcx, ky = pos[id(k)]
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.ELBOW, Emu(int(cx)), Emu(int(y + ph)),
                Emu(int(kcx)), Emu(int(ky)))
            conn.name = "OrgConnector"
            _style_connector(conn, connector_color)
            emit(k)

    emit(tree)


def draw_timeline(slide, steps, bbox, proto_el, where):
    rx, ry, rw, rh = bbox
    n = len(steps)
    proto_w, node_h = _proto_size(proto_el)
    node_w = min(proto_w, Emu(int(rw / n - Inches(0.15))))
    if node_w < MIN_NODE_W:
        fail(f"{where}: timeline does not fit ({n} steps in "
             f"{rw / 914400:.2f}in; each box would be < 1.1in wide)")
    y = ry + (rh - node_h) / 2
    for i, step in enumerate(steps):
        cx = rx + (i + 0.5) * (rw / n)
        el = copy.deepcopy(proto_el)
        _set_xfrm(el, cx - node_w / 2, y, node_w, node_h)
        _mark_el(el, GENERATED_MARK)
        lines = [step["label"], step["period"]]
        if step.get("detail"):
            lines.append(step["detail"])
        fill_node_text(el, lines)
        slide.shapes._spTree.append(el)


def fill_orgchart_sample(ctx, slide, slot_def, tree, where):
    # resolve geometry before removals: dropping the prototype can
    # collapse an ancestor group and detach the region element
    region_el = find_shape_el(slide, slot_def["target"], where)
    bbox = _region_bbox(region_el, where)
    proto_el = _prototype_element(ctx, slide, slot_def, where, "org")
    kind = _target_kind(region_el)
    if kind not in ("shape", "group", "smartart"):
        ctx.fb(_el_name(region_el) or slot_def["target"],
               f"region の対象が{_KIND_JA.get(kind, '図形')}",
               "描画範囲としてのみ使い、枠は除去", "theme")
    _remove_region(slide, region_el)
    draw_orgchart(slide, tree, bbox, proto_el,
                  slot_def.get("connector_color", "accent1"), where)


def fill_timeline_sample(ctx, slide, slot_def, steps, where):
    region_el = find_shape_el(slide, slot_def["target"], where)
    bbox = _region_bbox(region_el, where)
    proto_el = _prototype_element(ctx, slide, slot_def, where, "timeline")
    kind = _target_kind(region_el)
    if kind not in ("shape", "group", "smartart"):
        ctx.fb(_el_name(region_el) or slot_def["target"],
               f"region の対象が{_KIND_JA.get(kind, '図形')}",
               "描画範囲としてのみ使い、枠は除去", "theme")
    _remove_region(slide, region_el)
    draw_timeline(slide, steps, bbox, proto_el, where)


def fill_orgchart_layout(ctx, slide, slot_def, tree, ph, bbox, where):
    proto_el = _prototype_element(ctx, slide, slot_def, where, "org",
                                  layout_mode=True)
    ph._element.getparent().remove(ph._element)
    draw_orgchart(slide, tree, bbox, proto_el,
                  slot_def.get("connector_color", "accent1"), where)


def fill_timeline_layout(ctx, slide, slot_def, steps, ph, bbox, where):
    proto_el = _prototype_element(ctx, slide, slot_def, where, "timeline",
                                  layout_mode=True)
    ph._element.getparent().remove(ph._element)
    draw_timeline(slide, steps, bbox, proto_el, where)


# ---------- chart ----------

def _chart_data(value):
    cd = CategoryChartData()
    cd.categories = value["categories"]
    fmt = value.get("number_format")
    for s in value["series"]:
        if fmt:
            cd.add_series(s["name"], s["values"], number_format=fmt)
        else:
            cd.add_series(s["name"], s["values"])
    return cd


def _chart_of_frame(slide, frame_el):
    """Chart object of a graphicFrame (works even nested in a group)."""
    for e, attr, rid in _rel_ids(frame_el):
        rel = slide.part.rels.get(rid)
        if rel is not None and not rel.is_external \
                and rel.reltype.endswith("/chart"):
            return rel.target_part.chart
    return None


def _series_fills(chart):
    """solidFill elements of each series, in order (None when absent)."""
    fills = []
    try:
        for ser in chart.plots[0].series:
            spPr = ser._element.find(qn("c:spPr"))
            f = spPr.find(qn("a:solidFill")) if spPr is not None else None
            fills.append(copy.deepcopy(f) if f is not None else None)
    except Exception:
        return []
    return fills


def _apply_series_fills(chart, fills):
    for ser, f in zip(chart.plots[0].series, fills):
        if f is None:
            continue
        ser_el = ser._element
        spPr = ser_el.find(qn("c:spPr"))
        if spPr is None:
            spPr = etree.Element(qn("c:spPr"))
            ser_el.append(spPr)
        for old in spPr.findall(qn("a:solidFill")):
            spPr.remove(old)
        spPr.insert(0, copy.deepcopy(f))


def _finish_chart(chart, value):
    """Common post-processing for newly added charts."""
    try:
        plot = chart.plots[0]
        plot.has_data_labels = True
        if value.get("number_format"):
            dl = plot.data_labels
            dl.number_format = value["number_format"]
            dl.number_format_is_linked = False
    except Exception:
        pass  # some chart types reject data labels; keep the chart
    chart.has_legend = len(value["series"]) > 1


def _replace_chart(ctx, slide, el, value, where, reason, fills=None):
    """Remove a non-conforming chart target and draw a fresh chart of the
    deck-requested type at the same absolute bbox."""
    if fills and not any(fills):
        fills = None
    bbox = absolute_bbox(el)
    if bbox is None:
        fail(f"{where}: chart slot target '{slot_target_name(el)}' "
             f"has no geometry")
    _remove_region(slide, el)
    frame = slide.shapes.add_chart(
        CHART_TYPES[value["type"]], *bbox, _chart_data(value))
    _mark_el(frame._element, FALLBACK_MARK)
    chart = frame.chart
    if fills:
        _apply_series_fills(chart, fills)
    _finish_chart(chart, value)
    src = "template chart series fills" if fills else "theme"
    ctx.fb(_el_name(el) or "?", reason,
           "同じ位置に新しいグラフを作成（テンプレの系列色を引き継ぐ）"
           if fills else "同じ位置に新しいグラフを作成", src)
    return frame


def slot_target_name(el):
    return _el_name(el) or "?"


def fill_chart_sample(ctx, slide, slot_def, value, where):
    el = find_shape_el(slide, slot_def["target"], where)
    kind = _target_kind(el)
    if kind != "chart":
        # target exists but is not a chart (picture of a chart, OLE, ...)
        if "type" not in value:
            fail(f"{where}: chart slot target '{slot_def['target']}' "
                 f"is not a chart and the deck gives no 'type' to "
                 f"build one from")
        _replace_chart(
            ctx, slide, el, value, where,
            f"対象が{_KIND_JA.get(kind, '図形')}のためグラフを差し替え"
            f"られない")
        return
    chart = _chart_of_frame(slide, el)
    if chart is None:
        fail(f"{where}: cannot access chart in "
             f"'{slot_def['target']}'")
    if "type" in value:
        expect = CHART_TYPES[value["type"]]
        actual = chart.chart_type
        if actual != expect:
            fills = _series_fills(chart)
            _replace_chart(
                ctx, slide, el, value, where,
                f"要求されたグラフ種 '{value['type']}' がテンプレの"
                f"グラフ（{actual}）と異なる", fills)
            return
    chart.replace_data(_chart_data(value))


def fill_chart_layout(ctx, slide, slot_def, value, ph, bbox, where):
    if "type" not in value:
        fail(f"{where}: layout-mode chart requires 'type' "
             f"(one of {sorted(CHART_TYPES)})")
    cd = _chart_data(value)
    if hasattr(ph, "insert_chart"):
        frame = ph.insert_chart(CHART_TYPES[value["type"]], cd)
        chart = frame.chart
    else:
        frame = slide.shapes.add_chart(
            CHART_TYPES[value["type"]], *bbox, cd)
        chart = frame.chart
        ph._element.getparent().remove(ph._element)
    _finish_chart(chart, value)


# ---------- driver ----------

def _fill_layout_slots(ctx, slide, comp, slots, where):
    for name, value in slots.items():
        slot_def = comp["slots"][name]
        t = slot_def["type"]
        ctx.slot = name
        n0 = len(ctx.fallbacks)
        if not isinstance(slot_def["target"], int):
            fail(f"{where}, slot '{name}': layout-mode target must "
                 f"be a placeholder idx (int)")
        try:
            ph = slide.placeholders[slot_def["target"]]
        except KeyError:
            fail(f"{where}, slot '{name}': placeholder idx "
                 f"{slot_def['target']} not on slide")
        if t == "text":
            fill_text(ph._element, value)
        elif t == "list":
            fill_list(ph._element, value)
        else:
            bbox = _ph_bbox(slide, ph, f"{where}, slot '{name}'")
            if t == "table":
                fill_table_layout(ctx, slide, slot_def, value, ph,
                                  bbox, f"{where}, slot '{name}'")
            elif t == "orgchart":
                fill_orgchart_layout(ctx, slide, slot_def, value, ph,
                                     bbox, f"{where}, slot '{name}'")
            elif t == "timeline":
                fill_timeline_layout(ctx, slide, slot_def, value, ph,
                                     bbox, f"{where}, slot '{name}'")
            elif t == "chart":
                fill_chart_layout(ctx, slide, slot_def, value, ph, bbox,
                                  f"{where}, slot '{name}'")
        if len(ctx.fallbacks) == n0:
            ctx.native += 1


def _fill_sample_slots(ctx, slide, comp, slots, where):
    for name, value in slots.items():
        slot_def = comp["slots"][name]
        t = slot_def["type"]
        ctx.slot = name
        n0 = len(ctx.fallbacks)
        if t == "text":
            fill_text(find_shape_el(slide, slot_def["target"], where),
                      value)
        elif t == "list":
            fill_list(find_shape_el(slide, slot_def["target"], where),
                      value)
        elif t == "table":
            fill_table(ctx, slide, slot_def, value, where)
        elif t == "orgchart":
            fill_orgchart_sample(ctx, slide, slot_def, value, where)
        elif t == "timeline":
            fill_timeline_sample(ctx, slide, slot_def, value, where)
        elif t == "chart":
            fill_chart_sample(ctx, slide, slot_def, value, where)
        if len(ctx.fallbacks) == n0:
            ctx.native += 1


def build(deck_path, map_path, template_path, out_path, strict=False,
          report_path=None):
    deck = load_json(deck_path)
    tmap = load_json(map_path)
    validate_deck(deck, tmap)
    if template_path is None:
        template_path = Path(map_path).resolve().parent / tmap["template"]
    prs = Presentation(str(template_path))
    n_original = len(prs.slides)
    ctx = BuildCtx(prs, n_original)
    for i, spec in enumerate(deck["slides"], start=1):
        comp = tmap["components"][spec["component"]]
        where = f"slide {i} (component '{spec['component']}')"
        ctx.slide_i, ctx.component = i, spec["component"]
        src = comp["source"]
        if "layout" in src:
            slide = prs.slides.add_slide(find_layout(prs, src["layout"]))
        else:
            idx = src["sample_slide"]
            if idx < 1 or idx > n_original:
                fail(f"{where}: sample_slide {idx} out of range "
                     f"(template has {n_original} slides)")
            slide = clone_slide(prs, prs.slides[idx - 1])
        slots = effective_slots(spec, comp, i)
        if "layout" in src:
            _fill_layout_slots(ctx, slide, comp, slots, where)
            remove_unfilled_placeholders(slide)
            inherit_placeholder_geometry(slide)
        else:
            _fill_sample_slots(ctx, slide, comp, slots, where)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        renumber_shape_ids(slide)
    report = {"fallbacks": ctx.fallbacks, "native": ctx.native}
    rp = report_path or (str(out_path) + ".build-report.json")
    Path(rp).write_text(json.dumps(report, ensure_ascii=False, indent=2),
                        encoding="utf-8")
    if strict and ctx.fallbacks:
        lines = "; ".join(
            f"slide {f['slide']} slot '{f['slot']}': "
            f"{f['reason']} -> {f['substitute']}"
            for f in ctx.fallbacks)
        fail(f"strict mode: {len(ctx.fallbacks)} fallback(s) occurred: "
             f"{lines}")
    # drop original template slides
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst)[:n_original]:
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)
    prs.save(str(out_path))
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", required=True)
    ap.add_argument("--map", required=True, dest="map_path")
    ap.add_argument("--template", default=None)
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--strict", action="store_true",
                    help="fail (exit 2) if any fallback was needed")
    ap.add_argument("--report", dest="report_path", default=None,
                    help="build report path (default: OUT.build-report.json)")
    args = ap.parse_args(argv)
    try:
        out = build(args.deck, args.map_path, args.template, args.output,
                    strict=args.strict, report_path=args.report_path)
    except BuildError as e:
        print(f"build error: {e}", file=sys.stderr)
        return 2
    report = json.loads(
        (Path(args.report_path) if args.report_path
         else Path(str(out) + ".build-report.json")).read_text(
            encoding="utf-8"))
    n_fb = len(report["fallbacks"])
    if n_fb:
        print(f"built {out} : fallbacks: {n_fb}")
        for f in report["fallbacks"]:
            print(f"  fallback slide {f['slide']} slot '{f['slot']}': "
                  f"{f['reason']} -> {f['substitute']} "
                  f"[{f['style_source']}]")
    else:
        print(f"built {out} : fallbacks: 0")
    return 0


if __name__ == "__main__":
    sys.exit(main())
