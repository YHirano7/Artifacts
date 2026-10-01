#!/usr/bin/env python3
"""Inventory a .pptx template: slide size, layouts, placeholders, slides,
shapes, tables, speaker notes. Optionally render slide thumbnails.

Usage: inventory.py TEMPLATE.pptx [--json OUT] [--thumbs DIR]
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

EMU_PER_IN = 914400


def _in(v):
    return round(v / EMU_PER_IN, 3) if v is not None else None


def bbox(shape):
    return [_in(shape.left), _in(shape.top), _in(shape.width),
            _in(shape.height)]


def find_soffice():
    env = os.environ.get("SOFFICE")
    if env and Path(env).exists():
        return env
    on_path = shutil.which("soffice")
    if on_path:
        return on_path
    win_default = r"C:\Program Files\LibreOffice\program\soffice.exe"
    if Path(win_default).exists():
        return win_default
    raise FileNotFoundError(
        "soffice not found. Set SOFFICE env var, put soffice on PATH, "
        "or install LibreOffice.")


def pptx_to_pngs(pptx_path, out_dir, width_px=960):
    """Render pptx -> pdf (soffice) -> PNG per slide (pypdfium2)."""
    import pypdfium2 as pdfium

    pptx_path = Path(pptx_path)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        proc = subprocess.run(
            [find_soffice(), "--headless", "--norestore",
             "--convert-to", "pdf", "--outdir", td, str(pptx_path)],
            capture_output=True, text=True, timeout=300)
        pdf_path = Path(td) / (pptx_path.stem + ".pdf")
        if proc.returncode != 0 or not pdf_path.exists():
            raise RuntimeError(
                f"soffice PDF conversion failed: {proc.stderr.strip() or proc.stdout.strip()}")
        doc = pdfium.PdfDocument(str(pdf_path))
        paths = []
        for i in range(len(doc)):
            page = doc[i]
            scale = width_px / page.get_width()
            img = page.render(scale=scale).to_pil()
            p = out_dir / f"{pptx_path.stem}-{i + 1:02d}.png"
            img.save(str(p))
            paths.append(str(p))
        doc.close()
        return paths


GRAPHIC_URIS = {
    "http://schemas.openxmlformats.org/drawingml/2006/chart": "chart",
    "http://schemas.openxmlformats.org/drawingml/2006/diagram": "smartart",
    "http://schemas.openxmlformats.org/drawingml/2006/table": "table",
}


def _graphic_kind(shape):
    """Kind of a graphicFrame's content ('chart'/'smartart'/'table')."""
    el = shape._element
    if el.tag != qn("p:graphicFrame"):
        return None
    gd = el.find(f"{qn('a:graphic')}/{qn('a:graphicData')}")
    if gd is None:
        return None
    return GRAPHIC_URIS.get(gd.get("uri"), gd.get("uri"))


def _shape_info(shape):
    info = {"name": shape.name, "type": str(shape.shape_type),
            "bbox": bbox(shape)}
    if shape.is_placeholder:
        info["placeholder_idx"] = shape.placeholder_format.idx
    if shape.shape_type == 6:  # group
        info["children"] = [_shape_info(s) for s in shape.shapes]
        return info
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            info["text"] = text[:80]
    kind = _graphic_kind(shape)
    if kind:
        info["graphic"] = kind
    if getattr(shape, "has_table", False) and shape.has_table:
        info["table"] = {"rows": len(shape.table.rows),
                         "cols": len(shape.table.columns)}
    return info


def collect(template_path):
    prs = Presentation(str(template_path))
    data = {
        "file": str(template_path),
        "slide_size": [_in(prs.slide_width), _in(prs.slide_height)],
        "layouts": [],
        "slides": [],
    }
    for i, layout in enumerate(prs.slide_layouts):
        phs = []
        for ph in layout.placeholders:
            phs.append({
                "idx": ph.placeholder_format.idx,
                "type": str(ph.placeholder_format.type),
                "name": ph.name,
                "bbox": bbox(ph),
            })
        data["layouts"].append({"index": i, "name": layout.name,
                                "placeholders": phs})
    for i, slide in enumerate(prs.slides, start=1):
        entry = {"index": i, "layout": slide.slide_layout.name,
                 "shapes": [_shape_info(sh) for sh in slide.shapes]}
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                entry["notes"] = notes
        data["slides"].append(entry)
    return data


def print_report(data):
    print(f"template : {data['file']}")
    print(f"size     : {data['slide_size'][0]} x {data['slide_size'][1]} in")
    print("\n== slide layouts ==")
    for layout in data["layouts"]:
        print(f"[{layout['index']}] {layout['name']}")
        for ph in layout["placeholders"]:
            print(f"    ph idx={ph['idx']} type={ph['type']} "
                  f"name={ph['name']!r} bbox={ph['bbox']}")
    print("\n== slides ==")
    for slide in data["slides"]:
        print(f"slide {slide['index']} (layout: {slide['layout']})")
        for sh in slide["shapes"]:
            _print_shape(sh, "    ")
        if "notes" in slide:
            print(f"    notes: {slide['notes']}")


def _print_shape(sh, indent):
    extra = ""
    if "placeholder_idx" in sh:
        extra += f" ph_idx={sh['placeholder_idx']}"
    if "graphic" in sh:
        extra += f" graphic={sh['graphic']}"
    if "table" in sh:
        extra += f" table={sh['table']['rows']}x{sh['table']['cols']}"
    text = f" text={sh['text']!r}" if "text" in sh else ""
    print(f"{indent}{sh['name']!r} {sh['type']} bbox={sh['bbox']}"
          f"{extra}{text}")
    for child in sh.get("children", []):
        _print_shape(child, indent + "  ")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("--json", dest="json_out")
    ap.add_argument("--thumbs", dest="thumbs_dir")
    args = ap.parse_args(argv)

    data = collect(Path(args.template))
    print_report(data)
    if args.json_out:
        Path(args.json_out).write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nwrote {args.json_out}")
    if args.thumbs_dir:
        paths = pptx_to_pngs(args.template, args.thumbs_dir)
        print(f"\nwrote {len(paths)} thumbnails to {args.thumbs_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
