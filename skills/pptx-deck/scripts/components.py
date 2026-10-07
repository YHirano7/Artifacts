from dataclasses import dataclass
import math
import sys

import jsonschema
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from common import (CHART_TYPES, GENERATED_MARK, LINE_HEIGHT,
                    char_units, chart_data, fail, finish_chart, mark_el,
                    read_theme)

EMU_PER_IN = 914400
ROLE_DEFAULTS = {
    "primary": "accent1", "highlight": "accent2", "text": "tx1",
    "surface": "bg2", "muted": "text", "line": "bg1",
    "background": "bg1",
}
DEFAULT_SIZES = {"heading": 20, "body": 16, "caption": 13,
                 "number": 48, "min": 11}


@dataclass(frozen=True)
class ColorRef:
    scheme: str = None
    rgb: str = None
    transforms: tuple = ()


@dataclass
class Style:
    colors: dict
    fonts: dict
    sizes: dict
    theme_colors: dict


@dataclass
class Component:
    variants: tuple
    slots_schema: dict
    draw: object
    check: object = None


def _luminance(color, theme_colors):
    rgb = color.rgb or theme_colors.get(color.scheme, "000000")
    channels = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    for kind, amount in color.transforms:
        if kind == "lumMod":
            channels = [v * amount for v in channels]
        elif kind == "lumOff":
            channels = [v + amount for v in channels]

    def linear(v):
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4

    r, g, b = (linear(min(max(v, 0), 1)) for v in channels)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _contrast(a, b, theme_colors):
    hi, lo = sorted((_luminance(a, theme_colors),
                     _luminance(b, theme_colors)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def _scheme_ref(key, clr_map):
    return ColorRef(scheme=clr_map.get(key, key))


def resolve_style(prs, tmap):
    master = prs.slide_masters[0]
    theme_data = read_theme(prs)
    theme_colors = {
        key: value.lstrip("#") for key, value in theme_data["colors"].items()
    }
    clr_map_el = master.element.find(qn("p:clrMap"))
    clr_map = {
        "tx1": "dk1", "tx2": "dk2", "bg1": "lt1", "bg2": "lt2"
    }
    if clr_map_el is not None:
        clr_map.update(clr_map_el.attrib)
    palette = tmap.get("style", {}).get("palette", {})

    def color_for(role, fallback=None):
        value = palette.get(role, fallback or ROLE_DEFAULTS[role])
        if value.startswith("#"):
            return ColorRef(rgb=value[1:].upper())
        return _scheme_ref(value, clr_map)

    colors = {role: color_for(role) for role in ROLE_DEFAULTS
              if role not in ("surface", "muted")}
    if "line" not in palette:
        colors["line"] = ColorRef(
            colors["line"].scheme, transforms=(("lumMod", 0.85),))
    colors["surface"] = color_for("surface")
    if "muted" in palette:
        colors["muted"] = color_for("muted")
    else:
        colors["muted"] = ColorRef(
            colors["text"].scheme, colors["text"].rgb,
            colors["text"].transforms + (("lumMod", 0.65),
                                         ("lumOff", 0.35)))
    dark, light = _scheme_ref("dk1", clr_map), _scheme_ref("lt1", clr_map)
    colors["on_primary"] = max(
        (dark, light), key=lambda c: _contrast(
            colors["primary"], c, theme_colors))
    colors["on_highlight"] = max(
        (dark, light), key=lambda c: _contrast(
            colors["highlight"], c, theme_colors))

    theme_fonts = theme_data["fonts"]
    overrides = tmap.get("style", {}).get("fonts", {})
    fonts = {
        "heading_latin": overrides.get("latin")
        or theme_fonts["major"].get("latin", ""),
        "heading_ea": overrides.get("ea")
        or theme_fonts["major"].get("ea", ""),
        "body_latin": overrides.get("latin")
        or theme_fonts["minor"].get("latin", ""),
        "body_ea": overrides.get("ea")
        or theme_fonts["minor"].get("ea", ""),
    }
    fonts = {key: value or "Yu Gothic" for key, value in fonts.items()}
    sizes = dict(DEFAULT_SIZES)
    sizes.update(tmap.get("style", {}).get("sizes", {}))
    if sizes["min"] > min(sizes[k] for k in ("heading", "body", "caption")):
        fail("style.sizes.min cannot exceed heading, body, or caption")
    return Style(colors, fonts, sizes, theme_colors)


def _write_color(parent, color, tag="a:solidFill"):
    old = parent.find(qn(tag))
    index = parent.index(old) if old is not None else None
    if old is not None:
        parent.remove(old)
    fill = OxmlElement(tag)
    if color.rgb:
        node = OxmlElement("a:srgbClr")
        node.set("val", color.rgb)
    else:
        node = OxmlElement("a:schemeClr")
        node.set("val", color.scheme)
    fill.append(node)
    for name, value in color.transforms:
        transform = OxmlElement(f"a:{name}")
        transform.set("val", str(round(value * 100000)))
        node.append(transform)
    if index is None:
        before = ("a:ln", "a:effectLst", "a:latin", "a:ea", "a:cs")
        index = next((i for i, child in enumerate(parent)
                      if child.tag in {qn(tag) for tag in before}),
                     len(parent))
    parent.insert(index, fill)


def _prepare_shape(shape, w, h, radius=False):
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    if not radius:
        return
    geom = shape._element.spPr.find(qn("a:prstGeom"))
    if geom is None or geom.get("prst") != "roundRect":
        return
    av_lst = geom.find(qn("a:avLst"))
    if av_lst is None:
        av_lst = OxmlElement("a:avLst")
        geom.insert(0, av_lst)
    for child in list(av_lst):
        if child.get("name") == "adj":
            av_lst.remove(child)
    gd = OxmlElement("a:gd")
    gd.set("name", "adj")
    gd.set("fmla", f"val {max(1, min(100000, round(
        0.06 / max(min(w, h), 0.001) * 100000)))}")
    av_lst.append(gd)


def _text_color(run, color):
    rpr = run._r.get_or_add_rPr()
    _write_color(rpr, color)


def _set_fonts(run, fonts, heading=False):
    latin = fonts["heading_latin" if heading else "body_latin"]
    ea = fonts["heading_ea" if heading else "body_ea"]
    run.font.name = latin
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", latin), ("a:ea", ea)):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set("typeface", face)


def _box(shape, fill, line=None):
    _prepare_shape(shape, shape.width / EMU_PER_IN, shape.height / EMU_PER_IN,
                   shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    _write_color(shape._element.spPr, fill)
    shape.line.width = Pt(0.75)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    line_el = shape._element.spPr.find(qn("a:ln"))
    if line is None:
        shape.line.fill.background()
        if line_el is not None:
            shape._element.spPr.remove(line_el)
    else:
        _write_color(line_el, line)


def _estimated_height(text, width, size, pad_top=0.08, pad_bottom=0.08):
    is_bulleted = isinstance(text, list)
    inner_width = max(width - 0.24 - (0.22 if is_bulleted else 0), 0.05) \
        * EMU_PER_IN
    line_height = size * LINE_HEIGHT / 72 * EMU_PER_IN
    lines = 0
    paragraphs = text if is_bulleted else str(text).splitlines()
    for paragraph in paragraphs or [""]:
        units = sum(char_units(char) for char in paragraph)
        lines += max(1, math.ceil(
            units * size / 72 * EMU_PER_IN / inner_width - 1e-9))
    after = (max(1, len(paragraphs)) * 4 / 72 * EMU_PER_IN
             if is_bulleted else 0)
    return lines * line_height + after + (
        pad_top + pad_bottom) * EMU_PER_IN


def fit_size(text, width, height, base, min_size, where):
    for size in range(int(base), int(min_size) - 1, -1):
        if _estimated_height(text, width, size) <= height * EMU_PER_IN:
            return size
    needed = _estimated_height(text, width, min_size) / EMU_PER_IN
    chars = max(1, int((width - 0.24) * height * 72 /
                       max(min_size * LINE_HEIGHT, 1)))
    fail(f"{where}: text does not fit even at {min_size}pt "
         f"(needs ~{needed:.1f}in, box {height:.1f}in); "
         f"shorten to about {chars} chars")


def fit_siblings(texts, width, height, base, min_size, where):
    sizes = [fit_size(text, width, height, base, min_size,
                      f"{where}[{i}]") for i, text in enumerate(texts, 1)]
    return min(sizes) if sizes else base


class DrawContext:
    def __init__(self, slide, style, where):
        self.slide = slide
        self.style = style
        self.where = where
        self.frame = {}

    def shape(self, name, x, y, w, h, kind=MSO_SHAPE.RECTANGLE,
              fill=None, line=None, text=None, size=None, color=None,
              bold=False, align=PP_ALIGN.LEFT, heading=False,
              radius=False, anchor=MSO_ANCHOR.TOP):
        if radius:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE
        shape = self.slide.shapes.add_shape(
            kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        _prepare_shape(shape, w, h, radius)
        if fill:
            _box(shape, fill, line)
        else:
            shape.fill.background()
            if line is None:
                shape.line.fill.background()
            else:
                shape.line.width = Pt(0.75)
                shape.line.color.rgb = RGBColor(0, 0, 0)
                _write_color(shape._element.spPr.find(qn("a:ln")), line)
        if text is not None:
            self.text_frame(shape, text, size, color, bold, align, heading,
                            name, anchor)
        mark_el(shape._element, GENERATED_MARK)
        return shape

    def textbox(self, name, x, y, w, h, text, size, color=None, bold=False,
                align=PP_ALIGN.LEFT, heading=False,
                anchor=MSO_ANCHOR.TOP, margin=0.12):
        shape = self.slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        shape.fill.background()
        shape.line.fill.background()
        shape.text_frame.word_wrap = True
        shape.text_frame.vertical_anchor = anchor
        shape.text_frame.margin_left = Inches(margin)
        shape.text_frame.margin_right = Inches(margin)
        shape.text_frame.margin_top = Inches(0.08)
        shape.text_frame.margin_bottom = Inches(0.08)
        self.text_frame(shape, text, size, color, bold, align, heading, name,
                        anchor)
        mark_el(shape._element, GENERATED_MARK)
        return shape

    def rich_textbox(self, name, x, y, w, h, runs, align=PP_ALIGN.LEFT,
                     anchor=MSO_ANCHOR.TOP):
        shape = self.slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        shape.fill.background()
        shape.line.fill.background()
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = Inches(0.12)
        frame.margin_right = Inches(0.12)
        frame.margin_top = Inches(0.08)
        frame.margin_bottom = Inches(0.08)
        para = frame.paragraphs[0]
        para.alignment = align
        para.line_spacing = 1.0
        para.space_after = Pt(0)
        for spec in runs:
            run = para.add_run()
            run.text = spec["text"]
            run.font.size = Pt(spec["size"])
            run.font.bold = spec.get("bold", False)
            _set_fonts(run, self.style.fonts, spec.get("heading", False))
            _text_color(run, spec.get("color", self.style.colors["text"]))
        mark_el(shape._element, GENERATED_MARK)
        return shape

    def text_frame(self, shape, text, size, color, bold, align, heading, name,
                   anchor=MSO_ANCHOR.TOP):
        size = fit_size(text, shape.width / EMU_PER_IN,
                        shape.height / EMU_PER_IN,
                        size or self.style.sizes["body"],
                        self.style.sizes["min"],
                        f"{self.where}, {name}")
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = Inches(0.12)
        frame.margin_right = Inches(0.12)
        frame.margin_top = Inches(0.08)
        frame.margin_bottom = Inches(0.08)
        is_bulleted = isinstance(text, list)
        lines = text if is_bulleted else str(text).splitlines()
        for index, line in enumerate(lines or [""]):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.alignment = align
            para.line_spacing = 1.0
            para.space_after = Pt(4 if is_bulleted else 0)
            if is_bulleted:
                ppr = para._p.get_or_add_pPr()
                ppr.set("marL", str(Inches(0.22)))
                ppr.set("indent", str(-Inches(0.12)))
                bu_char = OxmlElement("a:buChar")
                bu_char.set("char", "•")
                ppr.append(bu_char)
            run = para.add_run()
            run.text = line
            run.font.size = Pt(size or self.style.sizes["body"])
            run.font.bold = bold
            _set_fonts(run, self.style.fonts, heading)
            _text_color(run, color or self.style.colors["text"])

    def group(self, elements, name):
        _group_elements(self.slide, elements, name)


def _group_elements(slide, elements, name):
    if not elements:
        return
    elements = [el._element if hasattr(el, "_element") else el
                for el in elements]
    tree = slide.shapes._spTree
    entries = []
    for el in elements:
        xfrm = el.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
        if xfrm is None:
            xfrm = el.find(f"{qn('p:xfrm')}")
        if xfrm is None:
            continue
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        if off is not None and ext is not None:
            entries.append((el, xfrm, int(off.get("x")), int(off.get("y")),
                            int(ext.get("cx")), int(ext.get("cy"))))
    if not entries:
        return
    left = min(e[2] for e in entries)
    top = min(e[3] for e in entries)
    right = max(e[2] + e[4] for e in entries)
    bottom = max(e[3] + e[5] for e in entries)
    width, height = right - left, bottom - top
    ids = [int(el.get("id")) for el in tree.iter(qn("p:cNvPr"))
           if el.get("id", "").isdigit()]
    group = OxmlElement("p:grpSp")
    nv = OxmlElement("p:nvGrpSpPr")
    c_nv = OxmlElement("p:cNvPr")
    c_nv.set("id", str(max(ids, default=0) + 1))
    c_nv.set("name", name)
    c_nv.set("descr", GENERATED_MARK)
    nv.append(c_nv)
    nv.append(OxmlElement("p:cNvGrpSpPr"))
    nv.append(OxmlElement("p:nvPr"))
    group.append(nv)
    grp_props = OxmlElement("p:grpSpPr")
    xfrm = OxmlElement("a:xfrm")
    off = OxmlElement("a:off")
    off.set("x", str(left))
    off.set("y", str(top))
    ext = OxmlElement("a:ext")
    ext.set("cx", str(width))
    ext.set("cy", str(height))
    ch_off = OxmlElement("a:chOff")
    ch_off.set("x", "0")
    ch_off.set("y", "0")
    ch_ext = OxmlElement("a:chExt")
    ch_ext.set("cx", str(width))
    ch_ext.set("cy", str(height))
    for child in (off, ext, ch_off, ch_ext):
        xfrm.append(child)
    grp_props.append(xfrm)
    group.append(grp_props)
    first = min(tree.index(e[0]) for e in entries)
    for el, child_xfrm, x, y, _, _ in entries:
        child_xfrm.find(qn("a:off")).set("x", str(x - left))
        child_xfrm.find(qn("a:off")).set("y", str(y - top))
        mark_el(el, GENERATED_MARK)
        tree.remove(el)
        group.append(el)
    tree.insert(first, group)


def _component_schema(properties, required):
    props = {
        "lead": {"type": "string"},
        "takeaway": {"type": "string"},
    }
    props.update(properties)
    return {
        "type": "object", "required": list(required),
        "properties": props, "additionalProperties": False,
    }


def _string_or_list():
    return {"oneOf": [
        {"type": "string"},
        {"type": "array", "items": {"type": "string"}},
    ]}


def _items_schema(properties, required, minimum, maximum):
    return {
        "type": "array", "minItems": minimum, "maxItems": maximum,
        "items": {
            "type": "object", "properties": properties,
            "required": required, "additionalProperties": False,
        },
    }


def _takeaway_height(dc, text, width):
    return max(0.5, _estimated_height(
        text, width - 0.28, dc.style.sizes["body"]) / EMU_PER_IN + 0.04)


def _draw_frame(dc, region, slots):
    x, y, w, h = region
    lead_h = 0.0
    if slots.get("lead"):
        lead_h = max(0.42, _estimated_height(
            slots["lead"], w, dc.style.sizes["heading"]) / EMU_PER_IN)
        dc.textbox("Component lead", x, y, w, lead_h, slots["lead"],
                   dc.style.sizes["heading"], dc.style.colors["text"],
                   bold=True, heading=True)
    top = y + lead_h + (0.12 if lead_h else 0)
    dc.frame = {
        "x": x, "y": y, "w": w, "h": h,
        "content_top": top, "content_bottom": y + h,
        "takeaway": slots.get("takeaway"),
    }
    available_h = max(0.2, y + h - top)
    if slots.get("takeaway"):
        available_h = max(
            0.2, available_h - 0.25 - _takeaway_height(dc, slots["takeaway"], w))
    return x, top, w, available_h


def _finish_frame(dc, bottom):
    text = dc.frame.get("takeaway")
    if not text:
        return
    x, y, w, h = (dc.frame[key] for key in ("x", "y", "w", "h"))
    ty = bottom + 0.25
    th = _takeaway_height(dc, text, w)
    if ty + th > y + h:
        th = max(0.5, y + h - ty)
    dc.shape("Component takeaway", x, ty, w, th,
             fill=dc.style.colors["surface"], line=dc.style.colors["line"],
             radius=True)
    dc.shape("Takeaway accent", x, ty, 0.09, th,
             fill=dc.style.colors["primary"], line=None)
    dc.textbox("Takeaway text", x + 0.16, ty, w - 0.22, th, text,
               dc.style.sizes["body"], dc.style.colors["text"], bold=True)


def _shape_c_nvpr(element):
    tags = {
        qn("p:sp"): "p:nvSpPr",
        qn("p:grpSp"): "p:nvGrpSpPr",
        qn("p:cxnSp"): "p:nvCxnSpPr",
        qn("p:graphicFrame"): "p:nvGraphicFramePr",
    }
    tag = tags.get(element.tag)
    return element.find(f"{qn(tag)}/{qn('p:cNvPr')}") if tag else None


def _shape_xfrm(element):
    if element.tag in (qn("p:sp"), qn("p:cxnSp")):
        return element.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
    if element.tag == qn("p:grpSp"):
        return element.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    if element.tag == qn("p:graphicFrame"):
        return element.find(qn("p:xfrm"))
    return None


def _center_component(dc):
    frame = dc.frame
    bounds = []
    elements = []
    for element in dc.slide._element.spTree:
        c_nvpr = _shape_c_nvpr(element)
        if (c_nvpr is None or c_nvpr.get("descr") != GENERATED_MARK
                or c_nvpr.get("name") == "Component lead"):
            continue
        xfrm = _shape_xfrm(element)
        if xfrm is None:
            continue
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        if off is None or ext is None:
            continue
        top = int(off.get("y"))
        bottom = top + int(ext.get("cy"))
        bounds.append((top, bottom))
        elements.append(off)
    if not bounds:
        return
    top = min(bound[0] for bound in bounds)
    bottom = max(bound[1] for bound in bounds)
    content_top = round(frame["content_top"] * EMU_PER_IN)
    content_bottom = round(frame["content_bottom"] * EMU_PER_IN)
    offset = round((content_top + content_bottom - top - bottom) / 2)
    if offset > 0:
        for off in elements:
            off.set("y", str(int(off.get("y")) + offset))


def _text_height(text, width, size, padding=0.22, minimum=0.9):
    return max(minimum, _estimated_height(text, width, size) / EMU_PER_IN
               + padding)


def _group(dc, els, name, index):
    dc.group(els, f"{name} {index}")


def _panel(dc, prefix, x, y, w, heading, body, fill, line,
           heading_fill=None, heading_color=None, body_color=None,
           min_height=0.9, radius=True, heading_size=None, body_size=None):
    heading_color = heading_color or dc.style.colors["text"]
    body_color = body_color or dc.style.colors["text"]
    heading_size = heading_size or dc.style.sizes["heading"]
    body_size = body_size or dc.style.sizes["body"]
    heading_h = max(0.34, _estimated_height(
        heading, w - 0.24, heading_size) / EMU_PER_IN + 0.12) \
        if heading else 0
    body_h = max(0.3, _estimated_height(
        body, w - 0.24, body_size) / EMU_PER_IN + 0.08) if body else 0
    height = max(min_height, heading_h + body_h
                 + (0.12 if heading and body else 0.16))
    els = [dc.shape(f"{prefix} panel", x, y, w, height, fill=fill,
                    line=line, radius=radius)]
    if heading and heading_fill is not None:
        els.append(dc.shape(
            f"{prefix} heading", x, y, w, heading_h, fill=heading_fill,
            line=None, text=heading, size=heading_size,
            color=heading_color, bold=True, heading=True,
            anchor=MSO_ANCHOR.TOP, radius=radius))
    elif heading:
        els.append(dc.textbox(
            f"{prefix} heading", x + 0.12, y + 0.08, w - 0.24,
            heading_h - 0.08, heading, heading_size,
            heading_color, bold=True, heading=True))
    if body:
        els.append(dc.textbox(
            f"{prefix} body", x + 0.12, y + (heading_h or 0.08),
            w - 0.24, body_h, body, body_size, body_color))
    return els, height


def _line(dc, name, x1, y1, x2, y2, color=None, width=0.75):
    line = dc.slide.shapes.add_connector(
        1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.name = name
    line.line.width = Pt(width)
    ln = line._element.spPr.find(qn("a:ln"))
    if ln is not None:
        _write_color(ln, color or dc.style.colors["line"])
    style = line._element.find(qn("p:style"))
    if style is not None:
        line._element.remove(style)
    mark_el(line._element, GENERATED_MARK)
    return line


def _emphasis(slots, count, where):
    index = slots.get("emphasis")
    if index is not None and not 1 <= index <= count:
        fail(f"{where}, emphasis: expected a number from 1 to {count}")
    return index


def _draw_cards(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    items = slots["items"]
    rows = 2 if len(items) >= 5 else (2 if len(items) == 4 else 1)
    cols = 3 if rows == 2 and len(items) >= 5 else (
        2 if rows == 2 else len(items))
    gap = 0.18
    cw = (w - gap * (cols - 1)) / cols
    emph = _emphasis(slots, len(items), dc.where)
    body_texts = [item["body"] for item in items]
    heading_h = max(_estimated_height(
        item["heading"], cw - 0.24, dc.style.sizes["heading"]) / EMU_PER_IN
        + 0.12 for item in items)
    body_h = max(_estimated_height(
        body, cw - 0.28, dc.style.sizes["body"]) / EMU_PER_IN + 0.08
        for body in body_texts)
    ch = max(0.9, heading_h + body_h + 0.2)
    if rows * ch + gap * (rows - 1) > h:
        ch = (h - gap * (rows - 1)) / rows
    body_size = fit_siblings(
        body_texts, cw - 0.28, max(0.3, ch - heading_h - 0.2),
        dc.style.sizes["body"],
        dc.style.sizes["min"], f"{dc.where}, items[].body")
    for i, (item, body) in enumerate(zip(items, body_texts), 1):
        col, row = (i - 1) % cols, (i - 1) // cols
        cx, cy = x + col * (cw + gap), y + row * (ch + gap)
        highlight = emph == i
        line = dc.style.colors["highlight"] if highlight \
            else dc.style.colors["line"]
        heading_color = dc.style.colors["highlight"] if highlight \
            else dc.style.colors["primary"]
        els = []
        if variant == "header":
            els.append(dc.shape(f"Card {i} body", cx, cy, cw, ch,
                                fill=dc.style.colors["background"],
                                line=line, radius=True))
            els.append(dc.shape(
                f"Card {i} heading", cx, cy, cw, heading_h,
                fill=dc.style.colors["primary"], line=None,
                text=item["heading"],
                size=dc.style.sizes["heading"],
                color=dc.style.colors["on_primary"],
                bold=True, heading=True))
            els.append(dc.textbox(
                f"Card {i} copy", cx + 0.12, cy + heading_h,
                cw - 0.24, ch - heading_h - 0.08, body, body_size,
                dc.style.colors["text"]))
        elif variant == "outline":
            els, _ = _panel(
                dc, f"Card {i}", cx, cy, cw, item["heading"], body,
                dc.style.colors["background"], line,
                heading_color=heading_color,
                min_height=ch, body_size=body_size)
            els.append(dc.shape(
                f"Card {i} accent", cx + 0.06, cy + 0.06, cw - 0.12, 0.06,
                fill=dc.style.colors["primary"], line=None))
        else:
            els.append(dc.shape(f"Card {i} panel", cx, cy, cw, ch,
                                fill=dc.style.colors["surface"],
                                line=line, radius=True))
            els.append(dc.shape(
                f"Card {i} number", cx + 0.14, cy + 0.12, 0.42, 0.42,
                kind=MSO_SHAPE.OVAL,
                fill=dc.style.colors["highlight"] if highlight
                else dc.style.colors["primary"], line=None, text=str(i),
                size=dc.style.sizes["caption"],
                color=dc.style.colors["on_highlight" if highlight
                                      else "on_primary"],
                bold=True, align=PP_ALIGN.CENTER))
            els.append(dc.textbox(
                f"Card {i} heading", cx + 0.64, cy + 0.1, cw - 0.78,
                heading_h, item["heading"], dc.style.sizes["heading"],
                heading_color, bold=True, heading=True))
            els.append(dc.textbox(
                f"Card {i} copy", cx + 0.16, cy + 0.7, cw - 0.32,
                ch - 0.78, body, body_size, dc.style.colors["text"]))
        if highlight:
            els[0].line.width = Pt(2)
        _group(dc, els, "Cards", i)
    _finish_frame(dc, y + rows * ch + gap * (rows - 1))


def _draw_kpi(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    items = slots["items"]
    gap = 0.16
    cw = (w - gap * (len(items) - 1)) / len(items)
    label_h = max(_estimated_height(
        item["label"], cw - 0.24, dc.style.sizes["body"]) / EMU_PER_IN
        + 0.08 for item in items)
    note_h = max((_estimated_height(
        item.get("note", ""), cw - 0.24, dc.style.sizes["caption"])
        / EMU_PER_IN + 0.04) if item.get("note") else 0 for item in items)
    value_h = 1.08
    block_h = max(1.2, 0.06 + value_h + 0.1 + label_h
                  + (0.08 + note_h if note_h else 0.04))
    if block_h > h:
        block_h = h
    for i, item in enumerate(items, 1):
        cx = x + (i - 1) * (cw + gap)
        els = []
        value_runs = [{"text": item["value"],
                       "size": dc.style.sizes["number"],
                       "color": dc.style.colors["primary"],
                       "bold": True, "heading": True}]
        if item.get("unit"):
            value_runs.append({
                "text": f" {item['unit']}",
                "size": round(dc.style.sizes["number"] * 0.45),
                "color": dc.style.colors["muted"],
                "bold": True,
            })
        if variant == "tiles":
            els.append(dc.shape(f"KPI {i} tile", cx, y, cw, block_h,
                                fill=dc.style.colors["surface"],
                                line=dc.style.colors["line"], radius=True))
            els.append(dc.shape(
                f"KPI {i} accent", cx + 0.06, y + 0.04, cw - 0.12, 0.06,
                fill=dc.style.colors["primary"], line=None))
            value_box = dc.rich_textbox(
                f"KPI {i} value", cx + 0.02, y + 0.1, cw - 0.04, value_h,
                value_runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            value_box.text_frame.margin_left = Inches(0.02)
            value_box.text_frame.margin_right = Inches(0.02)
            els.append(value_box)
            label_y = y + 0.1 + value_h + 0.02
            els.append(dc.textbox(
                f"KPI {i} label", cx + 0.12, label_y, cw - 0.24, label_h,
                item["label"], dc.style.sizes["body"],
                dc.style.colors["text"], bold=True, align=PP_ALIGN.CENTER,
                heading=True))
            if item.get("note"):
                els.append(dc.textbox(
                    f"KPI {i} note", cx + 0.12,
                    label_y + label_h + 0.08, cw - 0.24, note_h,
                    item["note"],
                    dc.style.sizes["caption"], dc.style.colors["muted"],
                    align=PP_ALIGN.CENTER))
        else:
            if i == 1:
                els.append(dc.shape("KPI band", x, y, w, block_h,
                                    fill=dc.style.colors["surface"],
                                    line=dc.style.colors["line"], radius=True))
            if i > 1:
                els.append(dc.shape(
                    f"KPI {i} divider", cx - gap / 2, y + 0.12, 0.01,
                    block_h - 0.24, fill=dc.style.colors["line"], line=None))
            value_box = dc.rich_textbox(
                f"KPI {i} value", cx + 0.08, y + 0.08, cw - 0.16, value_h,
                value_runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            value_box.text_frame.margin_left = Inches(0.02)
            value_box.text_frame.margin_right = Inches(0.02)
            els.append(value_box)
            label_y = y + 0.08 + value_h + 0.04
            els.append(dc.textbox(
                f"KPI {i} label", cx + 0.08, label_y, cw - 0.16, label_h,
                item["label"], dc.style.sizes["body"], dc.style.colors["text"],
                bold=True,
                align=PP_ALIGN.CENTER, heading=True))
            if item.get("note"):
                els.append(dc.textbox(
                    f"KPI {i} note", cx + 0.08,
                    label_y + label_h + 0.08,
                    cw - 0.16, note_h, item["note"],
                    dc.style.sizes["caption"], dc.style.colors["muted"],
                    align=PP_ALIGN.CENTER))
        _group(dc, els, "KPI", i)
    _finish_frame(dc, y + block_h)


def _draw_process(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    steps = slots["steps"]
    n = len(steps)
    gap = 0.24 if variant == "circles" else 0.3
    cw = (w - gap * (n - 1)) / n
    emphasis = _emphasis(slots, n, dc.where)
    label_size = fit_siblings(
        [step["label"] for step in steps], cw - 0.18, 0.6,
        dc.style.sizes["body"], dc.style.sizes["min"],
        f"{dc.where}, steps[].label")
    details = [step.get("detail", "") for step in steps]
    label_h = max(_estimated_height(
        step["label"], cw - 0.18, label_size) / EMU_PER_IN + 0.04
        for step in steps)
    detail_size = fit_siblings(
        details, cw - 0.18, max(0.4, h - 1.25),
        dc.style.sizes["body"],
        dc.style.sizes["min"], f"{dc.where}, steps[].detail")
    detail_h = max((_estimated_height(
        detail, cw - 0.18, detail_size) / EMU_PER_IN + 0.04)
        if detail else 0 for detail in details)
    if variant == "circles":
        diameter = min(0.9, max(0.7, h * 0.22))
        cy = y
        label_y = cy + diameter + 0.12
        detail_y = label_y + label_h + 0.04
        for i, step in enumerate(steps, 1):
            cx = x + (i - 1) * (cw + gap) + (cw - diameter) / 2
            if i < n:
                next_cx = x + i * (cw + gap) + (cw - diameter) / 2
                _line(dc, f"Process connector {i}",
                      cx + diameter, cy + diameter / 2,
                      next_cx, cy + diameter / 2,
                      dc.style.colors["line"], 1.5)
            color = dc.style.colors["highlight"] if emphasis == i \
                else dc.style.colors["primary"]
            els = [dc.shape(
                f"Process {i} node", cx, cy, diameter, diameter,
                kind=MSO_SHAPE.OVAL, fill=color, line=None, text=str(i),
                size=dc.style.sizes["body"],
                color=dc.style.colors["on_highlight" if emphasis == i
                                      else "on_primary"],
                bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)]
            els.append(dc.textbox(
                f"Process {i} label", x + (i - 1) * (cw + gap), label_y,
                cw, label_h, step["label"], label_size,
                dc.style.colors["text"], bold=True, align=PP_ALIGN.CENTER))
            if step.get("detail"):
                els.append(dc.textbox(
                    f"Process {i} detail",
                    x + (i - 1) * (cw + gap) + 0.02, detail_y, cw - 0.04,
                    max(0.3, detail_h), step["detail"],
                    detail_size, dc.style.colors["muted"],
                    align=PP_ALIGN.CENTER))
            _group(dc, els, "Process", i)
        _finish_frame(dc, y + diameter + 0.12 + label_h + detail_h)
        return
    for i, step in enumerate(steps, 1):
        cx = x + (i - 1) * (cw + gap)
        color = dc.style.colors["highlight"] if emphasis == i \
            else dc.style.colors["primary"]
        if variant == "chevron":
            shape_kind = MSO_SHAPE.CHEVRON
            node_h = min(0.8, max(0.62, h * 0.24))
            els = [dc.shape(
                f"Process {i} chevron", cx, y, cw + 0.08, node_h,
                kind=shape_kind, fill=color, line=None,
                text=step["label"], size=label_size,
                color=dc.style.colors["on_highlight" if emphasis == i
                                      else "on_primary"],
                bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)]
            if step.get("detail"):
                els.append(dc.textbox(
                    f"Process {i} detail", cx + 0.04, y + node_h + 0.15,
                    cw - 0.08, max(0.3, detail_h), step["detail"],
                    detail_size, dc.style.colors["muted"]))
        else:
            node_h = min(
                h, max(0.9, label_h + detail_h + 0.28))
            els = [dc.shape(f"Process {i} box", cx, y, cw, node_h,
                            fill=dc.style.colors["surface"], line=color,
                            radius=True)]
            els.append(dc.textbox(
                f"Process {i} label", cx + 0.12, y + 0.08, cw - 0.24,
                label_h, step["label"], label_size,
                dc.style.colors["text"], bold=True, heading=True))
            if step.get("detail"):
                els.append(dc.textbox(
                    f"Process {i} detail", cx + 0.12,
                    y + 0.12 + label_h, cw - 0.24,
                    max(0.3, node_h - label_h - 0.2),
                    step["detail"], detail_size,
                    dc.style.colors["text"]))
        _group(dc, els, "Process", i)
        if variant == "arrows" and i < n:
            arrow = dc.shape(
                f"Process arrow {i}", cx + cw + 0.04, y + node_h / 2 - 0.12,
                0.24, 0.24, kind=MSO_SHAPE.ISOSCELES_TRIANGLE,
                fill=dc.style.colors["highlight"] if emphasis == i
                else dc.style.colors["primary"], line=None)
            arrow.rotation = 90
    _finish_frame(dc, y + max(0.8, node_h)
                  + (0.15 + detail_h if variant == "chevron" else 0))


def _draw_comparison(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    left, right = slots["left"], slots["right"]
    gap = 0.72 if variant == "before_after" else 0.52
    panel_w = (w - gap) / 2
    emphasis = _emphasis(slots, 2, dc.where)
    panels = (left, right)
    panel_heights = []
    for item in panels:
        body = item["items"]
        panel_heights.append(_text_height(
            body, panel_w - 0.34, dc.style.sizes["body"],
            padding=0.82, minimum=1.35))
    panel_h = min(h, max(panel_heights))
    heading_h = 0.5
    for i, item in enumerate(panels, 1):
        px = x + (i - 1) * (panel_w + gap)
        line = dc.style.colors["highlight"] if emphasis == i \
            else dc.style.colors["line"]
        els = [dc.shape(
            f"Comparison {i} panel", px, y, panel_w, panel_h,
            fill=dc.style.colors["surface"], line=line, radius=True)]
        if emphasis == i:
            els[0].line.width = Pt(2)
        if variant == "before_after":
            header_color = dc.style.colors["muted"] if i == 1 \
                else dc.style.colors["primary"]
            header_text = dc.style.colors["on_primary"]
        else:
            header_color = dc.style.colors["primary"]
            header_text = dc.style.colors["on_primary"]
        heading_text_color = dc.style.colors["highlight"] if emphasis == i \
            else header_text
        els.append(dc.shape(
            f"Comparison {i} heading", px, y, panel_w, heading_h,
            fill=header_color, line=None, text=item["heading"],
            size=dc.style.sizes["heading"],
            color=heading_text_color, bold=True, heading=True))
        items = item["items"]
        size = fit_size(items, panel_w - 0.34, panel_h - heading_h - 0.2,
                        dc.style.sizes["body"], dc.style.sizes["min"],
                        f"{dc.where}, {'left' if i == 1 else 'right'}.items")
        els.append(dc.textbox(
            f"Comparison {i} items", px + 0.17, y + heading_h + 0.08,
            panel_w - 0.34, panel_h - heading_h - 0.14, items, size,
            dc.style.colors["text"]))
        _group(dc, els, "Comparison", i)
    center_x = x + panel_w + gap / 2
    if variant == "before_after":
        dc.shape("Comparison transition", center_x - 0.2, y + panel_h * 0.42,
                 0.4, 0.28, kind=MSO_SHAPE.RIGHT_ARROW,
                 fill=dc.style.colors["highlight"], line=None)
    else:
        diameter = 0.72
        cx, cy = center_x - diameter / 2, y + panel_h * 0.42
        dc.shape("Comparison VS", cx, cy, diameter, diameter,
                 kind=MSO_SHAPE.OVAL, fill=dc.style.colors["primary"],
                 line=None)
        dc.textbox("Comparison VS label", cx, cy, diameter, diameter,
                   "VS", dc.style.sizes["caption"],
                   dc.style.colors["on_primary"], bold=True,
                   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _finish_frame(dc, y + panel_h)


def _draw_matrix(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    left_pad, bottom_pad, top_pad = 0.9, 0.95, 0.2
    grid_x, grid_y = x + left_pad, y + top_pad
    grid_w = w - left_pad - 0.14
    grid_h = h - bottom_pad - top_pad
    half_w, half_h = grid_w / 2, grid_h / 2
    emphasis = _emphasis(slots, 4, dc.where)
    for i, quad in enumerate(slots["quadrants"], 1):
        col, row = (i - 1) % 2, (i - 1) // 2
        qx, qy = grid_x + col * half_w, grid_y + row * half_h
        line = dc.style.colors["highlight"] if emphasis == i \
            else dc.style.colors["line"]
        els = [dc.shape(
            f"Matrix quadrant {i}", qx + 0.025, qy + 0.025,
            half_w - 0.05, half_h - 0.05,
            fill=dc.style.colors["surface"],
            line=line)]
        if emphasis == i:
            els[0].line.width = Pt(2)
        els.append(dc.textbox(
            f"Matrix quadrant {i} title", qx + 0.12, qy + 0.08,
            half_w - 0.24, 0.38, quad["title"], dc.style.sizes["body"],
            dc.style.colors["highlight"] if emphasis == i
            else dc.style.colors["primary"], bold=True))
        if quad["items"]:
            body = quad["items"]
            size = fit_size(body, half_w - 0.24, half_h - 0.48,
                            dc.style.sizes["body"], dc.style.sizes["min"],
                            f"{dc.where}, quadrants[{i}].items")
            els.append(dc.textbox(
                f"Matrix quadrant {i} items", qx + 0.12, qy + 0.46,
                half_w - 0.24, half_h - 0.54, body, size,
                dc.style.colors["text"]))
        _group(dc, els, "Matrix", i)
    _line(dc, "Matrix vertical grid", grid_x + half_w, grid_y,
          grid_x + half_w, grid_y + grid_h, dc.style.colors["line"])
    _line(dc, "Matrix horizontal grid", grid_x, grid_y + half_h,
          grid_x + grid_w, grid_y + half_h, dc.style.colors["line"])
    dc.textbox("Matrix x low", grid_x, grid_y + grid_h + 0.04, half_w, 0.36,
               slots["x_axis"]["low"], dc.style.sizes["caption"],
               dc.style.colors["muted"])
    dc.textbox("Matrix x high", grid_x + half_w, grid_y + grid_h + 0.04,
               half_w, 0.36, slots["x_axis"]["high"],
               dc.style.sizes["caption"], dc.style.colors["muted"],
               align=PP_ALIGN.RIGHT)
    dc.textbox("Matrix x label", grid_x, grid_y + grid_h + 0.43,
               grid_w, 0.4, slots["x_axis"]["label"],
               dc.style.sizes["caption"], dc.style.colors["text"],
               bold=True, align=PP_ALIGN.CENTER)
    ylab = dc.textbox("Matrix y label", x, grid_y, 0.36, grid_h,
               slots["y_axis"]["label"], dc.style.sizes["caption"],
                dc.style.colors["text"], bold=True, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
    ylab.text_frame._txBody.bodyPr.set("vert", "vert270")
    dc.textbox("Matrix y high", x + 0.36, grid_y + 0.05, left_pad - 0.4,
               0.36, slots["y_axis"]["high"], dc.style.sizes["caption"],
               dc.style.colors["muted"], align=PP_ALIGN.CENTER)
    dc.textbox("Matrix y low", x + 0.36, grid_y + grid_h - 0.41,
               left_pad - 0.4, 0.36, slots["y_axis"]["low"],
               dc.style.sizes["caption"], dc.style.colors["muted"],
               align=PP_ALIGN.CENTER)
    _finish_frame(dc, grid_y + grid_h + 0.8)


def _draw_pyramid(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    levels = slots["levels"]
    shape_w = min(w * 0.55, h * 1.8)
    detail_x = x + shape_w + 0.42
    detail_w = w - shape_w - 0.5
    row_h = (h - 0.04 * (len(levels) - 1)) / len(levels)
    for i, level in enumerate(levels):
        display = i if variant == "pyramid" else len(levels) - i - 1
        frac = (display + 1) / len(levels)
        band_w = shape_w * (0.34 + 0.66 * frac)
        bx = x + (shape_w - band_w) / 2
        by = y + i * (row_h + 0.04)
        fill = ColorRef(
            dc.style.colors["primary"].scheme,
            dc.style.colors["primary"].rgb,
            (("lumMod", max(0.25, 1 - display * 0.12)),
             ("lumOff", min(0.75, display * 0.12))))
        kind = (MSO_SHAPE.ISOSCELES_TRIANGLE
                if variant == "pyramid" and display == 0
                else MSO_SHAPE.TRAPEZOID)
        els = [dc.shape(
            f"Level {i + 1}", bx, by, band_w, row_h,
            kind=kind, fill=fill, line=dc.style.colors["background"],
            text=level["label"] if variant == "pyramid" else None,
            size=dc.style.sizes["body"],
            color=dc.style.colors["on_primary"], bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)]
        if variant == "funnel":
            els[0].rotation = 180
            els.append(dc.textbox(
                f"Level {i + 1} label", bx + band_w * 0.1, by,
                band_w * 0.8, row_h, level["label"],
                dc.style.sizes["body"], dc.style.colors["on_primary"],
                bold=True, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE))
        if level.get("detail"):
            _line(dc, f"Level {i + 1} leader", bx + band_w,
                  by + row_h / 2, detail_x - 0.06, by + row_h / 2,
                  dc.style.colors["line"], 0.75)
            els.append(dc.textbox(
                f"Level {i + 1} detail", detail_x, by, detail_w, row_h,
                level["detail"], dc.style.sizes["body"],
                dc.style.colors["text"], anchor=MSO_ANCHOR.MIDDLE))
        _group(dc, els, "Pyramid", i + 1)
    _finish_frame(dc, y + len(levels) * row_h + 0.04 * (len(levels) - 1))


def _draw_cycle(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    steps = slots["steps"]
    node_w = min(2.0, w * 0.2)
    label_h = max(_estimated_height(
        step["label"], node_w - 0.24, dc.style.sizes["body"]
    ) / EMU_PER_IN + 0.03 for step in steps)
    details = [step.get("detail", "") for step in steps]
    detail_h = max((_estimated_height(
        detail, node_w, dc.style.sizes["body"]
    ) / EMU_PER_IN + 0.03) if detail else 0 for detail in details)
    node_h = min(h, max(0.82, label_h + detail_h
                        + (0.28 if detail_h else 0.24)))
    diameter = min(h, w - node_w + node_h)
    radius = max(0, (diameter - node_h) / 2)
    cx, cy = x + w / 2, y + h / 2
    rx = ry = radius
    nodes = []
    for i, step in enumerate(steps):
        angle = -math.pi / 2 + 2 * math.pi * i / len(steps)
        nx, ny = cx + rx * math.cos(angle), cy + ry * math.sin(angle)
        nodes.append((nx, ny))
    for i, (nx, ny) in enumerate(nodes):
        next_x, next_y = nodes[(i + 1) % len(nodes)]
        dx, dy = next_x - nx, next_y - ny
        length = math.hypot(dx, dy) or 1
        ux, uy = dx / length, dy / length
        edge = min(
            node_w / 2 / abs(ux) if ux else float("inf"),
            node_h / 2 / abs(uy) if uy else float("inf"))
        start_x, start_y = nx + ux * edge, ny + uy * edge
        end_x, end_y = next_x - ux * (edge + 0.05), next_y - uy * (edge + 0.05)
        line = dc.slide.shapes.add_connector(
            1, Inches(start_x), Inches(start_y), Inches(end_x), Inches(end_y))
        line.name = f"Cycle connector {i + 1}"
        line.line.width = Pt(1.25)
        ln = line._element.spPr.find(qn("a:ln"))
        _write_color(ln, dc.style.colors["primary"])
        arrow = OxmlElement("a:headEnd")
        arrow.set("type", "triangle")
        arrow.set("w", "med")
        arrow.set("len", "med")
        ln.append(arrow)
        style = line._element.find(qn("p:style"))
        if style is not None:
            line._element.remove(style)
        mark_el(line._element, GENERATED_MARK)
    for i, ((nx, ny), step) in enumerate(zip(nodes, steps), 1):
        color = dc.style.colors["primary"]
        els = [dc.shape(
            f"Cycle node {i}", nx - node_w / 2, ny - node_h / 2,
            node_w, node_h, fill=dc.style.colors["surface"],
            line=color, radius=True)]
        step_label_h = _estimated_height(
            step["label"], node_w - 0.24, dc.style.sizes["body"]
        ) / EMU_PER_IN + 0.03
        els.append(dc.textbox(
            f"Cycle node {i} label", nx - node_w / 2 + 0.12,
            ny - node_h / 2 + 0.12, node_w - 0.24, step_label_h,
            step["label"], dc.style.sizes["body"], dc.style.colors["text"],
            bold=True, align=PP_ALIGN.CENTER))
        if step.get("detail"):
            body = step["detail"]
            body_h = node_h - 0.24 - label_h
            size = fit_size(body, node_w, body_h,
                            dc.style.sizes["body"], dc.style.sizes["min"],
                            f"{dc.where}, steps[{i}].detail")
            els.append(dc.textbox(
                f"Cycle node {i} detail", nx - node_w / 2,
                ny - node_h / 2 + 0.12 + label_h, node_w, body_h, body, size,
                dc.style.colors["muted"], align=PP_ALIGN.CENTER))
        _group(dc, els, "Cycle", i)
    if slots.get("center"):
        center_diameter = min(1.15, diameter * 0.3)
        dc.shape("Cycle center", cx - center_diameter / 2,
                 cy - center_diameter / 2,
                 center_diameter, center_diameter, kind=MSO_SHAPE.OVAL,
                 fill=dc.style.colors["primary"], line=None,
                 text=slots["center"], size=dc.style.sizes["caption"],
                 color=dc.style.colors["on_primary"], bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _finish_frame(dc, cy + diameter / 2)


def _draw_layers(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    layers = slots["layers"]
    gap = 0.12
    row_h = min(0.7, (h - gap * (len(layers) - 1)) / len(layers))
    label_w = w * 0.25
    for i, layer in enumerate(layers, 1):
        ry = y + (i - 1) * (row_h + gap)
        els = [dc.shape(
            f"Layer {i} label", x, ry, label_w, row_h,
            fill=dc.style.colors["primary"], line=None,
            text=layer["label"], size=dc.style.sizes["body"],
            color=dc.style.colors["on_primary"], bold=True,
            align=PP_ALIGN.CENTER, radius=False, anchor=MSO_ANCHOR.MIDDLE)]
        items = layer["items"]
        item_gap = 0.12
        iw = (w - label_w - 0.2 - item_gap * (len(items) - 1)) / len(items)
        for j, item in enumerate(items):
            bx = x + label_w + 0.2 + j * (iw + item_gap)
            els.append(dc.shape(
                f"Layer {i} item {j + 1}", bx, ry, iw, row_h,
                fill=dc.style.colors["surface"], line=dc.style.colors["line"],
                text=item, size=dc.style.sizes["body"],
                color=dc.style.colors["text"], align=PP_ALIGN.LEFT,
                radius=False, anchor=MSO_ANCHOR.MIDDLE))
        _group(dc, els, "Layers", i)
    _finish_frame(dc, y + len(layers) * row_h + gap * (len(layers) - 1))


def _draw_roadmap(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    periods, tracks = slots["periods"], slots["tracks"]
    milestones = slots.get("milestones", [])
    label_w = min(2.2, w * 0.2)
    timeline_w = w - label_w
    col_w = timeline_w / len(periods)
    header_h = 0.48
    milestone_h = 0.64 if milestones else 0
    row_y = y + header_h + milestone_h
    row_h = min(0.8, (h - header_h - milestone_h) / len(tracks))
    for j, period in enumerate(periods):
        dc.shape(f"Roadmap period {j + 1}", x + label_w + j * col_w, y,
                 col_w, header_h, fill=dc.style.colors["primary"], line=None,
                 text=period, size=dc.style.sizes["caption"],
                 color=dc.style.colors["on_primary"], bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if milestones:
        for i, item in enumerate(milestones, 1):
            mx = x + label_w + (item["at"] - 0.5) * col_w
            dc.shape(f"Roadmap milestone {i}", mx - 0.09, y + header_h + 0.03,
                     0.18, 0.18, kind=MSO_SHAPE.DIAMOND,
                     fill=dc.style.colors["highlight"], line=None)
            mw = min(1.6, max(0.8, col_w * 1.8))
            dc.textbox(f"Roadmap milestone {i} label", mx - mw / 2,
                       y + header_h + 0.23, mw, 0.38, item["label"],
                       dc.style.sizes["caption"], dc.style.colors["text"],
                       align=PP_ALIGN.CENTER)
    for i in range(len(tracks)):
        ry = row_y + i * row_h
        if i % 2 == 1:
            dc.shape(f"Roadmap row {i + 1}", x, ry, w, row_h,
                     fill=dc.style.colors["surface"], line=None)
    for j in range(len(periods) + 1):
        gx = x + label_w + j * col_w
        _line(dc, f"Roadmap grid {j + 1}", gx, row_y, gx,
              row_y + len(tracks) * row_h, dc.style.colors["line"])
    for i, track in enumerate(tracks, 1):
        ry = row_y + (i - 1) * row_h
        dc.textbox(f"Roadmap track {i}", x + 0.04, ry + 0.03,
                   label_w - 0.1, row_h - 0.06, track["label"],
                   dc.style.sizes["body"], dc.style.colors["text"],
                   bold=True, anchor=MSO_ANCHOR.MIDDLE)
        for j, bar in enumerate(track["bars"], 1):
            bx = x + label_w + (bar["start"] - 1) * col_w + 0.04
            bw = (bar["end"] - bar["start"] + 1) * col_w - 0.08
            fill = dc.style.colors["highlight"] if bar.get("emphasis") \
                else dc.style.colors["primary"]
            els = [dc.shape(
                f"Roadmap track {i} bar {j}", bx, ry + (row_h - 0.42) / 2,
                bw, 0.42, fill=fill, line=None,
                text=bar.get("label", track["label"]),
                size=dc.style.sizes["caption"],
                color=dc.style.colors["on_highlight" if bar.get("emphasis")
                                      else "on_primary"],
                bold=True, align=PP_ALIGN.CENTER, radius=True,
                anchor=MSO_ANCHOR.MIDDLE)]
            _group(dc, els, "Roadmap", (i - 1) * 10 + j)
    _finish_frame(dc, row_y + len(tracks) * row_h)


def _draw_message(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    supports = slots.get("supports", [])
    if variant == "banner":
        statement_h = max(0.55, _estimated_height(
            slots["statement"], w - 0.4, 24) / EMU_PER_IN + 0.08)
        dc.shape("Message banner accent", x, y, 0.06, statement_h,
                 fill=dc.style.colors["primary"], line=None)
        dc.textbox("Message banner", x + 0.2, y, w - 0.2, statement_h,
                   slots["statement"], 24, dc.style.colors["primary"],
                   bold=True, heading=True)
        support_y = y + statement_h + 0.28
        support_h = 0
        if supports:
            gap = 0.16
            cw = (w - gap * (len(supports) - 1)) / len(supports)
            support_h = max(_text_height(
                text, cw - 0.24, dc.style.sizes["body"], padding=0.24,
                minimum=0.9) for text in supports)
            for i, text in enumerate(supports, 1):
                els, _ = _panel(
                    dc, f"Message support {i}",
                    x + (i - 1) * (cw + gap), support_y, cw,
                    None, text, dc.style.colors["background"],
                    dc.style.colors["line"], heading_color=dc.style.colors[
                        "primary"], min_height=support_h)
                _group(dc, els, "Message", i)
        _finish_frame(dc, support_y + support_h if supports
                      else y + statement_h)
    else:
        quote_h = 0.52
        dc.textbox("Message quote mark", x, y, 0.55, quote_h, "“",
                   40, dc.style.colors["highlight"], bold=True,
                   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        statement_h = max(0.8, _estimated_height(
            slots["statement"], w - 0.8, 22) / EMU_PER_IN + 0.08)
        dc.textbox("Message quote", x + 0.52, y, w - 0.7, statement_h,
                   slots["statement"], 22, dc.style.colors["text"],
                   bold=True, heading=True, align=PP_ALIGN.CENTER)
        support_y = y + statement_h + 0.22
        support_h = 0
        if supports:
            support_h = _estimated_height(
                supports, w - 0.8, dc.style.sizes["caption"]
            ) / EMU_PER_IN + 0.08
            dc.textbox("Message quote support", x + 0.4, support_y,
                       w - 0.8, support_h, supports,
                       dc.style.sizes["caption"], dc.style.colors["muted"],
                       align=PP_ALIGN.CENTER)
        _finish_frame(dc, support_y + support_h if supports
                      else y + statement_h)


def _draw_checklist(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    items = slots["items"]
    has_meta = any(item.get("owner") or item.get("due") for item in items)
    header_h = 0.38 if has_meta else 0
    if has_meta:
        for name, cx, cw in (
                ("項目", x + 0.52, w * 0.53),
                ("担当", x + w * 0.72, w * 0.15),
                ("期限", x + w * 0.88, w * 0.11)):
            dc.textbox(f"Checklist heading {name}", cx, y, cw, header_h,
                       name, dc.style.sizes["caption"],
                       dc.style.colors["text"], bold=True)
        _line(dc, "Checklist heading divider", x, y + header_h, x + w,
              y + header_h, dc.style.colors["line"], 0.75)
    row_h = min(0.64, (h - header_h) / len(items))
    text_w = w * (0.52 if has_meta else 0.82)
    text_size = fit_siblings(
        [item["text"] for item in items], text_w - 0.06,
        row_h - 0.04, dc.style.sizes["body"], dc.style.sizes["min"],
        f"{dc.where}, items[].text")
    for i, item in enumerate(items, 1):
        ry = y + header_h + (i - 1) * row_h
        status = item.get("status", "todo")
        color = dc.style.colors["primary"] if status == "done" \
            else dc.style.colors["highlight"] if status == "risk" else None
        icon = "✓" if status == "done" else "!" if status == "risk" else ""
        icon_size = 0.38
        status_shape = dc.shape(
            f"Checklist status {i}", x + 0.03,
            ry + (row_h - icon_size) / 2, icon_size, icon_size,
            kind=MSO_SHAPE.OVAL, fill=color,
            line=dc.style.colors["line"] if status == "todo" else None,
            text=icon or None, size=round(icon_size * 72 * 0.55),
            color=dc.style.colors["on_primary" if status == "done"
                                  else "on_highlight"],
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if icon:
            frame = status_shape.text_frame
            frame.margin_left = Inches(0.04)
            frame.margin_right = Inches(0.04)
            frame.margin_top = Inches(0.04)
            frame.margin_bottom = Inches(0.04)
            frame.paragraphs[0].runs[0].font.size = Pt(
                round(icon_size * 72 * 0.55))
        els = [status_shape]
        els.append(dc.textbox(
            f"Checklist item {i}", x + 0.42, ry + 0.02, text_w - 0.42,
            row_h - 0.04, item["text"], text_size,
            dc.style.colors["text"], anchor=MSO_ANCHOR.MIDDLE))
        if has_meta:
            els.append(dc.textbox(
                f"Checklist owner {i}", x + w * 0.72, ry + 0.02,
                w * 0.15, row_h - 0.04, item.get("owner", ""),
                dc.style.sizes["body"], dc.style.colors["text"],
                anchor=MSO_ANCHOR.MIDDLE))
            els.append(dc.textbox(
                f"Checklist due {i}", x + w * 0.88, ry + 0.02,
                w * 0.11, row_h - 0.04, item.get("due", ""),
                dc.style.sizes["body"], dc.style.colors["text"],
                anchor=MSO_ANCHOR.MIDDLE))
        if i > 1:
            _line(dc, f"Checklist divider {i}", x, ry, x + w, ry,
                  dc.style.colors["line"], 0.75)
        _group(dc, els, "Checklist", i)
    _finish_frame(dc, y + header_h + len(items) * row_h)


def _apply_chart_colors(chart, style):
    colors = [style.colors["primary"], style.colors["highlight"]]
    colors.extend(_scheme_ref(f"accent{i}", {}) for i in range(3, 7))
    for series, color in zip(chart.plots[0].series, colors):
        sp_pr = series._element.find(qn("c:spPr"))
        if sp_pr is None:
            sp_pr = OxmlElement("c:spPr")
            series._element.append(sp_pr)
        _write_color(sp_pr, color)
    typeface = style.fonts["body_latin"]
    for tx_pr in chart._element.iter(qn("c:txPr")):
        for latin in tx_pr.iter(qn("a:latin")):
            latin.set("typeface", typeface)
        for run in list(tx_pr.iter(qn("a:defRPr"))) + list(
                tx_pr.iter(qn("a:rPr"))) + list(
                tx_pr.iter(qn("a:endParaRPr"))):
            run.set("sz", str(style.sizes["body"] * 100))
            _write_color(run, style.colors["text"])


def _draw_chart(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    chart_value = slots["chart"]
    points = slots.get("points", [])
    panel_h = min(h, max(1.2, len(points) * 0.62 + 0.2)) if points else 0
    chart_h = h
    if variant == "full":
        frame = dc.slide.shapes.add_chart(
            CHART_TYPES[chart_value["type"]], Inches(x), Inches(y),
            Inches(w), Inches(chart_h), chart_data(chart_value))
        frame.name = "Library chart"
        _apply_chart_colors(frame.chart, dc.style)
        finish_chart(frame.chart, chart_value)
        mark_el(frame._element, GENERATED_MARK)
        _finish_frame(dc, y + chart_h)
        return
    chart_w = w * 0.62
    frame = dc.slide.shapes.add_chart(
        CHART_TYPES[chart_value["type"]], Inches(x), Inches(y),
        Inches(chart_w), Inches(chart_h), chart_data(chart_value))
    frame.name = "Library chart"
    _apply_chart_colors(frame.chart, dc.style)
    finish_chart(frame.chart, chart_value)
    mark_el(frame._element, GENERATED_MARK)
    panel_x = x + chart_w + 0.12
    panel_w = w - chart_w - 0.12
    panel_y = y + (chart_h - panel_h) / 2
    dc.shape("Chart key points panel", panel_x, panel_y, panel_w, panel_h,
             fill=dc.style.colors["surface"], line=dc.style.colors["line"],
             radius=True)
    row_h = (panel_h - 0.12) / len(points)
    for i, point in enumerate(points, 1):
        ry = panel_y + 0.06 + (i - 1) * row_h
        number = dc.shape(
            f"Chart point {i} number", panel_x + 0.13,
            ry + (row_h - 0.42) / 2, 0.42, 0.42,
            kind=MSO_SHAPE.OVAL, fill=dc.style.colors["primary"], line=None,
            text=str(i), size=dc.style.sizes["caption"],
            color=dc.style.colors["on_primary"], bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        label = dc.textbox(
            f"Chart point {i}", panel_x + 0.54, ry + 0.02,
            panel_w - 0.66, row_h - 0.04, point,
            dc.style.sizes["body"], dc.style.colors["text"],
            anchor=MSO_ANCHOR.MIDDLE)
        _group(dc, [number, label], "Chart", i)
    _finish_frame(dc, y + chart_h)


def _draw_component(name, dc, region, slots, variant):
    draw = {
        "cards": _draw_cards, "kpi": _draw_kpi,
        "process": _draw_process, "comparison": _draw_comparison,
        "matrix": _draw_matrix, "pyramid": _draw_pyramid,
        "cycle": _draw_cycle, "layers": _draw_layers,
        "roadmap": _draw_roadmap, "message": _draw_message,
        "checklist": _draw_checklist, "chart": _draw_chart,
    }[name]
    draw(dc, region, slots, variant)


_BODY_ITEMS = _items_schema(
    {"heading": {"type": "string", "minLength": 1},
     "body": _string_or_list()},
    ("heading", "body"), 2, 6)
_KPI_ITEMS = _items_schema(
    {"value": {"type": "string", "minLength": 1},
     "unit": {"type": "string"},
     "label": {"type": "string", "minLength": 1},
     "note": {"type": "string"}},
    ("value", "label"), 1, 4)
_PROCESS_ITEMS = _items_schema(
    {"label": {"type": "string", "minLength": 1},
     "detail": _string_or_list()},
    ("label",), 2, 6)
_COMPARISON_PANEL = {
    "type": "object",
    "required": ["heading", "items"],
    "properties": {
        "heading": {"type": "string", "minLength": 1},
        "items": {"type": "array", "minItems": 1, "maxItems": 6,
                  "items": {"type": "string"}},
    },
    "additionalProperties": False,
}
_LEVELS = _items_schema(
    {"label": {"type": "string", "minLength": 1},
     "detail": {"type": "string"}},
    ("label",), 2, 5)
_LAYERS = _items_schema(
    {"label": {"type": "string", "minLength": 1},
     "items": {"type": "array", "minItems": 1, "maxItems": 5,
               "items": {"type": "string"}}},
    ("label", "items"), 2, 6)
_CHECKLIST = _items_schema(
    {"text": {"type": "string", "minLength": 1},
     "status": {"enum": ["done", "todo", "risk"]},
     "owner": {"type": "string"}, "due": {"type": "string"}},
    ("text",), 1, 8)
_CHART = {
    "type": "object",
    "required": ["type", "categories", "series"],
    "properties": {
        "type": {"enum": sorted(CHART_TYPES)},
        "categories": {"type": "array", "minItems": 1,
                       "items": {"type": ["string", "number"]}},
        "series": {"type": "array", "minItems": 1, "items": {
            "type": "object", "required": ["name", "values"],
            "properties": {"name": {"type": "string"},
                           "values": {"type": "array", "items":
                                      {"type": "number"}}},
            "additionalProperties": False,
        }},
        "number_format": {"type": "string"},
    },
    "additionalProperties": False,
}
_ROADMAP_TRACKS = _items_schema(
    {"label": {"type": "string", "minLength": 1},
     "bars": {"type": "array", "minItems": 1, "items": {
         "type": "object", "required": ["start", "end", "label"],
         "properties": {
             "start": {"type": "integer", "minimum": 1},
             "end": {"type": "integer", "minimum": 1},
             "label": {"type": "string"},
             "emphasis": {"type": "boolean"},
         }, "additionalProperties": False,
     }},
     "required": ["start", "end"]},
    ("label", "bars"), 1, 8)


def _base_properties(extra, emphasis=False):
    result = dict(extra)
    if emphasis:
        result["emphasis"] = {"type": "integer", "minimum": 1}
    return result


LIBRARY = {
    "cards": Component(
        ("header", "outline", "numbered"),
        _component_schema(_base_properties({"items": _BODY_ITEMS}, True),
                          ("items",)),
        lambda dc, r, s, v: _draw_component("cards", dc, r, s, v)),
    "kpi": Component(
        ("tiles", "band"),
        _component_schema(_base_properties({"items": _KPI_ITEMS}),
                          ("items",)),
        lambda dc, r, s, v: _draw_component("kpi", dc, r, s, v)),
    "process": Component(
        ("chevron", "circles", "arrows"),
        _component_schema(_base_properties({"steps": _PROCESS_ITEMS}, True),
                          ("steps",)),
        lambda dc, r, s, v: _draw_component("process", dc, r, s, v)),
    "comparison": Component(
        ("before_after", "versus"),
        _component_schema(_base_properties({
            "left": _COMPARISON_PANEL, "right": _COMPARISON_PANEL,
            "emphasis": {"type": "integer", "minimum": 1, "maximum": 2},
        }), ("left", "right")),
        lambda dc, r, s, v: _draw_component("comparison", dc, r, s, v)),
    "matrix": Component(
        ("matrix",),
        _component_schema(_base_properties({
            "x_axis": {"type": "object", "required": ["label", "low", "high"],
                       "properties": {k: {"type": "string"} for k in
                                      ("label", "low", "high")},
                       "additionalProperties": False},
            "y_axis": {"type": "object", "required": ["label", "low", "high"],
                       "properties": {k: {"type": "string"} for k in
                                      ("label", "low", "high")},
                       "additionalProperties": False},
            "quadrants": {
                "type": "array", "minItems": 4, "maxItems": 4,
                "items": {"type": "object", "required": ["title", "items"],
                          "properties": {
                              "title": {"type": "string", "minLength": 1},
                              "items": {"type": "array", "maxItems": 3,
                                        "items": {"type": "string"}},
                          }, "additionalProperties": False},
            },
        }, True), ("x_axis", "y_axis", "quadrants")),
        lambda dc, r, s, v: _draw_component("matrix", dc, r, s, v)),
    "pyramid": Component(
        ("pyramid", "funnel"),
        _component_schema(_base_properties({"levels": _LEVELS}),
                          ("levels",)),
        lambda dc, r, s, v: _draw_component("pyramid", dc, r, s, v)),
    "cycle": Component(
        ("cycle",),
        _component_schema(_base_properties({
            "steps": {"type": "array", "minItems": 3, "maxItems": 6,
                      "items": {"type": "object",
                                "required": ["label"],
                                "properties": {
                                    "label": {"type": "string",
                                              "minLength": 1},
                                    "detail": _string_or_list()},
                                "additionalProperties": False}},
            "center": {"type": "string"},
        }), ("steps",)),
        lambda dc, r, s, v: _draw_component("cycle", dc, r, s, v)),
    "layers": Component(
        ("layers",),
        _component_schema(_base_properties({"layers": _LAYERS}),
                          ("layers",)),
        lambda dc, r, s, v: _draw_component("layers", dc, r, s, v)),
    "roadmap": Component(
        ("roadmap",),
        _component_schema(_base_properties({
            "periods": {"type": "array", "minItems": 2, "maxItems": 12,
                        "items": {"type": "string"}},
            "tracks": _ROADMAP_TRACKS,
            "milestones": {"type": "array", "items": {
                "type": "object", "required": ["at", "label"],
                "properties": {"at": {"type": "integer", "minimum": 1},
                               "label": {"type": "string"}},
                "additionalProperties": False}},
        }), ("periods", "tracks")),
        lambda dc, r, s, v: _draw_component("roadmap", dc, r, s, v)),
    "message": Component(
        ("banner", "quote"),
        _component_schema(_base_properties({
            "statement": {"type": "string", "minLength": 1},
            "supports": {"type": "array", "maxItems": 3,
                         "items": {"type": "string"}},
        }), ("statement",)),
        lambda dc, r, s, v: _draw_component("message", dc, r, s, v)),
    "checklist": Component(
        ("checklist",),
        _component_schema(_base_properties({"items": _CHECKLIST}),
                          ("items",)),
        lambda dc, r, s, v: _draw_component("checklist", dc, r, s, v)),
    "chart": Component(
        ("side", "full"),
        _component_schema(_base_properties({
            "chart": _CHART,
            "points": {"type": "array", "minItems": 1, "maxItems": 4,
                       "items": {"type": "string"}},
        }), ("chart",)),
        lambda dc, r, s, v: _draw_component("chart", dc, r, s, v)),
}


import components_business  # noqa: E402

components_business.register(sys.modules[__name__])


def resolve_component(name, tmap, prefer=None):
    if prefer == "library":
        if name in LIBRARY:
            return "library", LIBRARY[name]
        fail(f"component '{name}': prefer=library but the library has no "
             f"such component (library: {sorted(LIBRARY)})")
    if name in tmap.get("components", {}):
        return "template", tmap["components"][name]
    if name in LIBRARY:
        return "library", LIBRARY[name]
    fail(f"component '{name}' is not in template map or library "
         f"(template: {sorted(tmap.get('components', {}))}; "
         f"library: {sorted(LIBRARY)})")


def validate_component(name, slots, variant, slide):
    component = LIBRARY[name]
    where = f"slide {slide} (component '{name}')"
    variant = variant or component.variants[0]
    if variant not in component.variants:
        fail(f"{where}: unknown variant '{variant}' "
             f"(choose from {list(component.variants)})")
    validator = jsonschema.Draft7Validator(component.slots_schema)
    errors = sorted(validator.iter_errors(slots),
                    key=lambda error: list(map(str, error.absolute_path)))
    if errors:
        error = errors[0]
        path = ""
        for part in error.absolute_path:
            path += f"[{part}]" if isinstance(part, int) else (
                f".{part}" if path else str(part))
        if error.validator == "maxItems":
            fail(f"{where}{', ' + path if path else ''}: {error.message}; "
                 "split the content across slides")
        fail(f"{where}{', ' + path if path else ''}: {error.message}")
    emphasis = slots.get("emphasis")
    count = len(slots.get("items", [])) if name in ("cards", "checklist") \
        else len(slots.get("steps", [])) if name == "process" \
        else 4 if name == "matrix" else 0
    if emphasis is not None and count and emphasis > count:
        fail(f"{where}, emphasis: expected a number from 1 to {count}")
    if name == "roadmap":
        periods = len(slots["periods"])
        for i, track in enumerate(slots["tracks"], 1):
            for j, bar in enumerate(track["bars"], 1):
                for key in ("start", "end"):
                    if bar[key] > periods:
                        fail(f"{where}, tracks[{i}].bars[{j}].{key}: "
                             f"must be within 1..{periods}")
                if bar["start"] > bar["end"]:
                    fail(f"{where}, tracks[{i}].bars[{j}]: start "
                         "must be <= end")
        for i, milestone in enumerate(slots.get("milestones", []), 1):
            if milestone["at"] > periods:
                fail(f"{where}, milestones[{i}].at: must be within "
                     f"1..{periods}")
    if component.check is not None:
        component.check(slots, variant, where)
    if name == "chart":
        if variant == "side" and not slots.get("points"):
            fail(f"{where}, points: side variant requires 1..4 points")
        if variant == "full" and "points" in slots:
            fail(f"{where}, points: full variant does not accept points")
        value = slots["chart"]
        if any(len(series["values"]) != len(value["categories"])
               for series in value["series"]):
            fail(f"{where}, chart.series: each series must match categories")
    return variant


def resolve_canvas(prs, tmap, slide):
    canvas = tmap.get("canvas", {})
    layout_name = canvas.get("layout")
    if layout_name is None:
        layout_name = next((name for name in ("Title Only", "タイトルのみ")
                            if any(layout.name == name
                                   for layout in prs.slide_layouts)), None)
    layout = next((layout for layout in prs.slide_layouts
                   if layout.name == layout_name), None)
    if layout is None:
        fail(f"slide {slide}: canvas layout '{layout_name}' not found; "
             "map に canvas.layout を指定")
    title_idx = canvas.get("title")
    title_ph = None
    if title_idx is not None:
        title_ph = next((ph for ph in layout.placeholders
                         if ph.placeholder_format.idx == title_idx), None)
    else:
        title_ph = next((ph for ph in layout.placeholders
                         if str(ph.placeholder_format.type)
                         in ("TITLE (1)", "CENTER_TITLE (3)")), None)
        if title_ph is not None:
            title_idx = title_ph.placeholder_format.idx
    if title_idx is not None and title_ph is None:
        title_ph = next((ph for ph in layout.placeholders
                         if ph.placeholder_format.idx == title_idx), None)
    custom_region = canvas.get("region")
    if custom_region is not None:
        region = tuple(float(v) for v in custom_region)
        if region[0] < 0 or region[1] < 0 or region[2] <= 0 or region[3] <= 0:
            fail(f"slide {slide}: canvas.region requires nonnegative x/y "
                 "and positive width/height")
    else:
        if title_ph is None:
            fail(f"slide {slide}: layout '{layout_name}' has no title "
                 "placeholder; specify map.canvas.title and canvas.region")
        x, _, w, _ = (title_ph.left / EMU_PER_IN,
                      title_ph.top / EMU_PER_IN,
                      title_ph.width / EMU_PER_IN,
                      title_ph.height / EMU_PER_IN)
        y = (title_ph.top + title_ph.height) / EMU_PER_IN + 0.25
        region = (x, y, w, prs.slide_height / EMU_PER_IN - y - 0.75)
    return layout, title_idx, region


def is_text_only(source, component):
    """True for template content components that only carry text/list."""
    return (source == "template" and component.get("kind") == "content"
            and all(slot.get("type") in ("text", "list")
                    for slot in component.get("slots", {}).values()))


def draw_component(prs, slide, name, spec, tmap, slide_i, variant,
                   region, title_idx):
    component = LIBRARY[name]
    style = resolve_style(prs, tmap)
    title = spec.get("slots", {}).get("title") or spec.get("message")
    if title_idx is not None and title:
        try:
            title_shape = slide.placeholders[title_idx]
        except KeyError:
            fail(f"slide {slide_i}: canvas title placeholder "
                 f"{title_idx} not on slide")
        title_shape.text = title
    slots = dict(spec.get("slots") or {})
    slots.pop("title", None)
    validate_component(name, slots, variant, slide_i)
    dc = DrawContext(slide, style, f"slide {slide_i} (component '{name}')")
    component.draw(dc, region, slots, variant)
    _center_component(dc)
    return variant
