#!/usr/bin/env python3
"""Detect and use PowerPoint or LibreOffice to render presentation files."""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation


POWERSHELL_SCRIPT = r"""
param([string]$In, [string]$Out, [int]$Width = 960, [switch]$Pdf,
      [switch]$Probe)
$ErrorActionPreference = "Stop"
$app = $null
$presentation = $null
try {
    $app = New-Object -ComObject PowerPoint.Application
    if ($Probe) { Write-Output ("VERSION=" + $app.Version) }
    $presentation = $app.Presentations.Open($In, -1, 0, 0)
    if ($Pdf) {
        $presentation.SaveAs($Out, 32)
    } else {
        $height = [int][Math]::Round(
            $Width * $presentation.PageSetup.SlideHeight /
            $presentation.PageSetup.SlideWidth)
        for ($i = 1; $i -le $presentation.Slides.Count; $i++) {
            $path = Join-Path $Out ("{0}-{1:D2}.png" -f
                [IO.Path]::GetFileNameWithoutExtension($In), $i)
            $presentation.Slides.Item($i).Export($path, "PNG", $Width, $height)
        }
    }
} finally {
    if ($presentation) {
        try { $presentation.Close() } catch {}
        [void][Runtime.InteropServices.Marshal]::ReleaseComObject($presentation)
    }
    if ($app) {
        try {
            if ($app.Presentations.Count -eq 0) { $app.Quit() }
        } catch {}
        [void][Runtime.InteropServices.Marshal]::ReleaseComObject($app)
    }
}
"""
WINDOWS_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def _find_soffice():
    env = os.environ.get("SOFFICE")
    if env and Path(env).is_file():
        return env
    on_path = shutil.which("soffice")
    if on_path:
        return on_path
    if Path(WINDOWS_SOFFICE).is_file():
        return WINDOWS_SOFFICE
    return None


def find_soffice():
    path = _find_soffice()
    if path is None:
        raise RuntimeError(
            "soffice not found. Set SOFFICE env var, put soffice on PATH, "
            "or install LibreOffice.")
    return path


def _powershell():
    return shutil.which("powershell.exe") or shutil.which("pwsh")


def _run_powershell(args, timeout=300):
    powershell = _powershell()
    if not powershell:
        raise RuntimeError("PowerShell (powershell.exe or pwsh) not found")
    with tempfile.TemporaryDirectory() as td:
        script = Path(td) / "pptx-deck.ps1"
        script.write_text(POWERSHELL_SCRIPT, encoding="ascii")
        return subprocess.run(
            [powershell, "-NoProfile", "-NonInteractive",
             "-ExecutionPolicy", "Bypass", "-File", str(script), *args],
            capture_output=True, text=True, timeout=timeout)


def _probe_powerpoint():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        pptx_path = td / "probe.pptx"
        png_dir = td / "png"
        png_dir.mkdir()
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(str(pptx_path))
        proc = _run_powershell(
            ["-In", str(pptx_path), "-Out", str(png_dir), "-Probe"],
            timeout=90)
        if proc.returncode != 0:
            return {"ok": False, "version": None,
                    "error": proc.stderr.strip() or proc.stdout.strip()
                    or f"PowerShell exited with status {proc.returncode}"}
        match = re.search(r"VERSION=([^\r\n]+)", proc.stdout)
        png_path = png_dir / "probe-01.png"
        if not png_path.is_file():
            return {"ok": False, "version": match.group(1) if match else None,
                    "error": "PowerPoint probe did not export a PNG"}
        return {"ok": True, "version": match.group(1) if match else None,
                "error": None}


def detect():
    powerpoint = {"ok": False, "version": None, "error": None}
    forced = os.environ.get("PPTX_DECK_ENGINE", "").strip().lower()
    if forced in ("powerpoint", "libreoffice"):
        powerpoint["error"] = "not probed (PPTX_DECK_ENGINE)"
        return {"engine": forced, "reason": "PPTX_DECK_ENGINE",
                "powerpoint": powerpoint, "soffice": _find_soffice()}

    probe_error = None
    if os.name == "nt":
        if _powershell():
            try:
                powerpoint = _probe_powerpoint()
            except subprocess.TimeoutExpired:
                powerpoint["error"] = "PowerPoint probe timed out after 90 seconds"
            except Exception as e:
                powerpoint["error"] = str(e)
            if powerpoint["ok"]:
                return {"engine": "powerpoint",
                        "reason": f"PowerPoint {powerpoint['version'] or 'available'}",
                        "powerpoint": powerpoint, "soffice": _find_soffice()}
            probe_error = powerpoint["error"]
        else:
            powerpoint["error"] = "PowerShell (powershell.exe or pwsh) not found"
            probe_error = powerpoint["error"]

    soffice = _find_soffice()
    if soffice:
        reason = f"soffice found: {soffice}"
        if probe_error:
            reason = f"PowerPoint probe failed: {probe_error}; {reason}"
        return {"engine": "libreoffice", "reason": reason,
                "powerpoint": powerpoint, "soffice": soffice}
    reason = "no rendering engine found"
    if probe_error:
        reason = f"PowerPoint probe failed: {probe_error}; {reason}"
    return {"engine": "none", "reason": reason,
            "powerpoint": powerpoint, "soffice": None}


def resolve(engine_arg="auto"):
    if engine_arg == "auto":
        return detect()["engine"]
    if engine_arg not in ("powerpoint", "libreoffice", "none"):
        raise ValueError(f"unknown rendering engine: {engine_arg}")
    return engine_arg


def _libreoffice_pdf(pptx_path, pdf_path):
    pptx_path, pdf_path = Path(pptx_path).resolve(), Path(pdf_path).resolve()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        proc = subprocess.run(
            [find_soffice(), "--headless", "--norestore", "--convert-to",
             "pdf", "--outdir", td, str(pptx_path)],
            capture_output=True, text=True, timeout=300)
        converted = Path(td) / f"{pptx_path.stem}.pdf"
        if proc.returncode != 0 or not converted.is_file():
            raise RuntimeError(
                "soffice PDF conversion failed: "
                f"{proc.stderr.strip() or proc.stdout.strip()}")
        shutil.move(str(converted), str(pdf_path))
    return str(pdf_path)


def export_pdf(pptx, pdf_path, engine):
    engine = resolve(engine)
    if engine == "powerpoint":
        pdf_path = Path(pdf_path).resolve()
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        proc = _run_powershell(
            ["-In", str(Path(pptx).resolve()), "-Out", str(pdf_path), "-Pdf"])
        if proc.returncode != 0 or not pdf_path.is_file():
            raise RuntimeError(
                f"PowerPoint PDF export failed: "
                f"{proc.stderr.strip() or proc.stdout.strip()}")
        return str(pdf_path)
    if engine == "libreoffice":
        return _libreoffice_pdf(pptx, pdf_path)
    raise RuntimeError("no rendering engine is available")


def render_pngs(pptx, out_dir, engine, width_px=960):
    engine = resolve(engine)
    pptx_path = Path(pptx).resolve()
    out_dir = Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    if engine == "powerpoint":
        proc = _run_powershell(
            ["-In", str(pptx_path), "-Out", str(out_dir),
             "-Width", str(width_px)])
        if proc.returncode != 0:
            raise RuntimeError(
                f"PowerPoint PNG export failed: "
                f"{proc.stderr.strip() or proc.stdout.strip()}")
        return sorted(str(p) for p in out_dir.glob(f"{pptx_path.stem}-*.png"))
    if engine == "libreoffice":
        import pypdfium2 as pdfium

        with tempfile.TemporaryDirectory() as td:
            pdf_path = _libreoffice_pdf(pptx_path, Path(td) / "render.pdf")
            doc = pdfium.PdfDocument(pdf_path)
            paths = []
            try:
                for i in range(len(doc)):
                    page = doc[i]
                    img = page.render(
                        scale=width_px / page.get_width()).to_pil()
                    path = out_dir / f"{pptx_path.stem}-{i + 1:02d}.png"
                    img.save(str(path))
                    paths.append(str(path))
            finally:
                doc.close()
            return paths
    raise RuntimeError("no rendering engine is available")


def _summary(result):
    engine = result["engine"]
    if engine == "powerpoint":
        return f"engine: powerpoint (PowerPoint {result['powerpoint']['version']})"
    if engine == "libreoffice":
        suffix = (f"; {result['reason'].split('; ', 1)[0]}"
                  if result["powerpoint"]["error"] else "")
        return f"engine: libreoffice ({result['soffice']}{suffix})"
    return f"engine: none ({result['reason']})"


def main(argv=None):
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="command", required=True)
    detect_parser = sub.add_parser("detect")
    detect_parser.add_argument("--json", dest="json_path")
    render_parser = sub.add_parser("render")
    render_parser.add_argument("pptx")
    render_parser.add_argument("out_dir")
    render_parser.add_argument("--engine", default="auto",
                               choices=("auto", "powerpoint", "libreoffice"))
    pdf_parser = sub.add_parser("pdf")
    pdf_parser.add_argument("pptx")
    pdf_parser.add_argument("pdf_path")
    pdf_parser.add_argument("--engine", default="auto",
                            choices=("auto", "powerpoint", "libreoffice"))
    args = ap.parse_args(argv)
    try:
        if args.command == "detect":
            result = detect()
            print(_summary(result))
            if args.json_path:
                Path(args.json_path).write_text(
                    json.dumps(result, indent=2), encoding="utf-8")
            return 3 if result["engine"] == "none" else 0
        if args.command == "render":
            paths = render_pngs(args.pptx, args.out_dir, args.engine)
            print(f"rendered {len(paths)} slide(s) to {args.out_dir}")
        else:
            path = export_pdf(args.pptx, args.pdf_path, args.engine)
            print(f"wrote {path}")
    except Exception as e:
        print(f"engine error: {e}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
