"""Business diagram components: tables, scorecards, timelines, logic trees,
hub diagrams, swimlanes and action plans.

Registered into components.LIBRARY by components.py (``register``). Helpers
are injected from components.py at registration time so this module never
imports components.py at import time (no circular import).
"""
import math

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from common import GENERATED_MARK, char_units, fail, mark_el

_HELPERS = (
    "EMU_PER_IN", "ColorRef", "Component", "_component_schema",
    "_draw_frame", "_finish_frame", "_estimated_height", "fit_size",
    "fit_siblings", "_line", "_group", "_write_color", "_set_fonts",
    "_text_color",
)

RATING_SCORES = {"◎": 3, "○": 2, "△": 1, "×": 0, "-": None, "": None}
RACI_ROLES = ("R", "A", "C", "I")
RACI_LEGEND = "R 実行　A 説明責任　C 相談　I 報告"
TONES = ("positive", "caution", "negative", "neutral")


def register(module):
    """Bind helpers from components.py and add components to LIBRARY."""
    g = globals()
    for name in _HELPERS:
        g[name] = getattr(module, name)
    module.LIBRARY.update(_library())


# ---------- small helpers ----------

def _tint(color, amount=0.18):
    return ColorRef(color.scheme, color.rgb,
                    color.transforms + (("lumMod", amount),
                                        ("lumOff", 1 - amount)))


def _h(text, width, size):
    return _estimated_height(text, width, size) / EMU_PER_IN


def _arrow(dc, name, x1, y1, x2, y2, color=None, width=1.25, elbow=False):
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    line = dc.slide.shapes.add_connector(
        kind, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.name = name
    line.line.width = Pt(width)
    ln = line._element.spPr.find(qn("a:ln"))
    _write_color(ln, color or dc.style.colors["primary"])
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    ln.append(tail)
    style = line._element.find(qn("p:style"))
    if style is not None:
        line._element.remove(style)
    mark_el(line._element, GENERATED_MARK)
    return line


def _chip(dc, name, x, y, w, h, text, fill, color, size=None, line=None):
    """Small label (ID, badge, RACI role). Chips hold a few characters, so
    they skip the paragraph fit check and use tight margins instead."""
    shape = dc.shape(name, x, y, w, h, fill=fill, line=line, radius=True)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right", "margin_top",
                 "margin_bottom"):
        setattr(frame, side, Inches(0.01))
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size or dc.style.sizes["caption"])
    run.font.bold = True
    _set_fonts(run, dc.style.fonts, True)
    _text_color(run, color)
    return shape


def _fit_common(texts, width, height, base, dc, where):
    texts = [t for t in texts if t]
    if not texts:
        return base
    return fit_siblings(texts, width, height, base, dc.style.sizes["min"],
                        where)


def _single_line_size(texts, width, base, minimum, pad=0.24):
    """Largest size (<= base) at which every text stays on one line.

    ``pad`` is the horizontal text inset of the box the text goes into."""
    widest = max((sum(char_units(c) for c in t) for t in texts if t),
                 default=0)
    if not widest:
        return base
    size = int((width - pad) * 72 / widest)
    return max(minimum, min(base, size))


def _cell(value):
    if isinstance(value, dict):
        return value["text"], value.get("tone"), value.get("bold", False)
    return value, None, False


# ---------- table ----------

def _set_cell(dc, cell, text, size, fill, color, bold=False, heading=False,
              align=PP_ALIGN.LEFT, bottom=None):
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.05)
    cell.margin_bottom = Inches(0.05)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0, 0, 0)
    tc_pr = cell._tc.get_or_add_tcPr()
    _write_color(tc_pr, fill)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tc_pr.find(qn(tag))
        if old is not None:
            tc_pr.remove(old)
    fill_index = list(tc_pr).index(tc_pr.find(qn("a:solidFill")))
    for offset, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = OxmlElement(tag)
        if tag == "a:lnB" and bottom is not None:
            ln.set("w", str(Pt(0.75)))
            _write_color(ln, bottom)
        else:
            ln.set("w", "0")
            ln.append(OxmlElement("a:noFill"))
        tc_pr.insert(fill_index + offset, ln)
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, line in enumerate(str(text).splitlines() or [""]):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = 1.0
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        _set_fonts(run, dc.style.fonts, heading)
        _text_color(run, color)


def _table_layout(dc, header, rows, widths, h, where):
    """Pick the largest font size at which every row fits in ``h``."""
    base = dc.style.sizes["body"]
    for size in range(int(base), int(dc.style.sizes["min"]) - 1, -1):
        head_h = max(_h(t, cw - 0.08, size) for t, cw in zip(header, widths))
        row_hs = []
        for row in rows:
            row_hs.append(max(
                _h(_cell(v)[0], cw - 0.08, size) for v, cw in
                zip(row, widths)))
        if head_h + sum(row_hs) <= h:
            spare = h - head_h - sum(row_hs)
            pad = min(0.18, spare / (len(rows) + 1))
            return size, head_h + pad, [r + pad for r in row_hs]
    fail(f"{where}: table does not fit even at {dc.style.sizes['min']}pt; "
         "shorten cells or split rows across slides")


def _draw_table(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    header, rows = slots["header"], slots["rows"]
    n = len(header)
    ratios = slots.get("col_widths") or [1] * n
    widths = [w * r / sum(ratios) for r in ratios]
    emph_row = slots.get("emphasis_row")
    emph_col = slots.get("emphasis_col")
    size, head_h, row_hs = _table_layout(dc, header, rows, widths, h,
                                         dc.where)
    shape = dc.slide.shapes.add_table(
        len(rows) + 1, n, Inches(x), Inches(y), Inches(w),
        Inches(head_h + sum(row_hs)))
    shape.name = "Library table"
    tbl = shape.table
    tbl_pr = tbl._tbl.tblPr
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is not None:
        tbl_pr.remove(style_id)
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    tbl.rows[0].height = Inches(head_h)
    for i, rh in enumerate(row_hs, 1):
        tbl.rows[i].height = Inches(rh)
    colors = dc.style.colors
    for j, text in enumerate(header):
        hi = emph_col == j + 1
        _set_cell(dc, tbl.cell(0, j), text, size,
                  colors["highlight"] if hi else colors["primary"],
                  colors["on_highlight"] if hi else colors["on_primary"],
                  bold=True, heading=True)
    for i, row in enumerate(rows, 1):
        striped = variant == "striped" and i % 2 == 0
        for j, value in enumerate(row):
            text, tone, bold = _cell(value)
            fill = colors["surface"] if striped else colors["background"]
            color = colors["text"]
            if emph_row == i or emph_col == j + 1:
                fill = _tint(colors["highlight"], 0.14)
                bold = True
            if tone == "positive":
                fill, bold = _tint(colors["primary"], 0.2), True
            elif tone == "caution":
                fill, bold = _tint(colors["highlight"], 0.25), True
            elif tone == "negative":
                fill, color, bold = (colors["highlight"],
                                     colors["on_highlight"], True)
            elif tone == "neutral":
                fill = colors["surface"]
            if slots.get("row_header") and j == 0:
                bold = True
            _set_cell(dc, tbl.cell(i, j), text, size, fill, color,
                      bold=bold, bottom=colors["line"])
    mark_el(shape._element, GENERATED_MARK)
    _finish_frame(dc, y + head_h + sum(row_hs))


def _check_table(slots, variant, where):
    n = len(slots["header"])
    for i, row in enumerate(slots["rows"], 1):
        if len(row) != n:
            fail(f"{where}, rows[{i}]: expected {n} cells to match header")
    if "col_widths" in slots and len(slots["col_widths"]) != n:
        fail(f"{where}, col_widths: expected {n} values")
    if slots.get("emphasis_row") and slots["emphasis_row"] > len(
            slots["rows"]):
        fail(f"{where}, emphasis_row: expected 1..{len(slots['rows'])}")
    if slots.get("emphasis_col") and slots["emphasis_col"] > n:
        fail(f"{where}, emphasis_col: expected 1..{n}")


# ---------- scorecard (rating / RACI) ----------

def _draw_scorecard(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    columns, rows = slots["columns"], slots["rows"]
    recommend = set(slots.get("recommend", []))
    colors = dc.style.colors
    total = variant == "rating" and slots.get("show_total", True)
    n = len(columns)
    label_w = min(3.6, w * 0.34)
    cw = (w - label_w) / n
    legend_h = 0.46
    head_h = max(0.5, max(_h(c, cw - 0.1, dc.style.sizes["body"])
                          for c in columns) + 0.14)
    if recommend:
        head_h += 0.3
    body_rows = len(rows) + (1 if total else 0)
    row_h = min(0.66, (h - head_h - legend_h) / body_rows)
    label_size = _fit_common([r["label"] for r in rows], label_w - 0.1,
                             row_h - 0.04, dc.style.sizes["body"], dc,
                             f"{dc.where}, rows[].label")
    grid_bottom = y + head_h + body_rows * row_h
    for j, name in enumerate(columns, 1):
        cx = x + label_w + (j - 1) * cw
        rec = j in recommend
        if rec:
            dc.shape(f"Scorecard column {j} band", cx + 0.03, y, cw - 0.06,
                     grid_bottom - y, fill=_tint(colors["highlight"], 0.12),
                     line=colors["highlight"], radius=True)
            _chip(dc, f"Scorecard column {j} badge", cx + cw / 2 - 0.42,
                  y + 0.06, 0.84, 0.26, "推奨", colors["highlight"],
                  colors["on_highlight"])
        dc.textbox(f"Scorecard column {j}", cx, y + (0.3 if recommend else 0),
                   cw, head_h - (0.3 if recommend else 0), name,
                   dc.style.sizes["body"],
                   colors["highlight"] if rec else colors["primary"],
                   bold=True, heading=True, align=PP_ALIGN.CENTER,
                   anchor=MSO_ANCHOR.MIDDLE)
    _line(dc, "Scorecard header rule", x, y + head_h, x + w, y + head_h,
          colors["primary"], 1.5)
    scores = [0] * n
    for i, row in enumerate(rows, 1):
        ry = y + head_h + (i - 1) * row_h
        els = [dc.textbox(f"Scorecard row {i}", x, ry, label_w, row_h,
                          row["label"], label_size, colors["text"],
                          bold=True, anchor=MSO_ANCHOR.MIDDLE)]
        for j, value in enumerate(row["cells"], 1):
            cx = x + label_w + (j - 1) * cw + cw / 2
            if variant == "rating":
                score = RATING_SCORES[value]
                if score is not None:
                    scores[j - 1] += score
                tone = (colors["highlight"] if value == "×"
                        else colors["muted"] if value in ("△", "-", "")
                        else colors["primary"])
                els.append(dc.textbox(
                    f"Scorecard cell {i}-{j}", cx - cw / 2, ry, cw, row_h,
                    value or "－", min(24, max(label_size + 4, 16)), tone,
                    bold=True, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE))
            elif value:
                roles = value.split("/")
                chip_w = 0.42 if len(roles) == 1 else 0.7
                chip_h = min(0.42, row_h - 0.12)
                accountable = "A" in roles
                responsible = "R" in roles
                fill = (colors["highlight"] if accountable
                        else colors["primary"] if responsible else None)
                text_color = (colors["on_highlight"] if accountable
                              else colors["on_primary"] if responsible
                              else colors["primary"] if "C" in roles
                              else colors["muted"])
                line = (None if fill is not None
                        else colors["primary"] if "C" in roles else None)
                els.append(_chip(
                    dc, f"Scorecard cell {i}-{j}", cx - chip_w / 2,
                    ry + (row_h - chip_h) / 2, chip_w, chip_h, value,
                    fill, text_color, line=line))
        _group(dc, els, "Scorecard row", i)
        _line(dc, f"Scorecard rule {i}", x, ry + row_h, x + w, ry + row_h,
              colors["line"])
    if total:
        ry = y + head_h + len(rows) * row_h
        dc.textbox("Scorecard total label", x, ry, label_w, row_h,
                   "合計（◎3・○2・△1・×0）", dc.style.sizes["caption"],
                   colors["text"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
        best = max(scores)
        for j, score in enumerate(scores, 1):
            cx = x + label_w + (j - 1) * cw
            dc.textbox(f"Scorecard total {j}", cx, ry, cw, row_h,
                       str(score), dc.style.sizes["heading"],
                       colors["highlight"] if score == best
                       else colors["text"], bold=True, heading=True,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    legend = (RACI_LEGEND if variant == "raci"
              else "◎ 非常に適合　○ 適合　△ 一部適合　× 不適合")
    dc.textbox("Scorecard legend", x, grid_bottom + 0.04, w, legend_h - 0.04,
               legend, dc.style.sizes["caption"], colors["muted"],
               align=PP_ALIGN.RIGHT)
    _finish_frame(dc, grid_bottom + legend_h)


def _check_scorecard(slots, variant, where):
    n = len(slots["columns"])
    for i, row in enumerate(slots["rows"], 1):
        if len(row["cells"]) != n:
            fail(f"{where}, rows[{i}].cells: expected {n} cells "
                 "to match columns")
        for j, value in enumerate(row["cells"], 1):
            if variant == "rating" and value not in RATING_SCORES:
                fail(f"{where}, rows[{i}].cells[{j}]: rating cells must "
                     f"be one of {[k for k in RATING_SCORES if k]} or ''")
            if variant == "raci" and value and not all(
                    role in RACI_ROLES for role in value.split("/")):
                fail(f"{where}, rows[{i}].cells[{j}]: RACI cells must be "
                     "R/A/C/I or a combination such as 'A/R'")
    for index in slots.get("recommend", []):
        if not 1 <= index <= n:
            fail(f"{where}, recommend: expected numbers 1..{n}")
    if variant == "raci":
        for i, row in enumerate(slots["rows"], 1):
            if sum("A" in c.split("/") for c in row["cells"] if c) != 1:
                fail(f"{where}, rows[{i}]: RACI rows need exactly one 'A'")


# ---------- timeline ----------

def _draw_timeline(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    events = slots["events"]
    colors = dc.style.colors
    n = len(events)
    if variant == "vertical":
        date_w, axis_x = 1.7, x + 1.95
        body_x, body_w = axis_x + 0.35, x + w - axis_x - 0.35
        row_h = h / n
        label_size = _fit_common([e["label"] for e in events], body_w,
                                 row_h * 0.5, dc.style.sizes["body"], dc,
                                 f"{dc.where}, events[].label")
        detail_size = _fit_common([e.get("detail") for e in events], body_w,
                                  row_h * 0.5, dc.style.sizes["body"], dc,
                                  f"{dc.where}, events[].detail")
        need = max(_h(e["label"], body_w, label_size)
                   + (_h(e["detail"], body_w, detail_size)
                      if e.get("detail") else 0) for e in events) + 0.1
        row_h = min(row_h, max(0.75, need))
        _line(dc, "Timeline axis", axis_x, y + row_h / 2, axis_x,
              y + (n - 0.5) * row_h, colors["line"], 2)
        for i, event in enumerate(events, 1):
            ry = y + (i - 1) * row_h
            emph = event.get("emphasis")
            d = 0.3 if emph else 0.22
            els = [dc.textbox(f"Timeline date {i}", x, ry, date_w, row_h,
                              event["date"], dc.style.sizes["body"],
                              colors["highlight"] if emph
                              else colors["primary"], bold=True,
                              heading=True, align=PP_ALIGN.RIGHT,
                              anchor=MSO_ANCHOR.MIDDLE),
                   dc.shape(f"Timeline dot {i}", axis_x - d / 2,
                            ry + row_h / 2 - d / 2, d, d,
                            kind=MSO_SHAPE.OVAL,
                            fill=colors["highlight"] if emph
                            else colors["primary"],
                            line=colors["background"])]
            label_h = _h(event["label"], body_w, label_size)
            detail_h = (_h(event["detail"], body_w, detail_size)
                        if event.get("detail") else 0)
            top = ry + (row_h - label_h - detail_h) / 2
            els.append(dc.textbox(f"Timeline label {i}", body_x, top,
                                  body_w, label_h, event["label"],
                                  label_size, colors["text"], bold=True,
                                  heading=True))
            if event.get("detail"):
                els.append(dc.textbox(f"Timeline detail {i}", body_x,
                                      top + label_h, body_w, detail_h,
                                      event["detail"], detail_size,
                                      colors["muted"]))
            _group(dc, els, "Timeline", i)
        _finish_frame(dc, y + n * row_h)
        return
    col = w / n
    date_h = 0.46
    axis_y = y + date_h + 0.2
    label_size = _fit_common([e["label"] for e in events], col - 0.12, 0.8,
                             dc.style.sizes["body"], dc,
                             f"{dc.where}, events[].label")
    label_h = max(_h(e["label"], col - 0.12, label_size) for e in events)
    detail_room = max(0.4, h - (axis_y - y) - 0.35 - label_h)
    detail_size = _fit_common([e.get("detail") for e in events],
                              col - 0.12, detail_room,
                              dc.style.sizes["body"], dc,
                              f"{dc.where}, events[].detail")
    detail_h = max((_h(e["detail"], col - 0.12, detail_size)
                    if e.get("detail") else 0) for e in events)
    _arrow(dc, "Timeline axis", x, axis_y, x + w, axis_y, colors["line"], 2)
    for i, event in enumerate(events, 1):
        cx = x + (i - 0.5) * col
        emph = event.get("emphasis")
        d = 0.34 if emph else 0.26
        accent = colors["highlight"] if emph else colors["primary"]
        els = [dc.textbox(f"Timeline date {i}", cx - col / 2, y, col,
                          date_h, event["date"], dc.style.sizes["body"],
                          accent, bold=True, heading=True,
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM),
               dc.shape(f"Timeline dot {i}", cx - d / 2, axis_y - d / 2, d,
                        d, kind=MSO_SHAPE.OVAL, fill=accent,
                        line=colors["background"]),
               dc.textbox(f"Timeline label {i}", cx - col / 2 + 0.06,
                          axis_y + 0.3, col - 0.12, label_h, event["label"],
                          label_size, colors["text"], bold=True,
                          heading=True, align=PP_ALIGN.CENTER)]
        if event.get("detail"):
            els.append(dc.textbox(
                f"Timeline detail {i}", cx - col / 2 + 0.06,
                axis_y + 0.3 + label_h, col - 0.12, detail_h,
                event["detail"], detail_size, colors["muted"],
                align=PP_ALIGN.CENTER))
        _group(dc, els, "Timeline", i)
    _finish_frame(dc, axis_y + 0.3 + label_h + detail_h)


# ---------- logic tree ----------

def _tree_leaves(branches):
    return sum(len(b["children"]) for b in branches)


def _draw_tree(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    colors = dc.style.colors
    branches = slots["branches"]
    headers = slots.get("headers")
    head_h = 0.42 if headers else 0
    gap = 0.5
    root_w = w * 0.19
    branch_w = w * 0.29
    leaf_w = w - root_w - branch_w - 2 * gap
    leaves = _tree_leaves(branches)
    branch_gap = 0.16
    leaf_gap = 0.08
    avail = h - head_h - branch_gap * (len(branches) - 1) - leaf_gap * (
        leaves - len(branches))
    leaf_h = min(0.72, avail / leaves)
    has_tags = any(c.get("tag") for b in branches for c in b["children"])
    tag_w = 0.56 if has_tags else 0
    leaf_text_w = leaf_w - tag_w - (0.1 if has_tags else 0)
    leaf_size = _fit_common(
        [c["label"] for b in branches for c in b["children"]], leaf_text_w,
        leaf_h - 0.02, dc.style.sizes["body"], dc,
        f"{dc.where}, branches[].children[].label")
    branch_size = _fit_common([b["label"] for b in branches], branch_w,
                              max(leaf_h, 0.6), dc.style.sizes["body"], dc,
                              f"{dc.where}, branches[].label")
    if headers:
        for k, (hx, hw) in enumerate(((x, root_w),
                                      (x + root_w + gap, branch_w),
                                      (x + root_w + branch_w + 2 * gap,
                                       leaf_w)), 1):
            dc.textbox(f"Tree header {k}", hx, y, hw, head_h, headers[k - 1],
                       dc.style.sizes["caption"], colors["muted"], bold=True)
    top = y + head_h
    cursor = top
    branch_mid = []
    bx = x + root_w + gap
    lx = bx + branch_w + gap
    for i, branch in enumerate(branches, 1):
        kids = branch["children"]
        span = len(kids) * leaf_h + (len(kids) - 1) * leaf_gap
        mids = []
        els = []
        for j, child in enumerate(kids, 1):
            ly = cursor + (j - 1) * (leaf_h + leaf_gap)
            mids.append(ly + leaf_h / 2)
            els.append(dc.shape(f"Tree leaf {i}-{j}", lx, ly, leaf_w, leaf_h,
                                fill=colors["surface"], line=colors["line"],
                                radius=True))
            els.append(dc.textbox(f"Tree leaf {i}-{j} text", lx + 0.02, ly,
                                  leaf_text_w, leaf_h, child["label"],
                                  leaf_size, colors["text"],
                                  anchor=MSO_ANCHOR.MIDDLE))
            if child.get("tag"):
                els.append(_chip(
                    dc, f"Tree leaf {i}-{j} tag", lx + leaf_w - tag_w - 0.08,
                    ly + (leaf_h - 0.32) / 2, tag_w, 0.32, child["tag"],
                    colors["primary"], colors["on_primary"]))
        branch_h = max(min(span, 1.0), leaf_h)
        by = cursor + (span - branch_h) / 2
        els.append(dc.shape(f"Tree branch {i}", bx, by, branch_w, branch_h,
                            fill=colors["background"],
                            line=colors["primary"], radius=True,
                            text=branch["label"], size=branch_size,
                            color=colors["primary"], bold=True,
                            heading=True, anchor=MSO_ANCHOR.MIDDLE))
        els[-1].line.width = Pt(1.5)
        spine_x = bx + branch_w + gap / 2
        _line(dc, f"Tree stem {i}", bx + branch_w, by + branch_h / 2,
              spine_x, by + branch_h / 2, colors["primary"], 1.25)
        if len(mids) > 1:
            _line(dc, f"Tree spine {i}", spine_x, mids[0], spine_x,
                  mids[-1], colors["primary"], 1.25)
        for j, mid in enumerate(mids, 1):
            _line(dc, f"Tree twig {i}-{j}", spine_x, mid, lx, mid,
                  colors["primary"], 1.25)
        _group(dc, els, "Tree branch", i)
        branch_mid.append(by + branch_h / 2)
        cursor += span + branch_gap
    bottom = cursor - branch_gap
    root_h = min(bottom - top, max(1.1, _h(slots["root"], root_w,
                                           dc.style.sizes["heading"]) + 0.3))
    ry = top + (bottom - top - root_h) / 2
    dc.shape("Tree root", x, ry, root_w, root_h, fill=colors["primary"],
             line=None, text=slots["root"], size=dc.style.sizes["heading"],
             color=colors["on_primary"], bold=True, heading=True,
             radius=True, anchor=MSO_ANCHOR.MIDDLE)
    spine_x = x + root_w + gap / 2
    _line(dc, "Tree root stem", x + root_w, ry + root_h / 2, spine_x,
          ry + root_h / 2, colors["primary"], 1.5)
    _line(dc, "Tree root spine", spine_x, min(branch_mid + [ry + root_h / 2]),
          spine_x, max(branch_mid + [ry + root_h / 2]), colors["primary"], 1.5)
    for i, mid in enumerate(branch_mid, 1):
        _line(dc, f"Tree root twig {i}", spine_x, mid, bx, mid,
              colors["primary"], 1.5)
    _finish_frame(dc, bottom)


def _check_tree(slots, variant, where):
    leaves = _tree_leaves(slots["branches"])
    if leaves > 9:
        fail(f"{where}, branches: {leaves} leaves exceed 9; "
             "split the content across slides")
    headers = slots.get("headers")
    if headers is not None and len(headers) != 3:
        fail(f"{where}, headers: expected 3 labels (root, branch, leaf)")


# ---------- hub ----------

def _spoke_box(dc, name, bx, by, bw, bh, spoke, label_size, detail_size,
               accent):
    colors = dc.style.colors
    els = [dc.shape(f"{name} box", bx, by, bw, bh, fill=colors["surface"],
                    line=accent, radius=True)]
    label_h = _h(spoke["label"], bw - 0.1, label_size)
    detail_h = (_h(spoke["detail"], bw - 0.1, detail_size)
                if spoke.get("detail") else 0)
    top = by + max(0.02, (bh - label_h - detail_h) / 2)
    els.append(dc.textbox(f"{name} label", bx + 0.05, top, bw - 0.1,
                          label_h, spoke["label"], label_size,
                          colors["text"], bold=True, heading=True))
    if spoke.get("detail"):
        els.append(dc.textbox(f"{name} detail", bx + 0.05, top + label_h,
                              bw - 0.1, detail_h, spoke["detail"],
                              detail_size, colors["muted"]))
    return els


def _draw_hub(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    colors = dc.style.colors
    if variant == "flow":
        sides = [slots["left"], slots["right"]]
    else:
        spokes = slots["spokes"]
        half = math.ceil(len(spokes) / 2)
        sides = [{"items": spokes[:half]}, {"items": spokes[half:]}]
    title_h = 0.4 if any(side.get("title") for side in sides) else 0
    diameter = min(2.3, h - title_h - 0.2, w * 0.22)
    gap = 0.55
    box_w = (w - diameter - 2 * gap) / 2
    all_items = sides[0]["items"] + sides[1]["items"]
    k = max(len(side["items"]) for side in sides)
    box_gap = 0.16
    box_h = min(1.15, (h - title_h - box_gap * (k - 1)) / k)
    label_size = _fit_common([s["label"] for s in all_items], box_w - 0.1,
                             box_h * 0.5, dc.style.sizes["body"], dc,
                             f"{dc.where}, items[].label")
    detail_size = min(label_size, _fit_common(
        [s.get("detail") for s in all_items], box_w - 0.1, box_h * 0.5,
        dc.style.sizes["body"], dc, f"{dc.where}, items[].detail"))
    content_top = y + title_h
    block_h = k * box_h + (k - 1) * box_gap
    cx, cy = x + w / 2, content_top + block_h / 2
    radius = diameter / 2
    for s, side in enumerate(sides):
        items = side["items"]
        if not items:
            continue
        bx = x if s == 0 else x + w - box_w
        side_h = len(items) * box_h + (len(items) - 1) * box_gap
        top = content_top + (block_h - side_h) / 2
        if side.get("title"):
            dc.textbox(f"Hub side {s + 1} title", bx, y, box_w, title_h,
                       side["title"], dc.style.sizes["caption"],
                       colors["muted"], bold=True,
                       align=PP_ALIGN.LEFT if s == 0 else PP_ALIGN.RIGHT)
        for i, spoke in enumerate(items, 1):
            by = top + (i - 1) * (box_h + box_gap)
            name = f"Hub {'left' if s == 0 else 'right'} {i}"
            accent = colors["primary"]
            els = _spoke_box(dc, name, bx, by, box_w, box_h, spoke,
                             label_size, detail_size, accent)
            _group(dc, els, name, 0)
            ex = bx + box_w if s == 0 else bx
            ey = by + box_h / 2
            dx, dy = cx - ex, cy - ey
            dist = math.hypot(dx, dy) or 1
            px, py = cx - dx / dist * (radius + 0.06), \
                cy - dy / dist * (radius + 0.06)
            if variant == "flow":
                if s == 0:
                    _arrow(dc, f"{name} link", ex + 0.04, ey, px, py,
                           colors["primary"], 1.25)
                else:
                    _arrow(dc, f"{name} link", px, py, ex - 0.04, ey,
                           colors["primary"], 1.25)
            else:
                _line(dc, f"{name} link", ex, ey, px, py, colors["primary"],
                      1.25)
    dc.shape("Hub center ring", cx - radius - 0.08, cy - radius - 0.08,
             diameter + 0.16, diameter + 0.16, kind=MSO_SHAPE.OVAL,
             fill=_tint(colors["primary"], 0.2), line=None)
    dc.shape("Hub center", cx - radius, cy - radius, diameter, diameter,
             kind=MSO_SHAPE.OVAL, fill=colors["primary"], line=None)
    inner = diameter * 0.78
    # Renderers set CJK wider than the estimate; keep a 12% safety margin so
    # a word does not break inside the circle.
    center_size = min(
        fit_size(slots["center"], inner * 0.88, inner,
                 dc.style.sizes["heading"], dc.style.sizes["min"],
                 f"{dc.where}, center"),
        _single_line_size(str(slots["center"]).splitlines(), inner * 0.9,
                          dc.style.sizes["heading"], dc.style.sizes["min"],
                          pad=0.04))
    dc.textbox("Hub center text", cx - inner / 2, cy - inner / 2, inner,
               inner, slots["center"], center_size,
               colors["on_primary"], bold=True, heading=True,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    _finish_frame(dc, content_top + block_h)


def _check_hub(slots, variant, where):
    if variant == "flow":
        for key in ("left", "right"):
            if key not in slots:
                fail(f"{where}: flow variant requires '{key}'")
        if "spokes" in slots:
            fail(f"{where}: flow variant uses left/right, not spokes")
    else:
        if "spokes" not in slots:
            fail(f"{where}: hub variant requires 'spokes'")
        if "left" in slots or "right" in slots:
            fail(f"{where}: hub variant uses spokes, not left/right")


# ---------- swimlane ----------

def _draw_swimlane(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    colors = dc.style.colors
    lanes, steps = slots["lanes"], slots["steps"]
    cols = [s.get("col", i) for i, s in enumerate(steps, 1)]
    n_cols = max(cols)
    lane_w = min(1.7, w * 0.15)
    col_w = (w - lane_w) / n_cols
    lane_h = min(1.45, h / len(lanes))
    box_w = col_w - 0.28
    box_h = lane_h - 0.24
    emph = slots.get("emphasis")
    label_size = _fit_common([s["label"] for s in steps], box_w,
                             box_h - 0.12, dc.style.sizes["body"], dc,
                             f"{dc.where}, steps[].label")
    lane_size = _fit_common(lanes, lane_w - 0.1, lane_h - 0.1,
                            dc.style.sizes["body"], dc,
                            f"{dc.where}, lanes[]")
    for i, lane in enumerate(lanes, 1):
        ly = y + (i - 1) * lane_h
        dc.shape(f"Swimlane band {i}", x, ly, w, lane_h,
                 fill=colors["surface"] if i % 2 else colors["background"],
                 line=None)
        dc.shape(f"Swimlane lane {i}", x, ly, lane_w, lane_h,
                 fill=colors["primary"], line=colors["background"],
                 text=lane, size=lane_size, color=colors["on_primary"],
                 bold=True, heading=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
    _line(dc, "Swimlane bottom rule", x, y + len(lanes) * lane_h, x + w,
          y + len(lanes) * lane_h, colors["line"])
    centers = []
    for k, (step, col) in enumerate(zip(steps, cols), 1):
        bx = x + lane_w + (col - 1) * col_w + 0.14
        by = y + (step["lane"] - 1) * lane_h + 0.14
        hi = emph == k
        shape = dc.shape(f"Swimlane step {k}", bx, by, box_w, box_h,
                         fill=colors["background"],
                         line=colors["highlight"] if hi
                         else colors["primary"], radius=True,
                         text=step["label"], size=label_size,
                         color=colors["text"], bold=hi,
                         anchor=MSO_ANCHOR.MIDDLE)
        shape.line.width = Pt(2 if hi else 1.25)
        shape.text_frame.margin_top = Inches(0.16)
        chip = _chip(dc, f"Swimlane step {k} number", bx - 0.1, by - 0.13,
                     0.32, 0.32, str(k),
                     colors["highlight"] if hi else colors["primary"],
                     colors["on_highlight"] if hi else colors["on_primary"])
        chip.text_frame.paragraphs[0].runs[0].font.size = Pt(
            dc.style.sizes["min"])
        _group(dc, [shape, chip], "Swimlane step", k)
        centers.append((bx, by, col, step["lane"]))
    for k in range(len(steps) - 1):
        bx, by, col, lane = centers[k]
        nx, ny, ncol, nlane = centers[k + 1]
        if ncol == col:
            down = nlane > lane
            _arrow(dc, f"Swimlane flow {k + 1}", bx + box_w / 2,
                   by + (box_h if down else 0), nx + box_w / 2,
                   ny + (0 if down else box_h))
        else:
            _arrow(dc, f"Swimlane flow {k + 1}", bx + box_w, by + box_h / 2,
                   nx, ny + box_h / 2, elbow=nlane != lane)
    _finish_frame(dc, y + len(lanes) * lane_h)


def _check_swimlane(slots, variant, where):
    lanes = len(slots["lanes"])
    seen = set()
    for k, step in enumerate(slots["steps"], 1):
        if step["lane"] > lanes:
            fail(f"{where}, steps[{k}].lane: expected 1..{lanes}")
        key = (step["lane"], step.get("col", k))
        if key in seen:
            fail(f"{where}, steps[{k}]: two steps share lane {key[0]} "
                 f"and column {key[1]}")
        seen.add(key)
    cols = [s.get("col", i) for i, s in enumerate(slots["steps"], 1)]
    if any(b < a for a, b in zip(cols, cols[1:])):
        fail(f"{where}, steps[].col: columns must not decrease")
    if max(cols) > 8:
        fail(f"{where}, steps[].col: at most 8 columns")
    emph = slots.get("emphasis")
    if emph is not None and emph > len(slots["steps"]):
        fail(f"{where}, emphasis: expected 1..{len(slots['steps'])}")


# ---------- action plan ----------

_ACTION_LABELS = {"issue": "課題", "action": "アクション", "owner": "担当",
                  "due": "期限", "done_when": "完了条件"}


def _draw_action_plan(dc, region, slots, variant):
    x, y, w, h = _draw_frame(dc, region, slots)
    colors = dc.style.colors
    rows = slots["rows"]
    labels = dict(_ACTION_LABELS)
    labels.update(slots.get("labels", {}))
    has_done = any(r.get("done_when") for r in rows)
    has_id = any(r.get("id") for r in rows)
    arrow_w = 0.36
    ratios = ([0.23, 0.31, 0.15, 0.08, 0.23] if has_done
              else [0.31, 0.41, 0.17, 0.11])
    usable = w - arrow_w - 0.08 * (len(ratios) - 1)
    widths = [usable * r for r in ratios]
    xs = [x]
    for k, cw in enumerate(widths[:-1]):
        xs.append(xs[-1] + cw + 0.08 + (arrow_w if k == 0 else 0))
    head_h = 0.42
    id_w = 0.5 if has_id else 0
    row_gap = 0.1
    row_h = min(1.0, (h - head_h - row_gap * (len(rows) - 1)) / len(rows))
    keys = ["issue", "action", "owner", "due"] + (
        ["done_when"] if has_done else [])
    sizes = {}
    for key, cw in zip(keys, widths):
        inner = cw - (id_w + 0.04 if key == "action" else 0) - 0.04
        sizes[key] = _fit_common([r.get(key) for r in rows], inner,
                                 row_h - 0.06, dc.style.sizes["body"], dc,
                                 f"{dc.where}, rows[].{key}")
    common = min(sizes[k] for k in keys if k not in ("owner", "due"))
    for key, cw in zip(keys, widths):
        if key in ("owner", "due"):
            # short labels read badly when wrapped: shrink to one line
            sizes[key] = min(common, _single_line_size(
                [r.get(key, "") for r in rows], cw, common,
                dc.style.sizes["min"]))
    for key, hx, cw in zip(keys, xs, widths):
        dc.textbox(f"Action header {key}", hx, y, cw, head_h, labels[key],
                   dc.style.sizes["caption"], colors["muted"], bold=True,
                   anchor=MSO_ANCHOR.BOTTOM)
    _line(dc, "Action header rule", x, y + head_h, x + w, y + head_h,
          colors["primary"], 1.5)
    emph = slots.get("emphasis")
    for i, row in enumerate(rows, 1):
        ry = y + head_h + 0.08 + (i - 1) * (row_h + row_gap)
        hi = emph == i
        accent = colors["highlight"] if hi else colors["primary"]
        els = [dc.shape(f"Action {i} issue", xs[0], ry, widths[0], row_h,
                        fill=colors["surface"], line=None, radius=True,
                        text=row["issue"], size=common,
                        color=colors["text"], anchor=MSO_ANCHOR.MIDDLE)]
        els.append(dc.shape(f"Action {i} arrow",
                            xs[0] + widths[0] + 0.06,
                            ry + row_h / 2 - 0.13, arrow_w - 0.04, 0.26,
                            kind=MSO_SHAPE.RIGHT_ARROW, fill=accent,
                            line=None))
        box = dc.shape(f"Action {i} box", xs[1], ry, widths[1], row_h,
                       fill=colors["background"], line=accent, radius=True)
        box.line.width = Pt(2 if hi else 1.25)
        els.append(box)
        if row.get("id"):
            els.append(_chip(dc, f"Action {i} id", xs[1] + 0.08,
                             ry + (row_h - 0.34) / 2, id_w - 0.04, 0.34,
                             row["id"], accent,
                             colors["on_highlight" if hi else "on_primary"]))
        els.append(dc.textbox(f"Action {i} text", xs[1] + id_w + 0.04, ry,
                              widths[1] - id_w - 0.04, row_h, row["action"],
                              common, colors["text"], bold=True,
                              anchor=MSO_ANCHOR.MIDDLE))
        for key, hx, cw in list(zip(keys, xs, widths))[2:]:
            els.append(dc.textbox(f"Action {i} {key}", hx, ry, cw, row_h,
                                  row.get(key, ""),
                                  sizes[key] if key in ("owner", "due")
                                  else common,
                                  colors["text"],
                                  bold=key in ("owner", "due"),
                                  anchor=MSO_ANCHOR.MIDDLE))
        _group(dc, els, "Action", i)
        if i < len(rows):
            _line(dc, f"Action rule {i}", xs[2], ry + row_h + row_gap / 2,
                  x + w, ry + row_h + row_gap / 2, colors["line"])
    _finish_frame(dc, y + head_h + 0.08 + len(rows) * row_h
                  + (len(rows) - 1) * row_gap)


def _check_action_plan(slots, variant, where):
    emph = slots.get("emphasis")
    if emph is not None and emph > len(slots["rows"]):
        fail(f"{where}, emphasis: expected 1..{len(slots['rows'])}")
    if any(r.get("done_when") for r in slots["rows"]) and not all(
            r.get("done_when") for r in slots["rows"]):
        fail(f"{where}, rows[].done_when: give every row a completion "
             "condition or none of them")


# ---------- schemas ----------

def _library():
    text = {"type": "string", "minLength": 1}
    cell = {"oneOf": [
        {"type": "string"},
        {"type": "object", "required": ["text"],
         "properties": {"text": {"type": "string"},
                        "tone": {"enum": list(TONES)},
                        "bold": {"type": "boolean"}},
         "additionalProperties": False},
    ]}
    labelled = {"type": "object", "required": ["label"],
                "properties": {"label": text, "detail": {"type": "string"}},
                "additionalProperties": False}

    def side(minimum, maximum):
        return {"type": "object", "required": ["items"],
                "properties": {
                    "title": {"type": "string"},
                    "items": {"type": "array", "minItems": minimum,
                              "maxItems": maximum, "items": labelled}},
                "additionalProperties": False}

    def comp(variants, properties, required, draw, check):
        return Component(variants, _component_schema(properties, required),
                         draw, check)

    return {
        "table": comp(
            ("grid", "striped"),
            {"header": {"type": "array", "minItems": 2, "maxItems": 7,
                        "items": text},
             "rows": {"type": "array", "minItems": 1, "maxItems": 10,
                      "items": {"type": "array", "items": cell}},
             "col_widths": {"type": "array",
                            "items": {"type": "number",
                                      "exclusiveMinimum": 0}},
             "row_header": {"type": "boolean"},
             "emphasis_row": {"type": "integer", "minimum": 1},
             "emphasis_col": {"type": "integer", "minimum": 1}},
            ("header", "rows"), _draw_table, _check_table),
        "scorecard": comp(
            ("rating", "raci"),
            {"columns": {"type": "array", "minItems": 2, "maxItems": 6,
                         "items": text},
             "rows": {"type": "array", "minItems": 2, "maxItems": 8,
                      "items": {"type": "object",
                                "required": ["label", "cells"],
                                "properties": {
                                    "label": text,
                                    "cells": {"type": "array",
                                              "items": {"type": "string"}}},
                                "additionalProperties": False}},
             "recommend": {"type": "array", "maxItems": 3,
                           "items": {"type": "integer", "minimum": 1}},
             "show_total": {"type": "boolean"}},
            ("columns", "rows"), _draw_scorecard, _check_scorecard),
        "timeline": comp(
            ("horizontal", "vertical"),
            {"events": {"type": "array", "minItems": 2, "maxItems": 7,
                        "items": {"type": "object",
                                  "required": ["date", "label"],
                                  "properties": {
                                      "date": text, "label": text,
                                      "detail": {"type": "string"},
                                      "emphasis": {"type": "boolean"}},
                                  "additionalProperties": False}}},
            ("events",), _draw_timeline, None),
        "tree": comp(
            ("logic",),
            {"root": text,
             "headers": {"type": "array", "items": {"type": "string"}},
             "branches": {"type": "array", "minItems": 2, "maxItems": 4,
                          "items": {"type": "object",
                                    "required": ["label", "children"],
                                    "properties": {
                                        "label": text,
                                        "children": {
                                            "type": "array", "minItems": 1,
                                            "maxItems": 3, "items": {
                                                "type": "object",
                                                "required": ["label"],
                                                "properties": {
                                                    "label": text,
                                                    "tag": {"type": "string",
                                                            "maxLength": 4}},
                                                "additionalProperties":
                                                    False}}},
                                    "additionalProperties": False}}},
            ("root", "branches"), _draw_tree, _check_tree),
        "hub": comp(
            ("hub", "flow"),
            {"center": text,
             "spokes": {"type": "array", "minItems": 3, "maxItems": 8,
                        "items": labelled},
             "left": side(1, 4), "right": side(1, 4)},
            ("center",), _draw_hub, _check_hub),
        "swimlane": comp(
            ("swimlane",),
            {"lanes": {"type": "array", "minItems": 2, "maxItems": 5,
                       "items": text},
             "steps": {"type": "array", "minItems": 2, "maxItems": 8,
                       "items": {"type": "object",
                                 "required": ["lane", "label"],
                                 "properties": {
                                     "lane": {"type": "integer",
                                              "minimum": 1},
                                     "col": {"type": "integer",
                                             "minimum": 1},
                                     "label": text},
                                 "additionalProperties": False}},
             "emphasis": {"type": "integer", "minimum": 1}},
            ("lanes", "steps"), _draw_swimlane, _check_swimlane),
        "action_plan": comp(
            ("rows",),
            {"rows": {"type": "array", "minItems": 1, "maxItems": 6,
                      "items": {"type": "object",
                                "required": ["issue", "action", "owner",
                                             "due"],
                                "properties": {
                                    "id": {"type": "string", "maxLength": 4},
                                    "issue": text, "action": text,
                                    "owner": text, "due": text,
                                    "done_when": {"type": "string"}},
                                "additionalProperties": False}},
             "labels": {"type": "object",
                        "properties": {k: {"type": "string"}
                                       for k in _ACTION_LABELS},
                        "additionalProperties": False},
             "emphasis": {"type": "integer", "minimum": 1}},
            ("rows",), _draw_action_plan, _check_action_plan),
    }
