#!/usr/bin/env python3
"""QA a built .pptx deck.

Usage: qa.py out.pptx [--deck deck.json --map map.json]
            [--render DIR] [--report report.json]
Exit 0 when no errors, 1 otherwise.
"""
import argparse
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build import absolute_bbox  # noqa: E402
from common import (FALLBACK_MARK, GENERATED_MARK, LINE_HEIGHT,
                    char_units)  # noqa: E402
from components import resolve_component  # noqa: E402
from engine import render_pngs  # noqa: E402

EMU_PER_IN = 914400
GENERIC_MARKERS = ["Click to add", "クリックして", "Lorem", "TODO", "[insert"]
SLIDE_BOTTOM_MARGIN = Emu(int(0.25 * EMU_PER_IN))
OVERLAP_MIN_RATIO = 0.05  # warn when >5% of the drawn shape is covered

DRAWN_MARKS = (FALLBACK_MARK, GENERATED_MARK)
DRAWN_SHAPE_TAGS = {qn("p:sp"), qn("p:grpSp"), qn("p:pic"),
                    qn("p:graphicFrame")}
OVERFLOW_TOLERANCE = 1.05


# ---------- font size resolution ----------

def _sz_of_rpr(rpr):
    if rpr is not None and rpr.get("sz"):
        return int(rpr.get("sz")) / 100.0
    return None


def _para_default_sz(p_el):
    pPr = p_el.find(qn("a:pPr"))
    if pPr is not None:
        return _sz_of_rpr(pPr.find(qn("a:defRPr")))
    return None


def _lststyle_sz(sp_el):
    tx = sp_el.find(qn("p:txBody"))
    if tx is None:
        return None
    lst = tx.find(qn("a:lstStyle"))
    if lst is None:
        return None
    for lvl in ("a:lvl1pPr", "a:lvl2pPr", "a:lvl3pPr"):
        lp = lst.find(qn(lvl))
        if lp is not None:
            sz = _sz_of_rpr(lp.find(qn("a:defRPr")))
            if sz:
                return sz
    return None


def _inherited_placeholder_sz(slide, shape):
    if not shape.is_placeholder:
        return None
    idx = shape.placeholder_format.idx
    for container in (slide.slide_layout, slide.slide_layout.slide_master):
        for ph in container.placeholders:
            if ph.placeholder_format.idx == idx:
                sz = _lststyle_sz(ph._element)
                if sz:
                    return sz
    # master text styles (titleStyle / bodyStyle)
    master_el = slide.slide_layout.slide_master.element
    txstyles = master_el.find(qn("p:txStyles"))
    if txstyles is not None:
        styles = (qn("p:titleStyle") if idx == 0 else qn("p:bodyStyle"))
        st = txstyles.find(styles)
        if st is not None:
            lp = st.find(qn("a:lvl1pPr"))
            if lp is not None:
                return _sz_of_rpr(lp.find(qn("a:defRPr")))
    return None


def run_font_size(run, p_el, shape, slide):
    if run is not None:
        rpr = run.find(qn("a:rPr"))
        sz = _sz_of_rpr(rpr)
        if sz:
            return sz
    sz = _para_default_sz(p_el)
    if sz:
        return sz
    epr = p_el.find(qn("a:endParaRPr"))
    sz = _sz_of_rpr(epr)
    if sz:
        return sz
    sz = _lststyle_sz(shape._element)
    if sz:
        return sz
    sz = _inherited_placeholder_sz(slide, shape)
    if sz:
        return sz
    return 18.0


# ---------- overflow estimation ----------

def _insets(sp_el):
    tx = sp_el.find(qn("p:txBody"))
    if tx is None:
        tx = sp_el.find(qn("a:txBody"))
    if tx is None:
        return 0, 0, 0, 0
    body = tx.find(qn("a:bodyPr"))
    if body is None:
        return 0, 0, 0, 0
    g = lambda k, d: int(body.get(k, d))
    return (g("lIns", 91440), g("rIns", 91440), g("tIns", 45720),
            g("bIns", 45720))


def text_height_emu(shape, slide, box_w=None):
    """Estimated height needed by a text shape, in EMU."""
    el = shape._element
    tx = _txbody_el(el)
    if tx is None:
        return 0
    l_in, r_in, t_in, b_in = _insets(el)
    wrap = tx.find(qn("a:bodyPr"))
    no_wrap = wrap is not None and wrap.get("wrap") == "none"
    avail_w = max((box_w if box_w is not None else shape.width or 0)
                  - l_in - r_in, 1)
    total = 0.0
    for p in tx.findall(qn("a:p")):
        runs = p.findall(qn("a:r")) + p.findall(qn("a:fld"))
        text = "".join(
            (r.find(qn("a:t")).text or "") for r in runs
            if r.find(qn("a:t")) is not None)
        fs = 0.0
        for r in runs:
            s = run_font_size(r, p, shape, slide)
            fs = max(fs, s)
        if not fs:
            fs = run_font_size(None, p, shape, slide)
        line_h = fs * LINE_HEIGHT / 72.0 * EMU_PER_IN
        if not text:
            total += line_h
            continue
        if no_wrap:
            total += line_h
            continue
        units = sum(char_units(c) for c in text)
        char_w = fs / 72.0 * EMU_PER_IN
        lines = max(1, math.ceil(units * char_w / avail_w - 1e-9))
        total += lines * line_h
    return total + t_in + b_in


def _txbody_el(el):
    for tag in (qn("p:txBody"), qn("a:txBody")):
        tx = el.find(tag)
        if tx is not None:
            return tx
    return None


def _cell_insets(tc):
    l_in = int(tc.get("marL", 91440))
    r_in = int(tc.get("marR", 91440))
    t_in = int(tc.get("marT", 45720))
    b_in = int(tc.get("marB", 45720))
    return l_in, r_in, t_in, b_in


def cell_text_height(tc, col_w):
    tx = tc.find(qn("a:txBody"))
    l_in, r_in, t_in, b_in = _cell_insets(tc)
    avail_w = max(col_w - l_in - r_in, 1)
    total = 0.0
    for p in tx.findall(qn("a:p")):
        runs = p.findall(qn("a:r"))
        text = "".join((r.find(qn("a:t")).text or "")
                       for r in runs if r.find(qn("a:t")) is not None)
        fs = 0.0
        for r in runs:
            sz = _sz_of_rpr(r.find(qn("a:rPr")))
            if sz:
                fs = max(fs, sz)
        if not fs:
            sz = _para_default_sz(p)
            fs = sz if sz else 18.0
        line_h = fs * LINE_HEIGHT / 72.0 * EMU_PER_IN
        if not text:
            total += line_h
            continue
        units = sum(char_units(c) for c in text)
        char_w = fs / 72.0 * EMU_PER_IN
        lines = max(1, math.ceil(units * char_w / avail_w - 1e-9))
        total += lines * line_h
    return total + t_in + b_in


# ---------- checks ----------

def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # group
            yield from iter_shapes(shape.shapes)


def _drawn_mark(el):
    """descr marker on the element's own cNvPr, if any."""
    for child in el:
        for c in child.iter(qn("p:cNvPr")):
            return c.get("descr")
    return None


def _is_placeholder_el(el):
    for child in el:
        nv = child.find(qn("p:nvPr"))
        if nv is not None:
            return nv.find(qn("p:ph")) is not None
    return False


def _overlap_area(a, b):
    w = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    h = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0


def check_deck_qa(pptx_path, deck, tmap, render_dir, engine="auto"):
    errors = []
    warnings = []

    def err(slide, shape, msg):
        errors.append({"slide": slide, "shape": shape, "message": msg})

    def warn(slide, shape, msg):
        warnings.append({"slide": slide, "shape": shape,
                         "message": msg})

    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        return [{"slide": None, "shape": None,
                 "message": f"cannot reopen pptx: {e}"}], warnings
    sw, sh = prs.slide_width, prs.slide_height

    if deck is not None:
        if len(prs.slides) != len(deck["slides"]):
            err(None, None, f"slide count {len(prs.slides)} != deck slides "
                            f"{len(deck['slides'])}")

    markers = list(GENERIC_MARKERS)
    if tmap:
        markers += tmap.get("placeholder_markers", [])
    proto_names = set()
    if tmap:
        for comp in tmap.get("components", {}).values():
            for slot in comp.get("slots", {}).values():
                if slot.get("node_prototype"):
                    proto_names.add(slot["node_prototype"])
                elif slot.get("type") in ("orgchart", "timeline"):
                    proto_names.add("SynthNode")

    for s_idx, slide in enumerate(prs.slides, start=1):
        shapes = list(iter_shapes(slide.shapes))
        # 2. empty placeholders & leftover markers
        for shape in shapes:
            if shape.is_placeholder \
                    and shape._element.tag == qn("p:sp"):
                text = "".join(t.text or ""
                               for t in shape._element.iter(qn("a:t")))
                if not text.strip():
                    err(s_idx, shape.name,
                        "empty placeholder (unfilled layout slot)")
            if getattr(shape, "has_text_frame", False) \
                    and shape.has_text_frame:
                text = shape.text_frame.text
                for m in markers:
                    if m and m in text:
                        err(s_idx, shape.name,
                            f"placeholder/marker text remains: {m!r}")
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for m in markers:
                            if m and m in cell.text:
                                err(s_idx, shape.name,
                                    f"marker text in table cell: {m!r}")
        # geometry in absolute slide coordinates (groups transform)
        abs_boxes = {}
        for shape in shapes:
            ab = absolute_bbox(shape._element)
            if ab is not None:
                abs_boxes[id(shape)] = ab
        # 3. text overflow
        for shape in shapes:
            ab = abs_boxes.get(id(shape))
            if getattr(shape, "has_text_frame", False) \
                    and shape.has_text_frame and shape.text_frame.text.strip():
                need = text_height_emu(
                    shape, slide, box_w=ab[2] if ab else None)
                box_h = ab[3] if ab else shape.height
                if box_h and need > box_h * OVERFLOW_TOLERANCE:
                    err(s_idx, shape.name,
                        f"text overflow: needs ~{need / EMU_PER_IN:.2f}in, "
                        f"box is {box_h / EMU_PER_IN:.2f}in")
            if getattr(shape, "has_table", False) and shape.has_table:
                tbl = shape._element.graphic.graphicData.tbl
                col_w = [int(c.get("w")) for c in
                         tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))]
                est_rows = 0.0
                for tr in tbl.findall(qn("a:tr")):
                    row_h = 0.0
                    for tc, w in zip(tr.findall(qn("a:tc")), col_w):
                        row_h = max(row_h, cell_text_height(tc, w))
                    est_rows += row_h
                top = ab[1] if ab else (shape.top or 0)
                bottom = top + est_rows
                if bottom > sh - SLIDE_BOTTOM_MARGIN:
                    err(s_idx, shape.name,
                        f"table overflows slide bottom "
                        f"(est. bottom {bottom / EMU_PER_IN:.2f}in, "
                        f"margin { (sh - SLIDE_BOTTOM_MARGIN) / EMU_PER_IN:.2f}in)")
            # 3b. chart sanity: >=1 series and >=1 category
            if getattr(shape, "has_chart", False) and shape.has_chart:
                try:
                    series = list(shape.chart.plots[0].series)
                    cats = list(shape.chart.plots[0].categories)
                    if not series or not cats:
                        err(s_idx, shape.name,
                            "chart has no series or no categories")
                except Exception as e:
                    err(s_idx, shape.name, f"chart not readable: {e}")
        # 4a. shapes outside slide bounds
        for shape in shapes:
            ab = abs_boxes.get(id(shape))
            if ab is None:
                l, t, w, h = (shape.left, shape.top,
                              shape.width, shape.height)
                if l is None:
                    continue
            else:
                l, t, w, h = ab
            if l < -10000 or t < -10000 or l + w > sw + 10000 \
                    or t + h > sh + 10000:
                err(s_idx, shape.name,
                    f"shape outside slide bounds "
                    f"([{_in2(l)},{_in2(t)},{_in2(w)},{_in2(h)}])")
        # 4a2. shape ids must be unique positive integers
        seen_ids = {}
        tree = slide.shapes._spTree
        root_nv = tree.find(qn("p:nvGrpSpPr"))
        root_id = root_nv.find(qn("p:cNvPr")) if root_nv is not None \
            else None
        for e in tree.iter(qn("p:cNvPr")):
            if e is root_id:
                continue
            sid = e.get("id")
            name = e.get("name", "?")
            if sid is None or not sid.isdigit() or int(sid) <= 0:
                err(s_idx, name, f"invalid shape id {sid!r}")
                continue
            if sid in seen_ids:
                err(s_idx, name,
                    f"duplicate shape id {sid} (also on "
                    f"'{seen_ids[sid]}')")
            else:
                seen_ids[sid] = name
        # 4b. overlap between generated nodes (same prototype name)
        for name in proto_names:
            nodes = [s for s in shapes if s.name == name]
            for i in range(len(nodes)):
                for j in range(i + 1, len(nodes)):
                    a, b = nodes[i], nodes[j]
                    ba = abs_boxes.get(id(a))
                    bb = abs_boxes.get(id(b))
                    if ba and bb:
                        overlaps = _overlap_box(ba, bb)
                    else:
                        overlaps = _overlap(a, b)
                    if overlaps:
                        err(s_idx, name,
                            f"generated nodes overlap "
                            f"({_in2(a.left)},{_in2(a.top)} vs "
                            f"{_in2(b.left)},{_in2(b.top)})")
        # 4c. drawn (fallback/generated) shapes overlapping kept content
        tree_el = slide.shapes._spTree
        tagged = []
        tagged_ids = set()
        for el in tree_el.iter():
            if el.tag in DRAWN_SHAPE_TAGS \
                    and _drawn_mark(el) in DRAWN_MARKS:
                tagged.append(el)
                tagged_ids.add(id(el))

        def inside_tagged(el):
            p = el.getparent()
            while p is not None and p is not tree_el:
                if id(p) in tagged_ids:
                    return True
                p = p.getparent()
            return False

        if tagged:
            others = []
            for el in tree_el.iter():
                if el.tag not in DRAWN_SHAPE_TAGS \
                        or id(el) in tagged_ids or inside_tagged(el) \
                        or _is_placeholder_el(el):
                    continue
                ob = absolute_bbox(el)
                if ob:
                    name = ""
                    for c in el.iter(qn("p:cNvPr")):
                        name = c.get("name", "")
                        break
                    others.append((name or "?", ob))
            for el in tagged:
                tb = absolute_bbox(el)
                if not tb:
                    continue
                tname = ""
                for c in el.iter(qn("p:cNvPr")):
                    tname = c.get("name", "")
                    break
                hits = {}
                t_area = tb[2] * tb[3] or 1
                for oname, ob in others:
                    ratio = _overlap_area(tb, ob) / t_area
                    if ratio > OVERLAP_MIN_RATIO:
                        hits[oname] = max(hits.get(oname, 0), ratio)
                if hits:
                    desc = ", ".join(
                        f"'{n}' ({r * 100:.0f}%)" for n, r in
                        sorted(hits.items(), key=lambda kv: -kv[1]))
                    warn(s_idx, tname or "?",
                         f"drawn shape overlaps content: {desc}")

    # 5. deck-level checks
    if deck is not None and tmap is not None:
        _check_deck_consistency(prs, deck, tmap, err)
        _check_text_only_slides(deck, tmap, warn)

    # 6. render
    if render_dir:
        try:
            render_pngs(pptx_path, render_dir, engine)
        except Exception as e:
            err(None, None, f"render failed: {e}")

    return errors, warnings


def _in2(v):
    return round(v / EMU_PER_IN, 2) if v is not None else None


def _overlap(a, b, eps=9525):
    return (a.left + eps < b.left + b.width and
            b.left + eps < a.left + a.width and
            a.top + eps < b.top + b.height and
            b.top + eps < a.top + a.height)


def _overlap_box(ba, bb, eps=9525):
    return (ba[0] + eps < bb[0] + bb[2] and
            bb[0] + eps < ba[0] + ba[2] and
            ba[1] + eps < bb[1] + bb[3] and
            bb[1] + eps < ba[1] + ba[3])


def _norm_divider_title(title):
    return re.sub(r"^[0-9０-９]+[\s.、．]*", "", title).strip()


def _check_deck_consistency(prs, deck, tmap, err):
    story = deck.get("story", {})
    for k in ("purpose", "audience", "desired_action", "key_message"):
        if not story.get(k):
            err(None, None, f"story.{k} is empty")
    # content slides: message + notes
    for i, s in enumerate(deck["slides"], start=1):
        source, comp = resolve_component(s["component"], tmap)
        if source == "library" or comp.get("kind") == "content":
            if not s.get("message"):
                err(i, None, "content slide missing 'message'")
            if not s.get("notes"):
                err(i, None, "content slide missing 'notes'")
    # TOC items vs divider titles
    toc_slide = next((i for i, s in enumerate(deck["slides"])
                      if s["component"] == "toc"), None)
    if toc_slide is None:
        return
    items = deck["slides"][toc_slide]["slots"].get("items", [])
    structural = []
    for i, slide_spec in enumerate(deck["slides"]):
        source, component = resolve_component(slide_spec["component"], tmap)
        if (source == "template" and component.get("kind") == "structural"
                and i != toc_slide):
            structural.append((i, slide_spec))
    # drop title (first slide) and a trailing structural slide (closing)
    structural = [(i, s) for i, s in structural if i > toc_slide]
    titles = [str(s["slots"].get("title", "")) for _, s in structural]
    if len(titles) == len(items) + 1:
        titles = titles[:-1]
    norm = [_norm_divider_title(t) for t in titles]
    if norm != items:
        err(toc_slide + 1, None,
            f"TOC items {items} do not match divider titles {titles}")


def _check_text_only_slides(deck, tmap, warn):
    content = []
    text_only = []
    run = []
    long_runs = []
    for i, spec in enumerate(deck["slides"], start=1):
        source, component = resolve_component(spec["component"], tmap)
        is_content = (source == "library" or
                      component.get("kind") == "content")
        is_text = (source == "template" and is_content and
                   all(slot.get("type") in ("text", "list")
                       for slot in component.get("slots", {}).values()))
        if is_content:
            content.append(i)
        if is_text:
            text_only.append(i)
            run.append(i)
        else:
            if len(run) >= 3:
                long_runs.extend(run)
            run = []
    if len(run) >= 3:
        long_runs.extend(run)
    flagged = set(long_runs)
    if len(content) >= 4 and len(text_only) / len(content) > 0.5:
        flagged.update(text_only)
    for i in sorted(flagged):
        warn(i, None, "text-only slide; consider library components "
             "cards/process/kpi/comparison before using bullets")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--deck")
    ap.add_argument("--map", dest="map_path")
    ap.add_argument("--render", dest="render_dir")
    ap.add_argument("--engine", default="auto",
                    choices=("auto", "powerpoint", "libreoffice"))
    ap.add_argument("--report", dest="report_path")
    ap.add_argument("--build-report", dest="build_report",
                    help="build.py build-report.json; fallbacks become "
                         "warnings")
    args = ap.parse_args(argv)

    deck = json.loads(Path(args.deck).read_text(encoding="utf-8")) \
        if args.deck else None
    tmap = json.loads(Path(args.map_path).read_text(encoding="utf-8")) \
        if args.map_path else None
    fallbacks = []
    library = []
    if args.build_report:
        brep = json.loads(
            Path(args.build_report).read_text(encoding="utf-8"))
        fallbacks = brep.get("fallbacks", [])
        library = brep.get("library", [])

    errors, warnings = check_deck_qa(
        args.pptx, deck, tmap, args.render_dir, args.engine)
    for f in fallbacks:
        warnings.append({
            "slide": f.get("slide"), "shape": f.get("target"),
            "message": f"fallback in slot '{f.get('slot')}': "
                       f"{f.get('reason')} -> {f.get('substitute')} "
                       f"[{f.get('style_source')}]"})
    report = {"file": args.pptx, "errors": errors, "warnings": warnings,
              "fallbacks": fallbacks, "library": library}
    if args.report_path:
        Path(args.report_path).write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8")
    for e in errors:
        loc = f"slide {e['slide']}" if e["slide"] else "deck"
        shp = f" [{e['shape']}]" if e["shape"] else ""
        print(f"ERROR {loc}{shp}: {e['message']}")
    for w in warnings:
        loc = f"slide {w['slide']}" if w.get("slide") else "deck"
        print(f"WARN  {loc}: {w['message']}")
    print(f"qa: {len(errors)} error(s), {len(warnings)} warning(s)")
    if fallbacks:
        print(f"テンプレに準拠できなかった箇所: {len(fallbacks)}件")
        for f in fallbacks:
            print(f"  - スライド{f['slide']} component "
                  f"'{f['component']}' slot '{f['slot']}'"
                  f"（{f['target']}）: {f['reason']} -> {f['substitute']}"
                  f"［{f['style_source']}］")
    else:
        print("テンプレに準拠できなかった箇所: なし")
    print(f"スキルの部品で描いた箇所: {len(library)}件")
    for item in library:
        print(f"  - スライド{item['slide']}: {item['component']} / "
              f"{item['variant']}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
