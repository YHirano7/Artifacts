#!/usr/bin/env python3
"""Unit tests for the pptx-deck skill scripts (build/qa/inventory)."""
import base64
import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PackURI, Part
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
TEMPLATE = ROOT / "templates" / "sample-org" / "template.pptx"
MAP = ROOT / "templates" / "sample-org" / "template-map.json"
DECK = ROOT / "examples" / "redmine-migration" / "deck.json"

sys.path.insert(0, str(SCRIPTS))
from build import absolute_bbox  # noqa: E402

PY = sys.executable

DIAGRAM_URI = ("http://schemas.openxmlformats.org/drawingml/2006/diagram")
DIAGRAM_DATA_RT = ("http://schemas.openxmlformats.org/officeDocument/"
                   "2006/relationships/diagramData")


def run(script, *args, cwd=ROOT):
    return subprocess.run([PY, str(SCRIPTS / script), *args],
                          capture_output=True, text=True, cwd=cwd)


def make_deck(**overrides):
    return {
        "story": {"purpose": "p", "audience": "a", "desired_action": "d",
                  "key_message": "k"},
        "meta": {"title": "t"},
        "slides": [],
    }


def _shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def _scale_group(grp_shape, sx, sy, dx=0, dy=0):
    """Give a group a non-identity transform: scale child coords by
    (sx, sy) then translate by (dx, dy) EMU."""
    xfrm = grp_shape._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_ext = xfrm.find(qn("a:chExt"))
    off.set("x", str(int(off.get("x")) + dx))
    off.set("y", str(int(off.get("y")) + dy))
    ext.set("cx", str(int(int(ch_ext.get("cx")) * sx)))
    ext.set("cy", str(int(int(ch_ext.get("cy")) * sy)))


class TestPptxDeck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name) / "out.pptx"
        if not TEMPLATE.exists():
            r = run("make_sample_template.py")
            assert r.returncode == 0, r.stderr
        cls.deck = json.loads(DECK.read_text(encoding="utf-8"))
        cls.tmap = json.loads(MAP.read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _write_deck(self, deck, name="deck.json"):
        p = Path(self.tmp.name) / name
        p.write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
        return p

    def _build(self, deck, out=None):
        dp = self._write_deck(deck)
        out = out or self.out
        return run("build.py", "--deck", str(dp), "--map", str(MAP),
                   "-o", str(out))

    def _qa(self, pptx, deck=None):
        args = [str(pptx), "--map", str(MAP)]
        if deck is not None:
            args += ["--deck", str(self._write_deck(deck, "qa-deck.json"))]
        return run("qa.py", *args)

    def test_example_builds_and_passes_qa(self):
        out = Path(self.tmp.name) / "example.pptx"
        r = run("build.py", "--deck", str(DECK), "--map", str(MAP),
                "-o", str(out))
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        self.assertEqual(len(prs.slides), 15)
        all_text = "\n".join(
            t.text or ""
            for s in prs.slides for sh in s.shapes
            for t in sh._element.iter(qn("a:t")))
        self.assertNotIn("サンプル", all_text)
        r = run("qa.py", str(out), "--deck", str(DECK), "--map", str(MAP))
        self.assertEqual(r.returncode, 0, r.stdout)

    def test_unknown_component_rejected(self):
        deck = make_deck()
        deck["slides"] = [{"component": "nope", "message": None,
                           "slots": {"x": "y"}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 2)
        self.assertIn("not in template map", r.stderr)

    def test_missing_required_slot_rejected(self):
        deck = make_deck()
        deck["slides"] = [{"component": "bullets", "message": "m",
                           "notes": "n", "slots": {"title": "t"}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 2)
        self.assertIn("required slot 'body'", r.stderr)

    def test_missing_message_on_content_slide_rejected(self):
        deck = make_deck()
        deck["slides"] = [{"component": "bullets", "message": None,
                           "notes": "n",
                           "slots": {"title": "t", "body": ["a"]}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 2)
        self.assertIn("message", r.stderr)

    def test_missing_notes_on_content_slide_rejected(self):
        deck = make_deck()
        deck["slides"] = [{"component": "bullets", "message": "m",
                           "slots": {"body": ["a"]}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 2)
        self.assertIn("notes", r.stderr)

    def test_overflow_detected_by_qa(self):
        deck = make_deck()
        deck["slides"] = [{
            "component": "bullets", "message": "m",
            "notes": "n", "text_only_reason": "overflow test",
            "slots": {"body": ["これは非常に長い本文テキストで、"
                               "ボックスに収まらないことを確認するためのもの。"
                               ] * 60}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 0, r.stderr)
        r = self._qa(self.out, deck)
        self.assertEqual(r.returncode, 1)
        self.assertIn("overflow", r.stdout)

    def test_table_column_expansion(self):
        deck = make_deck()
        deck["slides"] = [{
            "component": "table", "message": "m", "notes": "n",
            "slots": {"table": {
                "header": ["a", "b", "c", "d", "e"],
                "rows": [["1", "2", "3", "4", "5"],
                         ["6", "7", "8", "9", "10"]]}}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(self.out))
        frame = next(sh for sh in prs.slides[0].shapes
                     if sh.name == "DataTable")
        tbl = frame._element.graphic.graphicData.tbl
        cols = tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))
        self.assertEqual(len(cols), 5)
        self.assertEqual(sum(int(c.get("w")) for c in cols), frame.width)
        self.assertEqual(len(tbl.findall(qn("a:tr"))), 3)

    def test_orgchart_too_wide_rejected(self):
        deck = make_deck()
        deck["slides"] = [{
            "component": "orgchart", "message": "m", "notes": "n",
            "slots": {"tree": {
                "name": "root",
                "children": [{"name": f"c{i}"} for i in range(12)]}}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 2)
        self.assertIn("orgchart", r.stderr)

    def test_toc_divider_mismatch_detected(self):
        deck = json.loads(json.dumps(self.deck))
        deck["slides"][1]["slots"]["items"] = ["A", "B", "C", "D"]
        out = Path(self.tmp.name) / "toc-mismatch.pptx"
        r = run("build.py", "--deck", str(self._write_deck(deck, "t.json")),
                "--map", str(MAP), "-o", str(out))
        self.assertEqual(r.returncode, 0, r.stderr)
        r = run("qa.py", str(out), "--deck",
                str(self._write_deck(deck, "t.json")), "--map", str(MAP))
        self.assertEqual(r.returncode, 1)
        self.assertIn("TOC", r.stdout)

    def test_run_formatting_preserved(self):
        deck = make_deck()
        deck["slides"] = [{
            "component": "toc", "message": None,
            "slots": {"title": "目次", "items": ["A", "B"]}}]
        r = self._build(deck)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(self.out))
        toc = next(sh for sh in prs.slides[0].shapes if sh.name == "TocBody")
        run = toc.text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.font.size.pt, 20)
        self.assertFalse(run.font.bold)
        node_deck = make_deck()
        node_deck["slides"] = [{
            "component": "orgchart", "message": "m", "notes": "n",
            "slots": {"tree": {"name": "r", "role": "role1",
                               "children": [{"name": "c1"}]}}}]
        out2 = Path(self.tmp.name) / "org.pptx"
        r = self._build(node_deck, out=out2)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out2))
        node = next(sh for sh in prs.slides[0].shapes if sh.name == "OrgNode")
        p1, p2 = node.text_frame.paragraphs
        self.assertTrue(p1.runs[0].font.bold)
        self.assertEqual(p1.runs[0].font.size.pt, 14)
        self.assertEqual(p2.runs[0].font.size.pt, 11)

    def test_inventory_runs(self):
        out_json = Path(self.tmp.name) / "inv.json"
        r = run("inventory.py", str(TEMPLATE), "--json", str(out_json))
        self.assertEqual(r.returncode, 0, r.stderr)
        data = json.loads(out_json.read_text(encoding="utf-8"))
        self.assertEqual(len(data["slides"]), 8)
        self.assertIn("Title Slide",
                      [l["name"] for l in data["layouts"]])
        chart_slide = data["slides"][7]
        self.assertIn("chart",
                      [sh.get("graphic") for sh in chart_slide["shapes"]])


class TestGroupsLayoutCharts(unittest.TestCase):
    """Groups, layout-mode slots, SmartArt regions and charts."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.count = [0]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _path(self, name):
        self.count[0] += 1
        return Path(self.tmp.name) / f"{self.count[0]}-{name}"

    def _write_json(self, obj, name):
        p = self._path(name)
        p.write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
        return p

    def _mutate_template(self, mutate, name="tpl"):
        """Copy the sample template, apply mutate(prs), return path."""
        tpath = self._path(f"{name}.pptx")
        shutil.copy(TEMPLATE, tpath)
        prs = Presentation(str(tpath))
        mutate(prs)
        prs.save(str(tpath))
        return tpath

    def _map_for(self, components, template):
        return self._write_json(
            {"template": str(template),
             "placeholder_markers": ["サンプル"],
             "components": components}, "map.json")

    def _build(self, deck, map_path):
        dp = self._write_json(deck, "deck.json")
        out = self._path("out.pptx")
        r = run("build.py", "--deck", str(dp), "--map", str(map_path),
                "-o", str(out))
        return r, out

    def _layout_map(self, comp_name, slots, template=TEMPLATE):
        return self._map_for(
            {comp_name: {"kind": "content",
                         "source": {"layout": "Title and Content"},
                         "slots": dict(
                             {"title": {"type": "text", "target": 0}},
                             **slots)}}, template)

    # ---------- groups ----------

    def test_grouped_text_target_filled(self):
        def mutate(prs):
            s = prs.slides[1]  # toc sample
            toc = _shape_by_name(s, "TocBody")
            grp = s.shapes.add_group_shape([toc])
            grp.name = "TocGroup"
        tpath = self._mutate_template(mutate)
        mpath = self._map_for({
            "toc": self_tmap()["components"]["toc"]}, tpath)
        deck = make_deck()
        deck["slides"] = [{"component": "toc", "message": None,
                           "slots": {"title": "目次", "items": ["A", "B"]}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        grp = _shape_by_name(prs.slides[0], "TocGroup")
        toc = _shape_by_name(grp, "TocBody")
        self.assertEqual(toc.text_frame.text.split("\n"), ["A", "B"])

    def test_grouped_region_and_group_prototype(self):
        def mutate(prs):
            s = prs.slides[3]  # orgchart sample
            region = _shape_by_name(s, "OrgRegion")
            region_el = region._element
            g1 = s.shapes.add_group_shape([region])
            g1.name = "RegionGroup"
            _scale_group(g1, 0.8, 0.9, dx=int(Inches(0.4)),
                         dy=int(Inches(0.5)))
            node = _shape_by_name(s, "OrgNode")
            node.name = "OrgBox"
            icon = s.shapes.add_shape(
                MSO_SHAPE.OVAL, node.left + node.width - Inches(0.4),
                node.top + Inches(0.08), Inches(0.3), Inches(0.3))
            icon.name = "OrgIcon"
            g2 = s.shapes.add_group_shape([node, icon])
            g2.name = "OrgNode"
            self.region_bbox = absolute_bbox(region_el)
        tpath = self._mutate_template(mutate)
        org_slot = dict(self_tmap()["components"]["orgchart"])
        org_slot["slots"] = dict(org_slot["slots"])
        org_slot["slots"]["tree"] = dict(org_slot["slots"]["tree"])
        org_slot["slots"]["tree"]["node_prototype"] = "OrgNode"
        mpath = self._map_for({"orgchart": org_slot}, tpath)
        deck = make_deck()
        deck["slides"] = [{
            "component": "orgchart", "message": "m", "notes": "n",
            "slots": {"title": "t", "tree": {
                "name": "部門長", "role": "情報システム部",
                "children": [{"name": "チームA"},
                             {"name": "チームB", "role": "受入"},
                             {"name": "チームC"}]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        names = [sh.name for sh in prs.slides[0].shapes]
        self.assertNotIn("RegionGroup", names)
        self.assertNotIn("OrgRegion", names)
        nodes = [sh for sh in prs.slides[0].shapes if sh.name == "OrgNode"]
        self.assertEqual(len(nodes), 4)
        rx, ry, rw, rh = self.region_bbox
        boxes = []
        for n in nodes:
            ab = absolute_bbox(n._element)
            boxes.append(ab)
            self.assertGreaterEqual(ab[0], rx - 1000)
            self.assertGreaterEqual(ab[1], ry - 1000)
            self.assertLessEqual(ab[0] + ab[2], rx + rw + 1000)
            self.assertLessEqual(ab[1] + ab[3], ry + rh + 1000)
            # group proto -> group node; children scaled, text inside
            texts = [t.text for t in n._element.iter(qn("a:t"))]
            self.assertTrue(any(texts))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                overlap = (a[0] < b[0] + b[2] and b[0] < a[0] + a[2]
                           and a[1] < b[1] + b[3] and b[1] < a[1] + a[3])
                self.assertFalse(overlap, f"nodes {i},{j} overlap")
        # text distribution: single text-bearing child gets all lines
        root = nodes[0]
        root_texts = [t.text for t in root._element.iter(qn("a:t"))]
        self.assertIn("部門長", root_texts)
        self.assertIn("情報システム部", root_texts)

    def test_duplicate_shape_name_errors(self):
        def mutate(prs):
            s = prs.slides[1]
            box = s.shapes.add_textbox(Inches(9), Inches(6),
                                       Inches(2), Inches(0.5))
            box.name = "Title"
            box.text_frame.text = "dup"
        tpath = self._mutate_template(mutate)
        mpath = self._map_for({"toc": self_tmap()["components"]["toc"]},
                              tpath)
        deck = make_deck()
        deck["slides"] = [{"component": "toc", "message": None,
                           "slots": {"title": "目次", "items": ["A"]}}]
        r, _ = self._build(deck, mpath)
        self.assertEqual(r.returncode, 2)
        self.assertIn("unique", r.stderr)

    # ---------- layout-mode slots ----------

    def test_layout_table(self):
        mpath = self._layout_map("ltable", {
            "table": {"type": "table", "target": 1, "required": True}})
        deck = make_deck()
        deck["slides"] = [{
            "component": "ltable", "message": "m", "notes": "n",
            "slots": {"title": "表題",
                      "table": {"header": ["a", "b", "c"],
                                "rows": [["1", "2", "3"],
                                         ["4", "5", "6"]]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        slide = prs.slides[0]
        self.assertEqual(
            [p.placeholder_format.idx for p in slide.placeholders], [0])
        frame = next(sh for sh in slide.shapes
                     if getattr(sh, "has_table", False) and sh.has_table)
        tbl = frame.table
        self.assertEqual(len(tbl.rows), 3)
        self.assertEqual(len(tbl.columns), 3)
        self.assertEqual(tbl.cell(0, 0).text, "a")
        self.assertEqual(tbl.cell(2, 2).text, "6")

    def test_layout_orgchart_no_prototype(self):
        mpath = self._layout_map("lorg", {
            "tree": {"type": "orgchart", "target": 1, "required": True}})
        deck = make_deck()
        deck["slides"] = [{
            "component": "lorg", "message": "m", "notes": "n",
            "slots": {"title": "体制",
                      "tree": {"name": "root",
                               "children": [{"name": "c1"},
                                            {"name": "c2"}]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        nodes = [sh for sh in prs.slides[0].shapes
                 if sh.name == "SynthNode"]
        self.assertEqual(len(nodes), 3)
        self.assertNotIn(1, [p.placeholder_format.idx
                             for p in prs.slides[0].placeholders])
        texts = [t.text for t in nodes[0]._element.iter(qn("a:t"))]
        self.assertEqual(texts, ["root"])

    def test_layout_timeline_no_prototype(self):
        mpath = self._layout_map("ltl", {
            "steps": {"type": "timeline", "target": 1, "required": True}})
        deck = make_deck()
        deck["slides"] = [{
            "component": "ltl", "message": "m", "notes": "n",
            "slots": {"title": "工程",
                      "steps": [{"label": "準備", "period": "10月"},
                                {"label": "実施", "period": "11月",
                                 "detail": "作業"}]}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        nodes = [sh for sh in prs.slides[0].shapes
                 if sh.name == "SynthNode"]
        self.assertEqual(len(nodes), 2)
        texts = [[t.text for t in n._element.iter(qn("a:t"))]
                 for n in nodes]
        self.assertEqual(texts[1], ["実施", "11月", "作業"])

    # ---------- SmartArt as region ----------

    def test_smartart_frame_as_region(self):
        rid_holder = {}

        def mutate(prs):
            s = prs.slides[3]  # orgchart sample
            region = _shape_by_name(s, "OrgRegion")
            x, y, w, h = (region.left, region.top,
                          region.width, region.height)
            # fake diagram part + rel for the frame to reference
            part = Part(PackURI("/ppt/diagrams/data1.xml"),
                        "application/vnd.openxmlformats-officedocument."
                        "drawingml.diagramData+xml",
                        prs.part.package, b"<dgm:dataModel/>")
            rid = s.part.relate_to(part, DIAGRAM_DATA_RT)
            rid_holder["rid"] = rid
            region._element.getparent().remove(region._element)
            ns = ('xmlns:p="http://schemas.openxmlformats.org/'
                  'presentationml/2006/main" '
                  f'xmlns:a="{qn("a:t").split("}")[0].strip("{")}" '
                  'xmlns:r="http://schemas.openxmlformats.org/'
                  'officeDocument/2006/relationships" '
                  f'xmlns:dgm="{DIAGRAM_URI}"')
            gf = etree.fromstring(
                f'<p:graphicFrame {ns}>'
                '<p:nvGraphicFramePr><p:cNvPr id="950" '
                'name="SmartRegion"/><p:cNvGraphicFramePr/><p:nvPr/>'
                '</p:nvGraphicFramePr>'
                f'<p:xfrm><a:off x="{int(x)}" y="{int(y)}"/>'
                f'<a:ext cx="{int(w)}" cy="{int(h)}"/></p:xfrm>'
                '<a:graphic><a:graphicData '
                f'uri="{DIAGRAM_URI}">'
                f'<dgm:relIds r:dm="{rid}"/>'
                '</a:graphicData></a:graphic></p:graphicFrame>')
            s.shapes._spTree.append(gf)
            self.region_bbox = (int(x), int(y), int(w), int(h))

        tpath = self._mutate_template(mutate)
        org = copy.deepcopy(self_tmap()["components"]["orgchart"])
        org["slots"]["tree"]["target"] = "SmartRegion"
        mpath = self._map_for({"orgchart": org}, tpath)
        deck = make_deck()
        deck["slides"] = [{
            "component": "orgchart", "message": "m", "notes": "n",
            "slots": {"title": "t", "tree": {
                "name": "root", "children": [{"name": "c1"},
                                             {"name": "c2"}]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        names = [sh.name for sh in prs.slides[0].shapes]
        self.assertNotIn("SmartRegion", names)
        nodes = [sh for sh in prs.slides[0].shapes if sh.name == "OrgNode"]
        self.assertEqual(len(nodes), 3)
        rx, ry, rw, rh = self.region_bbox
        for n in nodes:
            self.assertLessEqual(n.left + n.width, rx + rw + 1000)
        # diagram part and its rel are gone from the saved package
        partnames = [str(p.partname)
                     for p in prs.part.package.iter_parts()]
        self.assertFalse(any(p.startswith("/ppt/diagrams/")
                             for p in partnames))

    # ---------- shape ids ----------

    def _slide_shape_ids(self, slide):
        ids = []
        tree = slide.shapes._spTree
        root_nv = tree.find(qn("p:nvGrpSpPr"))
        root = root_nv.find(qn("p:cNvPr")) if root_nv is not None else None
        for e in tree.iter(qn("p:cNvPr")):
            if e is not root:
                ids.append(e.get("id"))
        return ids

    def test_shape_ids_unique_sample_orgchart(self):
        deck = make_deck()
        deck["slides"] = [{
            "component": "orgchart", "message": "m", "notes": "n",
            "slots": {"title": "t", "tree": {
                "name": "root", "children": [{"name": "c1"},
                                             {"name": "c2"}]}}}]
        r, out = self._build(deck, MAP)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        ids = self._slide_shape_ids(prs.slides[0])
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(int(i) > 0 for i in ids))

    def test_shape_ids_unique_layout_synth(self):
        mpath = self._layout_map("lorg", {
            "tree": {"type": "orgchart", "target": 1, "required": True}})
        deck = make_deck()
        deck["slides"] = [{
            "component": "lorg", "message": "m", "notes": "n",
            "slots": {"title": "t",
                      "tree": {"name": "r",
                               "children": [{"name": "a"},
                                            {"name": "b"}]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        ids = self._slide_shape_ids(prs.slides[0])
        self.assertEqual(len(ids), len(set(ids)))

    def test_qa_flags_duplicate_shape_ids(self):
        deck = make_deck()
        deck["slides"] = [{"component": "toc", "message": None,
                           "slots": {"title": "t", "items": ["A"]}}]
        r, out = self._build(deck, MAP)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        els = [e for e in prs.slides[0].shapes._spTree.iter(qn("p:cNvPr"))]
        els[-1].set("id", els[1].get("id"))  # collide two real shapes
        dup = self._path("dup.pptx")
        prs.save(str(dup))
        r = run("qa.py", str(dup), "--map", str(MAP))
        self.assertEqual(r.returncode, 1)
        self.assertIn("duplicate shape id", r.stdout)

    # ---------- charts ----------

    def _chart_deck(self, chart_value, n=1):
        deck = make_deck()
        deck["slides"] = [
            {"component": "chart", "message": "m", "notes": "n",
             "slots": {"title": "t", "chart": v}}
            for v in ([chart_value] if n == 1 else chart_value)]
        return deck

    def test_chart_replace_data_sample_mode(self):
        deck = self._chart_deck({
            "categories": ["1月", "2月", "3月"],
            "series": [{"name": "件数", "values": [10, 20, 15]}]})
        r, out = self._build(deck, MAP)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        frame = _shape_by_name(prs.slides[0], "DataChart")
        chart = frame.chart
        cats = list(chart.plots[0].categories)
        self.assertEqual([str(c) for c in cats], ["1月", "2月", "3月"])
        series = list(chart.plots[0].series)
        self.assertEqual(len(series), 1)
        self.assertEqual(list(series[0].values), [10, 20, 15])

    def test_chart_type_mismatch_fallback(self):
        deck = self._chart_deck({
            "type": "pie",
            "categories": ["a", "b"],
            "series": [{"name": "s", "values": [1, 2]}]})
        r, out = self._build(deck, MAP)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("fallbacks: 1", r.stdout)
        rep = json.loads(
            Path(str(out) + ".build-report.json").read_text(
                encoding="utf-8"))
        self.assertEqual(len(rep["fallbacks"]), 1)
        fb = rep["fallbacks"][0]
        self.assertEqual(fb["slot"], "chart")
        self.assertEqual(fb["target"], "DataChart")
        self.assertIn("グラフ", fb["reason"])
        prs = Presentation(str(out))
        charts = [sh.chart for sh in prs.slides[0].shapes if sh.has_chart]
        self.assertEqual(len(charts), 1)
        self.assertEqual(charts[0].chart_type, XL_CHART_TYPE.PIE)
        # --strict converts the same build into a failure
        strict_out = self._path("strict.pptx")
        dp = self._write_json(deck, "strict-deck.json")
        r = run("build.py", "--deck", str(dp), "--map", str(MAP),
                "-o", str(strict_out), "--strict")
        self.assertEqual(r.returncode, 2)
        self.assertIn("strict", r.stderr)
        self.assertIn("fallback", r.stderr)
        self.assertFalse(strict_out.exists())

    def test_chart_series_length_mismatch(self):
        deck = self._chart_deck({
            "categories": ["a", "b", "c"],
            "series": [{"name": "s", "values": [1, 2]}]})
        r, _ = self._build(deck, MAP)
        self.assertEqual(r.returncode, 2)
        self.assertIn("categories", r.stderr)

    def test_chart_sample_used_twice_independent_parts(self):
        deck = self._chart_deck([
            {"categories": ["a"], "series": [{"name": "s1",
                                              "values": [1]}]},
            {"categories": ["x", "y"], "series": [{"name": "s2",
                                                  "values": [7, 9]}]}],
            n=2)
        r, out = self._build(deck, MAP)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        self.assertEqual(len(prs.slides), 2)
        c0 = _shape_by_name(prs.slides[0], "DataChart").chart
        c1 = _shape_by_name(prs.slides[1], "DataChart").chart
        self.assertEqual([str(c) for c in c0.plots[0].categories], ["a"])
        self.assertEqual([str(c) for c in c1.plots[0].categories],
                         ["x", "y"])
        self.assertEqual(list(c0.plots[0].series[0].values), [1])
        self.assertEqual(list(c1.plots[0].series[0].values), [7, 9])
        chart_parts = [str(p.partname)
                       for p in prs.part.package.iter_parts()
                       if str(p.partname).startswith("/ppt/charts/")]
        self.assertEqual(len(chart_parts), 2)

    def test_layout_chart_builds(self):
        mpath = self._layout_map("lchart", {
            "chart": {"type": "chart", "target": 1, "required": True}})
        deck = make_deck()
        deck["slides"] = [{
            "component": "lchart", "message": "m", "notes": "n",
            "slots": {"title": "実績",
                      "chart": {"type": "line",
                                "categories": ["a", "b"],
                                "series": [{"name": "s",
                                            "values": [3, 5]}]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        slide = prs.slides[0]
        charts = [sh for sh in slide.shapes
                  if getattr(sh, "has_chart", False) and sh.has_chart]
        self.assertEqual(len(charts), 1)
        self.assertNotIn(1, [p.placeholder_format.idx
                             for p in slide.placeholders])
        # missing 'type' must fail in layout mode
        deck["slides"][0]["slots"]["chart"] = {
            "categories": ["a"], "series": [{"name": "s",
                                             "values": [1]}]}
        r, _ = self._build(deck, mpath)
        self.assertEqual(r.returncode, 2)
        self.assertIn("type", r.stderr)

    # ---------- fallbacks ----------

    def _add_frame(self, slide, name, uri, x, y, w, h, inner=""):
        ns = ('xmlns:p="http://schemas.openxmlformats.org/'
              'presentationml/2006/main" '
              f'xmlns:a="{qn("a:t").split("}")[0].strip("{")}"')
        gf = etree.fromstring(
            f'<p:graphicFrame {ns}>'
            '<p:nvGraphicFramePr><p:cNvPr id="960" '
            f'name="{name}"/><p:cNvGraphicFramePr/><p:nvPr/>'
            '</p:nvGraphicFramePr>'
            f'<p:xfrm><a:off x="{int(x)}" y="{int(y)}"/>'
            f'<a:ext cx="{int(w)}" cy="{int(h)}"/></p:xfrm>'
            '<a:graphic><a:graphicData '
            f'uri="{uri}">{inner}</a:graphicData></a:graphic>'
            '</p:graphicFrame>')
        slide.shapes._spTree.append(gf)
        return gf

    def _table_fallback_fixture(self, mutate_extra=None):
        """table slot target replaced by an OLE graphicFrame at a
        known bbox; returns (map_path, ole_bbox)."""
        box = [int(Inches(6)), int(Inches(3)), int(Inches(5)),
               int(Inches(2))]

        def mutate(prs):
            s = prs.slides[2]  # table sample
            self._add_frame(s, "OleTable", OLE_URI, *box)
            if mutate_extra:
                mutate_extra(prs)
        tpath = self._mutate_template(mutate)
        comp = copy.deepcopy(self_tmap()["components"]["table"])
        comp["slots"]["table"]["target"] = "OleTable"
        return self._map_for({"table": comp}, tpath), box

    def _table_deck(self):
        deck = make_deck()
        deck["slides"] = [{
            "component": "table", "message": "m", "notes": "n",
            "slots": {"table": {"header": ["列1", "列2"],
                                "rows": [["a", "b"], ["c", "d"]]}}}]
        return deck

    def _report(self, out):
        return json.loads(Path(str(out) + ".build-report.json")
                          .read_text(encoding="utf-8"))

    def test_table_ole_target_fallback(self):
        mpath, box = self._table_fallback_fixture()
        r, out = self._build(self._table_deck(), mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("fallbacks: 1", r.stdout)
        rep = self._report(out)
        self.assertEqual(len(rep["fallbacks"]), 1)
        fb = rep["fallbacks"][0]
        self.assertEqual(fb["slot"], "table")
        self.assertEqual(fb["target"], "OleTable")
        self.assertIn("OLE", fb["reason"])
        self.assertEqual(fb["style_source"],
                         "slide 3 'DataTable' (explicit fills)")
        prs = Presentation(str(out))
        # the OLE frame is gone; a native table sits at its bbox
        ole = [sh for sh in prs.slides[0].shapes if sh.name == "OleTable"]
        self.assertEqual(len(ole), 1)
        f = ole[0]
        self.assertTrue(f.has_table)
        self.assertEqual((int(f.left), int(f.top),
                          int(f.width), int(f.height)), tuple(box))
        # header fill borrowed from the DataTable donor (navy 1B3A66)
        tbl = f._element.graphic.graphicData.tbl
        trs = tbl.findall(qn("a:tr"))
        hdr = trs[0].findall(qn("a:tc"))[0]
        fill = hdr.find(qn("a:tcPr")).find(qn("a:solidFill"))
        self.assertEqual(fill.find(qn("a:srgbClr")).get("val"),
                         "1B3A66")
        # donor typography carried over: 14pt bold white header,
        # 13pt body, 0.8in rows
        hrpr = hdr.find(qn("a:txBody")).find(qn("a:p")) \
            .find(qn("a:r")).find(qn("a:rPr"))
        self.assertEqual(hrpr.get("sz"), "1400")
        self.assertEqual(hrpr.get("b"), "1")
        hf = hrpr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        self.assertEqual(hf.get("val"), "FFFFFF")
        body = trs[1].findall(qn("a:tc"))[0]
        brpr = body.find(qn("a:txBody")).find(qn("a:p")) \
            .find(qn("a:r")).find(qn("a:rPr"))
        self.assertEqual(brpr.get("sz"), "1300")
        self.assertIsNone(brpr.get("b"))
        self.assertEqual(trs[1].get("h"), "731520")

    def test_table_fallback_theme_when_no_donor(self):
        def wipe(prs):
            # remove every template table
            for s in prs.slides:
                for el in s.shapes._spTree.findall(qn("p:graphicFrame")):
                    gd = el.find(f"{qn('a:graphic')}/"
                                 f"{qn('a:graphicData')}")
                    if gd is not None and \
                            (gd.get("uri") or "").endswith("/table"):
                        el.getparent().remove(el)
            # and clear the default table style
            for p in prs.part.package.iter_parts():
                if str(p.partname) == "/ppt/tableStyles.xml":
                    root = etree.fromstring(p.blob)
                    root.attrib.pop("def", None)
                    p._blob = etree.tostring(
                        root, xml_declaration=True,
                        encoding="UTF-8", standalone=True)
        mpath, _ = self._table_fallback_fixture(wipe)
        r, out = self._build(self._table_deck(), mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        fb = self._report(out)["fallbacks"][0]
        self.assertEqual(fb["style_source"], "theme")
        prs = Presentation(str(out))
        ole = [sh for sh in prs.slides[0].shapes if sh.name == "OleTable"]
        self.assertEqual(len(ole), 1)
        self.assertTrue(ole[0].has_table)

    def test_table_plain_shape_fallback(self):
        def mutate(prs):
            s = prs.slides[2]
            box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6),
                                     Inches(3), Inches(5), Inches(2))
            box.name = "PlainBox"
        tpath = self._mutate_template(mutate)
        comp = copy.deepcopy(self_tmap()["components"]["table"])
        comp["slots"]["table"]["target"] = "PlainBox"
        mpath = self._map_for({"table": comp}, tpath)
        r, out = self._build(self._table_deck(), mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        fb = self._report(out)["fallbacks"][0]
        self.assertEqual(fb["target"], "PlainBox")
        self.assertEqual(fb["style_source"],
                         "slide 3 'DataTable' (explicit fills)")
        prs = Presentation(str(out))
        box_shapes = [sh for sh in prs.slides[0].shapes
                      if sh.name == "PlainBox"]
        self.assertEqual(len(box_shapes), 1)
        self.assertTrue(box_shapes[0].has_table)

    def test_picture_prototype_fallback(self):
        pic_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")

        def mutate(prs):
            s = prs.slides[3]
            part = Part(PackURI("/ppt/media/proto.png"),
                        "image/png", prs.part.package, pic_bytes)
            rid = s.part.relate_to(part, RT.IMAGE)
            ns = ('xmlns:p="http://schemas.openxmlformats.org/'
                  'presentationml/2006/main" '
                  f'xmlns:a="{qn("a:t").split("}")[0].strip("{")}" '
                  'xmlns:r="http://schemas.openxmlformats.org/'
                  'officeDocument/2006/relationships"')
            pic = etree.fromstring(
                f'<p:pic {ns}>'
                '<p:nvPicPr><p:cNvPr id="961" name="PicProto"/>'
                '<p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
                f'<p:blipFill><a:blip r:embed="{rid}"/>'
                '<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
                '<p:spPr><a:xfrm>'
                f'<a:off x="{int(Inches(1))}" y="{int(Inches(1))}"/>'
                f'<a:ext cx="{int(Inches(1))}" '
                f'cy="{int(Inches(1))}"/></a:xfrm>'
                '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                '</p:spPr></p:pic>')
            s.shapes._spTree.append(pic)
        tpath = self._mutate_template(mutate)
        comp = copy.deepcopy(self_tmap()["components"]["orgchart"])
        comp["slots"]["tree"]["node_prototype"] = "PicProto"
        mpath = self._map_for({"orgchart": comp}, tpath)
        deck = make_deck()
        deck["slides"] = [{
            "component": "orgchart", "message": "m", "notes": "n",
            "slots": {"title": "t", "tree": {
                "name": "root",
                "children": [{"name": "a"}, {"name": "b"}]}}}]
        r, out = self._build(deck, mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        fb = self._report(out)["fallbacks"][0]
        self.assertEqual(fb["slot"], "tree")
        self.assertEqual(fb["target"], "PicProto")
        self.assertIn("画像", fb["reason"])
        prs = Presentation(str(out))
        self.assertEqual(
            len([sh for sh in prs.slides[0].shapes
                 if sh.name == "SynthNode"]), 3)

    def test_contract_error_not_a_fallback(self):
        tpath = self._mutate_template(lambda prs: None)
        comp = copy.deepcopy(self_tmap()["components"]["table"])
        comp["slots"]["table"]["target"] = "DoesNotExist"
        mpath = self._map_for({"table": comp}, tpath)
        r, out = self._build(self._table_deck(), mpath)
        self.assertEqual(r.returncode, 2)
        self.assertIn("not found", r.stderr)
        rep = Path(str(out) + ".build-report.json")
        self.assertTrue(
            not rep.exists()
            or json.loads(rep.read_text(encoding="utf-8"))
            ["fallbacks"] == [])

    def test_fallback_overlap_warned_not_error(self):
        # the fallback table drawn at the OleTable bbox overlaps the
        # DataTable donor still sitting on the sample slide
        mpath, _ = self._table_fallback_fixture()
        r, out = self._build(self._table_deck(), mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        r = run("qa.py", str(out))
        self.assertEqual(r.returncode, 0, r.stdout)
        self.assertIn("overlap", r.stdout)
        self.assertIn("DataTable", r.stdout)

    def test_qa_build_report_warnings(self):
        mpath, _ = self._table_fallback_fixture()
        r, out = self._build(self._table_deck(), mpath)
        self.assertEqual(r.returncode, 0, r.stderr)
        rep_path = str(out) + ".build-report.json"
        r = run("qa.py", str(out), "--build-report", rep_path)
        self.assertEqual(r.returncode, 0, r.stdout)
        self.assertIn("テンプレに準拠できなかった箇所: 1件", r.stdout)
        qrep = self._path("qrep.json")
        r = run("qa.py", str(out), "--build-report", rep_path,
                "--report", str(qrep))
        self.assertEqual(r.returncode, 0, r.stdout)
        data = json.loads(qrep.read_text(encoding="utf-8"))
        self.assertEqual(len(data["fallbacks"]), 1)
        self.assertTrue(any("fallback" in w["message"]
                            for w in data["warnings"]))
        self.assertEqual(data["errors"], [])

    def test_redmine_build_report_zero_fallbacks(self):
        out = self._path("redmine.pptx")
        r = run("build.py", "--deck", str(DECK), "--map", str(MAP),
                "-o", str(out))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("fallbacks: 0", r.stdout)
        rep = self._report(out)
        self.assertEqual(rep["fallbacks"], [])
        self.assertGreater(rep["native"], 0)
        r = run("qa.py", str(out), "--build-report",
                str(out) + ".build-report.json")
        self.assertEqual(r.returncode, 0, r.stdout)
        self.assertIn("テンプレに準拠できなかった箇所: なし", r.stdout)


OLE_URI = ("http://schemas.openxmlformats.org/presentationml/2006/ole")


def self_tmap():
    return json.loads(MAP.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
