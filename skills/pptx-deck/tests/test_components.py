import copy
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build import validate_deck
import common
import components as components_module
import inventory as inventory_module
from common import BuildError, GENERATED_MARK, read_theme
from components import (EMU_PER_IN, LIBRARY, ROLE_DEFAULTS, _scheme_ref,
                        _write_color, fit_siblings, fit_size, resolve_canvas,
                        resolve_component, resolve_style, validate_component)
from inventory import collect
from make_sample_template import build_minimal_template
from qa import _check_text_only_slides

GALLERY = ROOT / "examples" / "component-gallery" / "deck.json"
MINIMAL_MAP = ROOT / "templates" / "minimal-org" / "template-map.json"


class TestComponents(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.template = Path(cls.tmp.name) / "minimal.pptx"
        build_minimal_template().save(cls.template)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_gallery_builds_every_variant_and_passes_qa(self):
        output = Path(self.tmp.name) / "gallery.pptx"
        build_report = Path(str(output) + ".build-report.json")
        qa_report = Path(self.tmp.name) / "qa.json"
        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "utf-8"
        build = subprocess.run(
            [sys.executable, str(SCRIPTS / "build.py"),
             "--deck", str(GALLERY), "--map", str(MINIMAL_MAP),
             "--template", str(self.template), "-o", str(output), "--strict"],
            capture_output=True, text=True, env=env)
        self.assertEqual(build.returncode, 0, build.stderr + build.stdout)
        prs = Presentation(str(output))
        generated_shapes = []
        round_rectangles = 0
        for slide in prs.slides:
            for shape_el in slide._element.xpath(".//p:sp"):
                c_nvpr = shape_el.find(
                    f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
                if c_nvpr is None or c_nvpr.get("descr") != GENERATED_MARK:
                    continue
                generated_shapes.append(shape_el)
                self.assertEqual(shape_el.xpath(".//p:style"), [])
                geometry = shape_el.find(
                    f"{qn('p:spPr')}/{qn('a:prstGeom')}")
                if geometry is None or geometry.get("prst") != "roundRect":
                    continue
                round_rectangles += 1
                ext = shape_el.find(
                    f"{qn('p:spPr')}/{qn('a:xfrm')}/{qn('a:ext')}")
                width = int(ext.get("cx"))
                height = int(ext.get("cy"))
                expected = max(
                    1, min(100000, round(0.06 * EMU_PER_IN
                                         / min(width, height) * 100000)))
                av_lst = geometry.find(qn("a:avLst"))
                adjustment = av_lst.find(
                    f"{qn('a:gd')}[@name='adj']")
                self.assertEqual(adjustment.get("fmla"), f"val {expected}")
        self.assertTrue(generated_shapes)
        self.assertGreater(round_rectangles, 0)

        def find_shape(name):
            for slide in prs.slides:
                for shape_el in slide._element.xpath(".//p:sp"):
                    c_nvpr = shape_el.find(
                        f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
                    if c_nvpr is not None and c_nvpr.get("name") == name:
                        return shape_el
            return None

        def generated_bounds(slide):
            positions = []
            for shape in slide.shapes:
                props = shape._element.xpath(
                    "./p:nvSpPr/p:cNvPr | ./p:nvGrpSpPr/p:cNvPr | "
                    "./p:nvCxnSpPr/p:cNvPr | "
                    "./p:nvGraphicFramePr/p:cNvPr")
                if (props and props[0].get("descr") == GENERATED_MARK
                        and props[0].get("name") != "Component lead"):
                    positions.append((shape.top, shape.top + shape.height))
            self.assertTrue(positions)
            return min(top for top, _ in positions), max(
                bottom for _, bottom in positions)

        lead_slide = prs.slides[1]
        lead = next(shape for shape in lead_slide.shapes
                    if shape.name == "Component lead")
        self.assertAlmostEqual(lead.top / EMU_PER_IN, 1.8, delta=0.01)
        content_top, content_bottom = generated_bounds(lead_slide)
        expected_center = (lead.top + lead.height + Inches(0.12)
                           + Inches(6.75)) / 2
        self.assertAlmostEqual((content_top + content_bottom) / 2,
                               expected_center, delta=Inches(0.03))

        takeaway_top, takeaway_bottom = generated_bounds(prs.slides[2])
        self.assertAlmostEqual(
            (takeaway_top + takeaway_bottom) / 2,
            Inches((1.8 + 6.75) / 2), delta=Inches(0.03))

        charts = [shape for slide in prs.slides for shape in slide.shapes
                  if shape.name == "Library chart"]
        self.assertEqual(len(charts), 2)
        for chart in charts:
            self.assertAlmostEqual(
                chart.height / EMU_PER_IN, 4.95, delta=0.01)
        side_chart = prs.slides[20]
        chart = next(shape for shape in side_chart.shapes
                     if shape.name == "Library chart")
        panel = next(shape for shape in side_chart.shapes
                     if shape.name == "Chart key points panel")
        self.assertAlmostEqual(panel.top + panel.height / 2,
                               chart.top + chart.height / 2,
                               delta=Inches(0.01))

        for name in ("Layer 1 item 1", "Roadmap track 1",
                     "Checklist item 1", "Checklist owner 2",
                     "Checklist due 2", "Chart point 1"):
            body_pr = find_shape(name).find(
                f"{qn('p:txBody')}/{qn('a:bodyPr')}")
            self.assertEqual(body_pr.get("anchor"), "ctr", name)
        check_run = find_shape("Checklist status 1").find(
            f"{qn('p:txBody')}/{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}")
        self.assertEqual(check_run.get("sz"), "1500")

        card_body = find_shape("Card 1 copy")
        body_paragraphs = card_body.xpath("./p:txBody/a:p")
        self.assertGreater(len(body_paragraphs), 1)
        for paragraph in body_paragraphs:
            ppr = paragraph.find(qn("a:pPr"))
            self.assertIsNotNone(ppr.find(qn("a:buChar")))
            self.assertEqual(ppr.get("marL"), str(round(0.22 * EMU_PER_IN)))
            self.assertEqual(ppr.get("indent"),
                             str(round(-0.12 * EMU_PER_IN)))
            self.assertEqual(
                ppr.find(f"{qn('a:spcAft')}/{qn('a:spcPts')}").get("val"),
                "400")
        process_detail = find_shape("Process 1 detail")
        for paragraph in process_detail.xpath("./p:txBody/a:p"):
            ppr = paragraph.find(qn("a:pPr"))
            self.assertIsNone(ppr.find(qn("a:buChar")) if ppr is not None
                              else None)
        kpi_value = find_shape("KPI 1 value")
        value_paragraphs = kpi_value.xpath("./p:txBody/a:p")
        self.assertEqual(len(value_paragraphs), 1)
        self.assertEqual(
            value_paragraphs[0].xpath("./a:r/a:t/text()"),
            ["1", " 窓口"])

        cycle_nodes = [shape for shape in prs.slides[14].shapes
                       if shape.name in {
                           "Cycle 1", "Cycle 2", "Cycle 3", "Cycle 4"}]
        self.assertEqual(len(cycle_nodes), 4)
        for index, node in enumerate(cycle_nodes):
            for other in cycle_nodes[index + 1:]:
                overlaps = (
                    node.left < other.left + other.width
                    and other.left < node.left + node.width
                    and node.top < other.top + other.height
                    and other.top < node.top + node.height)
                self.assertFalse(
                    overlaps, f"{node.name} overlaps {other.name}")

        connectors = [shape for slide in prs.slides for shape in slide.shapes
                      if shape.name.startswith("Cycle connector")]
        self.assertEqual(len(connectors), 4)
        for connector in connectors:
            ln = connector._element.spPr.find(qn("a:ln"))
            arrow = ln.find(qn("a:headEnd"))
            self.assertEqual(arrow.get("type"), "triangle")
            self.assertEqual(arrow.get("w"), "med")
        report = json.loads(build_report.read_text(encoding="utf-8"))
        variants = {(item["component"], item["variant"])
                    for item in report["library"]}
        self.assertEqual(
            variants,
            {(name, variant) for name, component in LIBRARY.items()
             for variant in component.variants})
        qa = subprocess.run(
            [sys.executable, str(SCRIPTS / "qa.py"), str(output),
             "--deck", str(GALLERY), "--map", str(MINIMAL_MAP),
             "--build-report", str(build_report), "--report", str(qa_report)],
            capture_output=True, text=True, env=env)
        self.assertEqual(qa.returncode, 0, qa.stderr + qa.stdout)
        result = json.loads(qa_report.read_text(encoding="utf-8"))
        self.assertEqual(result["errors"], [])
        self.assertEqual(len(result["library"]), len(variants))
        self.assertIn("スキルの部品で描いた箇所", qa.stdout)

    def test_minimal_cover_and_section_layouts(self):
        prs = Presentation(str(self.template))

        def get_layout(name):
            return next(layout for layout in prs.slide_layouts
                        if layout.name == name)

        def get_placeholder(layout, idx):
            return next(ph for ph in layout.placeholders
                        if ph.placeholder_format.idx == idx)

        def title_style(layout):
            placeholder = get_placeholder(layout, 0)
            level = placeholder._element.find(
                f"{qn('p:txBody')}/{qn('a:lstStyle')}/{qn('a:lvl1pPr')}")
            run = level.find(qn("a:defRPr"))
            self.assertEqual(level.get("algn"), "l")
            self.assertEqual(run.get("sz"), "4000")
            self.assertEqual(run.get("b"), "1")
            self.assertEqual(
                run.find(f"{qn('a:solidFill')}/{qn('a:schemeClr')}")
                .get("val"), "dk2")
            self.assertAlmostEqual(placeholder.left / EMU_PER_IN, 0.9,
                                   delta=0.01)
            self.assertAlmostEqual(placeholder.top / EMU_PER_IN, 2.6,
                                   delta=0.01)
            return placeholder

        cover = get_layout("Title Slide")
        title_style(cover)
        subtitle = get_placeholder(cover, 1)
        self.assertAlmostEqual(subtitle.left / EMU_PER_IN, 0.9, delta=0.01)
        self.assertAlmostEqual(subtitle.top / EMU_PER_IN, 3.47, delta=0.01)
        subtitle_level = subtitle._element.find(
            f"{qn('p:txBody')}/{qn('a:lstStyle')}/{qn('a:lvl1pPr')}")
        subtitle_run = subtitle_level.find(qn("a:defRPr"))
        self.assertEqual(subtitle_level.get("algn"), "l")
        self.assertEqual(subtitle_run.get("sz"), "1800")
        self.assertEqual(
            subtitle_run.find(
                f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val"),
            "595F6B")

        for name in ("Title Slide", "Section Header"):
            layout = get_layout(name)
            if name == "Section Header":
                title_style(layout)
            accent = next(shape for shape in layout.shapes
                          if shape.name == f"{name} Accent")
            self.assertAlmostEqual(accent.left / EMU_PER_IN, 0.9,
                                   delta=0.01)
            self.assertAlmostEqual(accent.top / EMU_PER_IN, 2.28,
                                   delta=0.01)
            self.assertAlmostEqual(accent.width / EMU_PER_IN, 1.2,
                                   delta=0.01)
            self.assertAlmostEqual(accent.height / EMU_PER_IN, 0.06,
                                   delta=0.01)
            self.assertEqual(
                accent._element.spPr.find(
                    f"{qn('a:solidFill')}/{qn('a:schemeClr')}")
                .get("val"), "accent2")

    def test_gallery_content_density(self):
        deck = json.loads(GALLERY.read_text(encoding="utf-8"))
        slides = deck["slides"]
        for slide in slides:
            if slide.get("component") == "cards":
                for item in slide["slots"]["items"]:
                    self.assertEqual(len(item["body"]), 3)
            elif slide.get("component") == "comparison":
                for side in ("left", "right"):
                    self.assertEqual(len(slide["slots"][side]["items"]), 5)
            elif slide.get("component") == "layers":
                for layer in slide["slots"]["layers"]:
                    self.assertIn(len(layer["items"]), (3, 4))
            elif (slide.get("component") == "process"
                  and slide.get("variant") == "circles"):
                for step in slide["slots"]["steps"]:
                    self.assertIn(step["detail"].count("。"), (1, 2))
            elif slide.get("component") == "cycle":
                for step in slide["slots"]["steps"]:
                    self.assertIn(step["detail"].count("。"), (1, 2))

    def test_theme_style_serialization_and_automatic_contrast(self):
        prs = Presentation(str(self.template))
        style = resolve_style(prs, {})
        default_sizes = {
            "heading": 20, "body": 16, "caption": 13,
            "number": 48, "min": 11,
        }
        self.assertEqual(style.sizes, default_sizes)
        schema = json.loads(
            (ROOT / "schema" / "template-map.schema.json").read_text(
                encoding="utf-8"))
        size_schema = schema["properties"]["style"]["properties"]["sizes"][
            "properties"]
        for name, value in default_sizes.items():
            self.assertEqual(size_schema[name]["minimum"], 11)
            self.assertEqual(size_schema[name]["default"], value)
        self.assertEqual(style.colors["primary"].scheme, "accent1")
        self.assertEqual(style.colors["on_primary"].scheme, "lt1")
        self.assertEqual(ROLE_DEFAULTS["surface"], "bg2")
        clr_map = {"tx1": "dk1", "tx2": "dk2",
                   "bg1": "lt1", "bg2": "lt2"}
        clr_map.update(prs.slide_masters[0].element.find(qn("p:clrMap")).attrib)
        expected_surface = _scheme_ref("bg2", clr_map)
        self.assertEqual(style.colors["surface"], expected_surface)
        self.assertEqual(style.fonts["heading_ea"], "Yu Gothic")
        self.assertEqual(style.fonts["body_latin"], "Yu Gothic")
        surface_override = resolve_style(
            prs, {"style": {"palette": {"surface": "#123456"}}})
        self.assertEqual(surface_override.colors["surface"].rgb, "123456")

        themed = OxmlElement("p:spPr")
        _write_color(themed, style.colors["primary"])
        scheme = themed.find(f"{qn('a:solidFill')}/{qn('a:schemeClr')}")
        self.assertIsNotNone(scheme)
        self.assertEqual(scheme.get("val"), "accent1")

        explicit = resolve_style(
            prs, {"style": {"palette": {"primary": "#123456"}}})
        shape = OxmlElement("p:spPr")
        _write_color(shape, explicit.colors["primary"])
        color = shape.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
        self.assertIsNotNone(color)
        self.assertEqual(color.get("val"), "123456")

    def test_fit_and_sibling_size_consistency(self):
        short = "Short label"
        long = "This longer label needs a smaller font size to fit"
        short_size = fit_size(short, 2.0, 0.7, 18, 9, "short")
        long_size = fit_size(long, 2.0, 0.7, 18, 9, "long")
        self.assertLess(long_size, short_size)
        self.assertEqual(
            fit_siblings([short, long], 2.0, 0.7, 18, 9, "cards"), long_size)
        with self.assertRaisesRegex(BuildError, "shorten to about"):
            fit_size("x" * 120, 1.0, 0.25, 12, 9,
                     "slide 3 (component 'cards'), items[2].body")

    def test_variants_limits_emphasis_and_roadmap_ranges(self):
        with self.assertRaisesRegex(BuildError, "unknown variant"):
            validate_component("cards", {"items": []}, "unknown", 2)
        with self.assertRaisesRegex(BuildError, "split the content across slides"):
            validate_component(
                "cards",
                {"items": [{"heading": "h", "body": "b"}] * 7},
                "outline", 2)
        with self.assertRaisesRegex(BuildError, "emphasis"):
            validate_component(
                "cards",
                {"items": [{"heading": "h", "body": "b"}] * 2,
                 "emphasis": 3},
                "outline", 2)
        with self.assertRaisesRegex(BuildError, "within 1..2"):
            validate_component(
                "roadmap",
                {"periods": ["Oct", "Nov"],
                 "tracks": [{"label": "Work",
                             "bars": [{"start": 1, "end": 3,
                                       "label": "Milestone"}]}]},
                "roadmap", 2)

    def test_template_resolution_canvas_and_template_variant_rejection(self):
        tmap = json.loads(MINIMAL_MAP.read_text(encoding="utf-8"))
        self.assertEqual(resolve_component("title", tmap)[0], "template")
        self.assertEqual(resolve_component("cards", tmap)[0], "library")
        custom = copy.deepcopy(tmap)
        custom["components"]["cards"] = custom["components"]["bullets"]
        self.assertEqual(resolve_component("cards", custom)[0], "template")

        prs = Presentation(str(self.template))
        layout, title_idx, region = resolve_canvas(prs, tmap, 1)
        self.assertEqual(layout.name, "Title Only")
        self.assertIsNotNone(title_idx)
        self.assertAlmostEqual(region[2], 12.33, places=2)
        self.assertGreater(region[3], 0)
        with self.assertRaisesRegex(BuildError, "map に canvas.layout を指定"):
            resolve_canvas(prs, {"canvas": {"layout": "Missing"}}, 1)

        deck = json.loads(GALLERY.read_text(encoding="utf-8"))
        deck["slides"][0]["variant"] = "header"
        with self.assertRaisesRegex(BuildError, "does not support variant"):
            validate_deck(deck, tmap)

    def test_inventory_theme_explicit_colors_and_canvas_candidates(self):
        data = collect(self.template)
        self.assertIs(components_module.read_theme, common.read_theme)
        self.assertIs(inventory_module.read_theme, common.read_theme)
        self.assertEqual(data["theme"], read_theme(Presentation(
            str(self.template))))
        self.assertEqual(data["theme"]["colors"]["accent1"], "#0F5B4F")
        self.assertEqual(data["theme"]["colors"]["accent2"], "#E07A1F")
        self.assertEqual(data["theme"]["fonts"]["major"]["ea"], "Yu Gothic")
        self.assertTrue(data["used_colors"])
        self.assertIn("Title Only", data["canvas_candidates"])
        self.assertLessEqual(len(data["used_colors"]), 8)

    def test_minimal_template_cli_generates_three_sample_slides(self):
        output = Path(self.tmp.name) / "minimal-cli.pptx"
        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "utf-8"
        result = subprocess.run(
            [sys.executable, str(SCRIPTS / "make_sample_template.py"),
             "--variant", "minimal", "-o", str(output)],
            capture_output=True, text=True, env=env)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        prs = Presentation(str(output))
        self.assertEqual(len(prs.slides), 3)
        self.assertEqual(
            [slide.slide_layout.name for slide in prs.slides],
            ["Title Slide", "Title and Content", "Title and Content"])
        self.assertIn("Section Header", [layout.name
                                         for layout in prs.slide_layouts])
        self.assertIn("Title Only", [layout.name
                                     for layout in prs.slide_layouts])
        for name in ("Title Slide", "Section Header", "Title Only",
                     "Title and Content"):
            layout = next(item for item in prs.slide_layouts
                          if item.name == name)
            title = next(ph for ph in layout.placeholders
                         if ph.placeholder_format.idx == 0)
            cover_or_divider = name in ("Title Slide", "Section Header")
            self.assertAlmostEqual(
                title.top / EMU_PER_IN,
                2.6 if cover_or_divider else 0.4, places=2)
            level = title._element.find(
                f"{qn('p:txBody')}/{qn('a:lstStyle')}/"
                f"{qn('a:lvl1pPr')}")
            run_props = level.find(qn("a:defRPr"))
            self.assertEqual(
                run_props.get("sz"), "4000" if cover_or_divider else "2800")
            self.assertEqual(run_props.get("b"), "1")
            self.assertEqual(
                run_props.find(
                    f"{qn('a:solidFill')}/{qn('a:schemeClr')}").get("val"),
                "dk2")

    def test_text_only_warning_threshold(self):
        tmap = {"components": {
            "bullets": {
                "kind": "content",
                "slots": {"title": {"type": "text"},
                          "body": {"type": "list"}}
            },
            "chart": {
                "kind": "content",
                "slots": {"chart": {"type": "chart"}}
            }
        }}
        deck = {"slides": [
            {"component": "bullets"} for _ in range(3)
        ] + [{"component": "chart"}]}
        warnings = []
        _check_text_only_slides(
            deck, tmap,
            lambda slide, shape, message:
            warnings.append((slide, shape, message)))
        self.assertEqual([warning[0] for warning in warnings], [1, 2, 3])
        self.assertIn("cards/process/kpi/comparison", warnings[0][2])


if __name__ == "__main__":
    unittest.main()
