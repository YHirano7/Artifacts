import copy
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build import check_visual_policy, validate_deck  # noqa: E402
from common import BuildError  # noqa: E402
from components import LIBRARY, resolve_component, validate_component  # noqa: E402,E501
from make_sample_template import build_minimal_template  # noqa: E402

MINIMAL_MAP = ROOT / "templates" / "minimal-org" / "template-map.json"
STORY = {"purpose": "p", "audience": "a", "desired_action": "d",
         "key_message": "k"}


def _deck(slides):
    return {"story": STORY, "meta": {"title": "t"}, "slides": slides}


def _run(script, *args):
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    return subprocess.run([sys.executable, str(SCRIPTS / script), *args],
                          capture_output=True, text=True, env=env)


class TestBusinessValidation(unittest.TestCase):
    def assertRejects(self, name, slots, variant, fragment):
        with self.assertRaises(BuildError) as ctx:
            validate_component(name, slots, variant, 1)
        self.assertIn(fragment, str(ctx.exception))

    def test_new_components_are_registered(self):
        for name in ("table", "scorecard", "timeline", "tree", "hub",
                     "swimlane", "action_plan"):
            self.assertIn(name, LIBRARY)

    def test_table_cells_must_match_header(self):
        self.assertRejects("table", {"header": ["a", "b"],
                                     "rows": [["1"]]}, None, "expected 2")
        self.assertRejects("table", {"header": ["a", "b"],
                                     "rows": [["1", "2"]],
                                     "col_widths": [1]}, None, "col_widths")
        self.assertRejects("table", {"header": ["a", "b"],
                                     "rows": [["1", {"text": "x",
                                                     "tone": "loud"}]]},
                           None, "rows")

    def test_scorecard_rating_and_raci_rules(self):
        base = {"columns": ["A", "B"],
                "rows": [{"label": "x", "cells": ["◎", "○"]},
                         {"label": "y", "cells": ["△", "×"]}]}
        validate_component("scorecard", base, "rating", 1)
        bad = copy.deepcopy(base)
        bad["rows"][0]["cells"][0] = "5"
        self.assertRejects("scorecard", bad, "rating", "rating cells")
        bad = copy.deepcopy(base)
        bad["recommend"] = [3]
        self.assertRejects("scorecard", bad, "rating", "recommend")
        raci = {"columns": ["A", "B"],
                "rows": [{"label": "x", "cells": ["A/R", "C"]},
                         {"label": "y", "cells": ["R", "I"]}]}
        self.assertRejects("scorecard", raci, "raci", "exactly one 'A'")
        raci["rows"][1]["cells"][1] = "A"
        validate_component("scorecard", raci, "raci", 1)

    def test_tree_limits_leaves(self):
        branches = [{"label": f"b{i}", "children": [
            {"label": f"l{i}{j}"} for j in range(3)]} for i in range(4)]
        self.assertRejects("tree", {"root": "r", "branches": branches},
                           None, "exceed 9")

    def test_hub_variant_slots(self):
        spokes = [{"label": str(i)} for i in range(3)]
        validate_component("hub", {"center": "c", "spokes": spokes}, "hub", 1)
        self.assertRejects("hub", {"center": "c", "spokes": spokes}, "flow",
                           "requires 'left'")

    def test_swimlane_lane_and_column_rules(self):
        lanes = ["a", "b"]
        self.assertRejects("swimlane", {"lanes": lanes, "steps": [
            {"lane": 1, "label": "x"}, {"lane": 3, "label": "y"}]},
            None, "lane")
        self.assertRejects("swimlane", {"lanes": lanes, "steps": [
            {"lane": 1, "label": "x", "col": 2},
            {"lane": 2, "label": "y", "col": 1}]}, None, "must not decrease")
        self.assertRejects("swimlane", {"lanes": lanes, "steps": [
            {"lane": 1, "label": "x", "col": 1},
            {"lane": 1, "label": "y", "col": 1}]}, None, "share lane")

    def test_action_plan_requires_consistent_done_when(self):
        rows = [{"issue": "i", "action": "a", "owner": "o", "due": "d",
                 "done_when": "w"},
                {"issue": "i", "action": "a", "owner": "o", "due": "d"}]
        self.assertRejects("action_plan", {"rows": rows}, None, "done_when")


class TestPolicyAndTargets(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.template = Path(cls.tmp.name) / "minimal.pptx"
        build_minimal_template().save(cls.template)
        cls.tmap = json.loads(MINIMAL_MAP.read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_prefer_library_overrides_template_component(self):
        self.assertEqual(resolve_component("table", self.tmap)[0],
                         "template")
        self.assertEqual(resolve_component("table", self.tmap,
                                           "library")[0], "library")
        with self.assertRaises(BuildError):
            resolve_component("bullets", self.tmap, "library")

    def test_text_only_slide_needs_a_reason(self):
        bullets = {"component": "bullets", "message": "m", "notes": "n",
                   "slots": {"body": ["a"]}}
        with self.assertRaises(BuildError) as ctx:
            validate_deck(_deck([bullets]), self.tmap)
        self.assertIn("text-only", str(ctx.exception))
        kept = dict(bullets, text_only_reason="引用文をそのまま示す")
        records = validate_deck(_deck([kept]), self.tmap)
        self.assertEqual(records[0]["slide"], 1)
        relaxed = dict(self.tmap, policy={"text_only": "warn"})
        self.assertEqual(len(check_visual_policy(_deck([bullets]), relaxed)),
                         1)

    def test_reason_on_visual_slide_is_rejected(self):
        slide = {"component": "kpi", "message": "m", "notes": "n",
                 "text_only_reason": "x",
                 "slots": {"items": [{"value": "1", "label": "a"}]}}
        with self.assertRaises(BuildError):
            validate_deck(_deck([slide]), self.tmap)

    def test_message_slide_cap(self):
        msg = {"component": "message", "message": "m", "notes": "n",
               "slots": {"statement": "s"}}
        with self.assertRaises(BuildError) as ctx:
            validate_deck(_deck([msg] * 3), self.tmap)
        self.assertIn("max_message_slides", str(ctx.exception))

    def test_google_slides_target_reports_charts_and_fonts(self):
        deck = _deck([{"component": "chart", "variant": "full",
                       "message": "m", "notes": "n",
                       "slots": {"chart": {"type": "column",
                                           "categories": ["a", "b"],
                                           "series": [{"name": "s",
                                                       "values": [1, 2]}]}}}])
        deck_path = Path(self.tmp.name) / "chart.json"
        deck_path.write_text(json.dumps(deck, ensure_ascii=False),
                             encoding="utf-8")
        out = Path(self.tmp.name) / "chart.pptx"
        r = _run("build.py", "--deck", str(deck_path), "--map",
                 str(MINIMAL_MAP), "--template", str(self.template), "-o",
                 str(out), "--target", "google_slides")
        self.assertEqual(r.returncode, 0, r.stderr)
        report = json.loads(Path(str(out) + ".build-report.json").read_text(
            encoding="utf-8"))
        self.assertEqual(report["target"], "google_slides")
        messages = [c["message"] for c in report["compat"]]
        self.assertTrue(any("静止画" in m for m in messages))
        self.assertTrue(any("Yu Gothic" in m for m in messages))
        r = _run("build.py", "--deck", str(deck_path), "--map",
                 str(MINIMAL_MAP), "--template", str(self.template), "-o",
                 str(out))
        report = json.loads(Path(str(out) + ".build-report.json").read_text(
            encoding="utf-8"))
        self.assertEqual(report["compat"], [])

    def test_library_table_uses_explicit_fills(self):
        deck = _deck([{"component": "table", "prefer": "library",
                       "message": "m", "notes": "n",
                       "slots": {"header": ["a", "b"],
                                 "rows": [["1", {"text": "2",
                                                 "tone": "negative"}]]}}])
        deck_path = Path(self.tmp.name) / "table.json"
        deck_path.write_text(json.dumps(deck), encoding="utf-8")
        out = Path(self.tmp.name) / "table.pptx"
        r = _run("build.py", "--deck", str(deck_path), "--map",
                 str(MINIMAL_MAP), "--template", str(self.template), "-o",
                 str(out))
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(out))
        frame = next(sh for sh in prs.slides[0].shapes
                     if sh.name == "Library table")
        tbl = frame._element.graphic.graphicData.tbl
        self.assertIsNone(tbl.tblPr.find(qn("a:tableStyleId")))
        for tc in tbl.iter(qn("a:tc")):
            self.assertIsNotNone(tc.tcPr.find(qn("a:solidFill")))
        report = json.loads(Path(str(out) + ".build-report.json").read_text(
            encoding="utf-8"))
        self.assertTrue(report["library"][0]["forced"])


if __name__ == "__main__":
    unittest.main()
